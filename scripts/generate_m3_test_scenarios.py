"""Create reproducible, synthetic-only M3-A workbook scenarios and fixtures.

This script does not create product features or manipulate user workbooks. It
creates fixed test assets for moderated usability sessions, then scans them
with the same local scanner used by the API so the documented expectations
remain traceable to an actual result.
"""

from __future__ import annotations

import json
import sys
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
API_ROOT = ROOT / "apps" / "api"
SCENARIO_ROOT = ROOT / "samples" / "m3"
FIXTURE_ROOT = SCENARIO_ROOT / "fixtures"

sys.path.insert(0, str(API_ROOT))

from app.config import Settings  # noqa: E402
from app.recommendation_engine import enrich_scan_result  # noqa: E402
from app.scanner import scan_workbook  # noqa: E402


def save_workbook(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def build_low_or_zero_findings(path: Path) -> None:
    """Make a simple, formula-free workbook that has zero current findings."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Summary"
    sheet.append(["Month", "Revenue", "Status"])
    sheet.append(["January", 1200, "Synthetic test data"])
    sheet.append(["February", 1350, "Synthetic test data"])
    sheet.append(["March", 1280, "Synthetic test data"])
    save_workbook(workbook, path)


def build_revalidation_before(path: Path) -> None:
    """Make the first scan state: external link, broken reference, numeric text."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revalidation"
    sheet.append(["Amount", "Source", "Formula check"])
    sheet.append([1200, "='[Reference.xlsx]Data'!B2", "=#REF!+1"])
    sheet.append([1300, "Synthetic only", None])
    sheet.append([1400, "Synthetic only", None])
    sheet.append(["1500", "Synthetic only", None])
    save_workbook(workbook, path)


def build_revalidation_after(path: Path) -> None:
    """Make the re-scan state: two prior signals removed, one continues, one is new."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revalidation"
    sheet.append(["Amount", "Source", "Formula check", "New check"])
    sheet.append([1200, "Internal value", "=#REF!+1", '=INDIRECT("A2")'])
    sheet.append([1300, "Synthetic only", None, None])
    sheet.append([1400, "Synthetic only", None, None])
    sheet.append([1500, "Synthetic only", None, None])
    save_workbook(workbook, path)


def build_fixture(path: Path) -> dict[str, object]:
    result = enrich_scan_result(scan_workbook(path.name, path.read_bytes(), Settings()))
    return {
        "fixture_metadata": {
            "generated_at": datetime.now(UTC).isoformat(),
            "generated_from": path.relative_to(ROOT).as_posix(),
            "generator": "scripts/generate_m3_test_scenarios.py",
            "scanner_version": result.scanner_version,
            "rule_set_version": result.rule_set_version,
        },
        "scan_result": result.model_dump(mode="json"),
    }


def write_fixture(path: Path) -> None:
    fixture_path = FIXTURE_ROOT / f"{path.stem}.scan-result.json"
    fixture_path.parent.mkdir(parents=True, exist_ok=True)
    fixture_path.write_text(
        json.dumps(build_fixture(path), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Generated {fixture_path.relative_to(ROOT)}")


def main() -> None:
    scenarios = {
        "low-or-zero-findings.xlsx": build_low_or_zero_findings,
        "revalidation-before.xlsx": build_revalidation_before,
        "revalidation-after.xlsx": build_revalidation_after,
    }

    for filename, builder in scenarios.items():
        path = SCENARIO_ROOT / filename
        builder(path)
        print(f"Generated {path.relative_to(ROOT)}")
        write_fixture(path)


if __name__ == "__main__":
    main()
