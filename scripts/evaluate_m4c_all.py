"""Evaluate the two immutable synthetic M4-C packs as one final baseline.

This developer-only evaluator is intentionally fail-closed. It uses the
actual isolated formula-audit engine but writes no formula text or workbook
cell values to its JSON/CSV artifacts. Sheet/cell locations are retained only
because every input workbook is an approved synthetic test fixture.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from time import perf_counter
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries


ROOT = Path(__file__).resolve().parents[1]
API_ROOT = ROOT / "apps" / "api"
sys.path.insert(0, str(API_ROOT))

from app.config import Settings, formula_audit_is_available  # noqa: E402
from app.m4_release import (  # noqa: E402
    M4_C_PRODUCT_OWNER_WAIVER_ID,
    M4_C_WAIVED_SOURCE_CONFLICTS,
    M4_FORMULA_AUDIT_RELEASE_CANDIDATE_VERSION,
)
from app.scanner import (  # noqa: E402
    FORMULA_AUDIT_RULE_SET_VERSION,
    RULE_SET_VERSION,
    SCANNER_VERSION,
    run_formula_audit,
    scan_workbook,
)


AUDIT_SETTINGS = Settings(
    app_env="internal_beta",
    formula_pattern_audit_enabled=True,
    scan_cell_limit=250_000,
    finding_limit=5_000,
)
DEFAULT_OUTPUT_DIRECTORY = ROOT / "artifacts" / "m4c-final"
PACK_ROOTS = (
    ROOT
    / "samples"
    / "WorkbookCare_M4C_Sample_Pack_2026-09-01"
    / "WorkbookCare_M4C_Sample_Pack",
    ROOT
    / "samples"
    / "WorkbookCare_M4C_Additional_Pack_2026-09-01"
    / "WorkbookCare_M4C_Additional_Pack",
)
ALLOWED_RULE_CODES = frozenset({"FORMULA_PATTERN_OUTLIER", "FORMULA_PATTERN_GAP"})


@dataclass(frozen=True, slots=True)
class CandidateKey:
    file: str
    sheet: str | None
    cell: str | None
    rule_code: str
    subtype: str | None

    @property
    def location_key(self) -> tuple[str, str | None, str | None]:
        return (self.file, self.sheet, self.cell)

    @property
    def top_rule_key(self) -> tuple[str, str | None, str | None, str]:
        return (self.file, self.sheet, self.cell, self.rule_code)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _manifest_path(pack_root: Path) -> Path:
    matches = sorted((pack_root / "expected").glob("*manifest.json"))
    if len(matches) != 1:
        raise RuntimeError(f"Expected one manifest in {pack_root}, found {len(matches)}")
    return matches[0]


def _check_checksums(pack_root: Path) -> list[str]:
    checksum_file = pack_root / "CHECKSUMS.sha256"
    if not checksum_file.is_file():
        return ["CHECKSUMS.sha256 is missing"]

    failures: list[str] = []
    for line in checksum_file.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        expected_hash, relative_path = line.split(maxsplit=1)
        target = pack_root / relative_path.strip()
        if not target.is_file():
            failures.append(f"missing checksum target: {relative_path.strip()}")
            continue
        if _sha256(target) != expected_hash.casefold():
            failures.append(f"checksum mismatch: {relative_path.strip()}")
    return failures


def _candidate_key(filename: str, finding: Any) -> CandidateKey:
    evidence = finding.formula_pattern
    return CandidateKey(
        file=filename,
        sheet=finding.sheet,
        cell=finding.cell,
        rule_code=finding.rule_code,
        subtype=evidence.pattern_subtype if evidence else None,
    )


def _expected_candidate_key(item: dict[str, Any]) -> CandidateKey:
    return CandidateKey(
        file=item["file"],
        sheet=item["sheet"],
        cell=item["cell"],
        rule_code=item["expected_rule"],
        subtype=item.get("expected_subtype"),
    )


def _is_in_range(cell: str | None, cell_or_range: str) -> bool:
    if cell is None:
        return False
    min_column, min_row, max_column, max_row = range_boundaries(cell_or_range)
    row, column = coordinate_to_tuple(cell)
    return min_row <= row <= max_row and min_column <= column <= max_column


def _base_contract(result: Any) -> dict[str, Any]:
    return {
        "workbook": result.workbook.model_dump(),
        "summary": result.summary.model_dump(),
        "quote": result.quote.model_dump(),
        "limitations": result.limitations,
        "findings": [finding.model_dump(exclude={"id"}) for finding in result.findings],
    }


def _formula_text_exposure_count(payload: str, workbook: Any) -> int:
    formula_texts = {
        cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }
    return sum(formula in payload for formula in formula_texts)


def _forbidden_raw_field_count(value: Any) -> int:
    if isinstance(value, dict):
        forbidden = {"formula", "formula_text", "cell_value", "raw_formula", "raw_value"}
        return sum(key.casefold() in forbidden for key in value) + sum(
            _forbidden_raw_field_count(child) for child in value.values()
        )
    if isinstance(value, list):
        return sum(_forbidden_raw_field_count(child) for child in value)
    return 0


def _csv_candidate_keys(path: Path) -> set[CandidateKey]:
    # Pack CSV files are UTF-8 with a BOM. utf-8-sig preserves Korean text
    # and normalizes the first header to `file`.
    with path.open(encoding="utf-8-sig", newline="") as source:
        return {
            CandidateKey(
                file=row["file"],
                sheet=row["sheet"],
                cell=row["cell"],
                rule_code=row["expected_rule"],
                subtype=row.get("expected_subtype"),
            )
            for row in csv.DictReader(source)
        }


def _csv_normal_keys(path: Path) -> set[tuple[str, str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as source:
        return {
            (row["file"], row["sheet"], row["cell"])
            for row in csv.DictReader(source)
        }


def _write_csv(path: Path, file_results: list[dict[str, Any]]) -> None:
    fieldnames = [
        "pack",
        "file",
        "usage",
        "audit_status",
        "expected_positive_count",
        "actual_candidate_count",
        "location_true_positive_count",
        "top_rule_match_count",
        "subtype_match_count",
        "unexpected_candidate_count",
        "normal_exception_false_positive_count",
        "first_run_ms",
        "second_run_ms",
        "finding_keys_stable",
        "baseline_unchanged",
        "formula_text_exposure_count",
        "raw_value_field_count",
        "is_clean_control",
    ]
    with path.open("w", encoding="utf-8", newline="") as destination:
        writer = csv.DictWriter(destination, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(file_results)


def _conflict_key(conflict: dict[str, str]) -> tuple[str, str, str, str, str]:
    return (
        conflict["pack"],
        conflict["file"],
        conflict["sheet"],
        conflict["cell"],
        conflict["normal_range"],
    )


def evaluate(
    output_directory: Path,
    *,
    accept_product_owner_fixture_waiver: bool = False,
) -> tuple[dict[str, Any], int]:
    output_directory.mkdir(parents=True, exist_ok=True)
    started_at = perf_counter()
    file_results: list[dict[str, Any]] = []
    validation_failures: list[str] = []
    fixture_conflicts: list[dict[str, str]] = []
    expected_location_keys: set[tuple[str, str | None, str | None]] = set()
    expected_top_rule_keys: set[tuple[str, str | None, str | None, str]] = set()
    expected_full_keys: set[CandidateKey] = set()
    actual_location_keys: set[tuple[str, str | None, str | None]] = set()
    actual_top_rule_keys: set[tuple[str, str | None, str | None, str]] = set()
    actual_full_keys: set[CandidateKey] = set()
    normal_false_positive_keys: set[CandidateKey] = set()
    clean_control_results: list[dict[str, Any]] = []
    pack_results: list[dict[str, Any]] = []

    default_settings = Settings()
    feature_flag_status = {
        "default_formula_pattern_audit_enabled": default_settings.formula_pattern_audit_enabled,
        "default_formula_audit_available": formula_audit_is_available(default_settings),
        "internal_formula_pattern_audit_enabled": AUDIT_SETTINGS.formula_pattern_audit_enabled,
        "internal_formula_audit_available": formula_audit_is_available(AUDIT_SETTINGS),
    }
    if feature_flag_status["default_formula_pattern_audit_enabled"]:
        validation_failures.append("formula audit feature flag is enabled by default")
    if feature_flag_status["default_formula_audit_available"]:
        validation_failures.append("formula audit is available in the default environment")
    if not feature_flag_status["internal_formula_audit_available"]:
        validation_failures.append("formula audit is unavailable in internal_beta")

    for pack_root in PACK_ROOTS:
        manifest_path = _manifest_path(pack_root)
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        checksum_failures_before = _check_checksums(pack_root)
        if checksum_failures_before:
            validation_failures.extend(
                f"{pack_root.name}: {failure}" for failure in checksum_failures_before
            )

        expected = [_expected_candidate_key(item) for item in manifest["expected_findings"]]
        expected_full_keys.update(expected)
        expected_location_keys.update(item.location_key for item in expected)
        expected_top_rule_keys.update(item.top_rule_key for item in expected)
        normal_exceptions = manifest["normal_exceptions"]
        # A source fixture cannot simultaneously call the same location an
        # expected positive and a normal exception. Detect this from the
        # immutable source labels before running the engine, so a future
        # detector regression cannot hide a fixture-contract problem.
        for expected_item in expected:
            for normal_item in normal_exceptions:
                if (
                    expected_item.file == normal_item["file"]
                    and expected_item.sheet == normal_item["sheet"]
                    and _is_in_range(expected_item.cell, normal_item["cell"])
                ):
                    fixture_conflicts.append(
                        {
                            "pack": pack_root.name,
                            "file": expected_item.file,
                            "sheet": expected_item.sheet or "",
                            "cell": expected_item.cell or "",
                            "normal_range": normal_item["cell"],
                        }
                    )
        expected_csv = next((pack_root / "expected").glob("*expected_findings.csv"))
        normal_csv = next((pack_root / "expected").glob("*normal_exceptions.csv"))
        expected_artifacts_match = (
            expected_csv is not None
            and normal_csv is not None
            and _csv_candidate_keys(expected_csv) == set(expected)
            and _csv_normal_keys(normal_csv)
            == {(item["file"], item["sheet"], item["cell"]) for item in normal_exceptions}
        )
        if not expected_artifacts_match:
            validation_failures.append(f"{pack_root.name}: manifest and CSV labels differ")

        pack_actual: set[CandidateKey] = set()
        for scenario in manifest["scenarios"]:
            filename = scenario["file"]
            payload = (pack_root / "samples" / filename).read_bytes()
            expected_for_file = {item for item in expected if item.file == filename}
            normal_for_file = [
                item for item in normal_exceptions if item["file"] == filename
            ]

            before_base = scan_workbook(filename, payload, Settings())
            with_flag_base = scan_workbook(filename, payload, AUDIT_SETTINGS)
            first_started_at = perf_counter()
            first_audit = run_formula_audit(filename, payload, AUDIT_SETTINGS)
            first_run_ms = round((perf_counter() - first_started_at) * 1000, 3)
            second_started_at = perf_counter()
            second_audit = run_formula_audit(filename, payload, AUDIT_SETTINGS)
            second_run_ms = round((perf_counter() - second_started_at) * 1000, 3)
            after_base = scan_workbook(filename, payload, Settings())

            first_candidates = {_candidate_key(filename, item) for item in first_audit.candidates}
            second_candidates = {_candidate_key(filename, item) for item in second_audit.candidates}
            first_finding_keys = sorted(item.finding_key for item in first_audit.candidates)
            second_finding_keys = sorted(item.finding_key for item in second_audit.candidates)
            finding_keys_stable = first_finding_keys == second_finding_keys
            baseline_unchanged = (
                _base_contract(before_base) == _base_contract(with_flag_base)
                and _base_contract(before_base) == _base_contract(after_base)
                and not any(
                    finding.rule_code.startswith("FORMULA_PATTERN_")
                    for finding in before_base.findings
                )
                and not any(
                    finding.rule_code.startswith("FORMULA_PATTERN_")
                    for finding in after_base.findings
                )
            )
            serialized_audit = first_audit.model_dump_json()
            # The serialized response must not contain an original formula
            # string. Load only this synthetic fixture in-memory for the
            # comparison; neither values nor formulas are written out.
            privacy_workbook = load_workbook(
                BytesIO(payload), read_only=True, data_only=False
            )
            try:
                formula_text_exposure_count = _formula_text_exposure_count(
                    serialized_audit, privacy_workbook
                )
            finally:
                privacy_workbook.close()
            raw_value_field_count = _forbidden_raw_field_count(first_audit.model_dump())

            if first_audit.status != "COMPLETED" or second_audit.status != "COMPLETED":
                validation_failures.append(
                    f"{pack_root.name}/{filename}: audit did not complete"
                )
            if not finding_keys_stable:
                validation_failures.append(
                    f"{pack_root.name}/{filename}: finding keys are unstable"
                )
            if not baseline_unchanged:
                validation_failures.append(
                    f"{pack_root.name}/{filename}: base diagnosis changed after audit"
                )
            if formula_text_exposure_count or raw_value_field_count:
                validation_failures.append(
                    f"{pack_root.name}/{filename}: raw formula/value exposure"
                )
            if any(item.rule_code not in ALLOWED_RULE_CODES for item in first_candidates):
                validation_failures.append(
                    f"{pack_root.name}/{filename}: unsupported rule code emitted"
                )

            expected_top_rules_for_file = {
                item.top_rule_key for item in expected_for_file
            }
            for candidate in first_candidates:
                for normal in normal_for_file:
                    if candidate.sheet != normal["sheet"] or not _is_in_range(
                        candidate.cell, normal["cell"]
                    ):
                        continue
                    if candidate.top_rule_key not in expected_top_rules_for_file:
                        normal_false_positive_keys.add(candidate)

            actual_location_keys.update(item.location_key for item in first_candidates)
            actual_top_rule_keys.update(item.top_rule_key for item in first_candidates)
            actual_full_keys.update(first_candidates)
            pack_actual.update(first_candidates)
            is_clean_control = len(expected_for_file) == 0
            if is_clean_control:
                clean_control_results.append(
                    {
                        "file": filename,
                        "audit_status": first_audit.status,
                        "candidate_count": len(first_candidates),
                    }
                )

            file_results.append(
                {
                    "pack": pack_root.name,
                    "file": filename,
                    "usage": scenario.get("usage", ""),
                    "audit_status": first_audit.status,
                    "expected_positive_count": len(expected_for_file),
                    "actual_candidate_count": len(first_candidates),
                    "location_true_positive_count": len(
                        {item.location_key for item in first_candidates}
                        & {item.location_key for item in expected_for_file}
                    ),
                    "top_rule_match_count": len(
                        {item.top_rule_key for item in first_candidates}
                        & {item.top_rule_key for item in expected_for_file}
                    ),
                    "subtype_match_count": len(first_candidates & expected_for_file),
                    "unexpected_candidate_count": len(first_candidates - expected_for_file),
                    "normal_exception_false_positive_count": len(
                        {
                            item
                            for item in normal_false_positive_keys
                            if item.file == filename
                        }
                    ),
                    "first_run_ms": first_run_ms,
                    "second_run_ms": second_run_ms,
                    "finding_keys_stable": finding_keys_stable,
                    "baseline_unchanged": baseline_unchanged,
                    "formula_text_exposure_count": formula_text_exposure_count,
                    "raw_value_field_count": raw_value_field_count,
                    "is_clean_control": is_clean_control,
                }
            )

        checksum_failures_after = _check_checksums(pack_root)
        if checksum_failures_after:
            validation_failures.extend(
                f"{pack_root.name}: input changed during evaluation: {failure}"
                for failure in checksum_failures_after
            )
        pack_results.append(
            {
                "pack": pack_root.name,
                "manifest": manifest_path.name,
                "manifest_version": manifest.get("version"),
                "scenario_count": len(manifest["scenarios"]),
                "expected_positive_count": len(expected),
                "actual_candidate_count": len(pack_actual),
                "normal_exception_label_count": len(normal_exceptions),
                "checksum_failures_before": checksum_failures_before,
                "checksum_failures_after": checksum_failures_after,
                "expected_artifacts_match": expected_artifacts_match,
            }
        )

    missing_locations = expected_location_keys - actual_location_keys
    missing_top_rules = expected_top_rule_keys - actual_top_rule_keys
    missing_subtypes = expected_full_keys - actual_full_keys
    unexpected_top_rules = actual_top_rule_keys - expected_top_rule_keys
    if missing_locations:
        validation_failures.append(f"missing locations: {len(missing_locations)}")
    if missing_top_rules:
        validation_failures.append(f"missing top-level rules: {len(missing_top_rules)}")
    if unexpected_top_rules:
        validation_failures.append(f"unexpected candidates: {len(unexpected_top_rules)}")
    if normal_false_positive_keys:
        validation_failures.append(
            f"normal-exception false positives: {len(normal_false_positive_keys)}"
        )
    if len(clean_control_results) != 2 or any(
        item["audit_status"] != "COMPLETED" or item["candidate_count"] != 0
        for item in clean_control_results
    ):
        validation_failures.append("clean-control candidate/status check failed")
    if fixture_conflicts:
        validation_failures.append(
            f"source manifest has {len(fixture_conflicts)} target/normal label conflicts"
        )

    source_conflict_failure = (
        f"source manifest has {len(fixture_conflicts)} target/normal label conflicts"
    )
    actual_conflict_keys = frozenset(_conflict_key(item) for item in fixture_conflicts)
    waiver_applies = (
        accept_product_owner_fixture_waiver
        and bool(fixture_conflicts)
        and actual_conflict_keys == M4_C_WAIVED_SOURCE_CONFLICTS
    )
    unwaived_validation_failures = list(validation_failures)
    if waiver_applies:
        unwaived_validation_failures = [
            failure
            for failure in validation_failures
            if failure != source_conflict_failure
        ]
    completion_status = (
        "PASS_WITH_PRODUCT_OWNER_WAIVER"
        if waiver_applies and not unwaived_validation_failures
        else "PASS"
        if not validation_failures
        else "BLOCKED"
    )

    result = {
        "evaluation_kind": "M4-C synthetic final validation",
        "evaluated_at_utc": datetime.now(UTC).isoformat(),
        "git": {"repository": False, "commit_hash": None},
        "release_candidate_version": M4_FORMULA_AUDIT_RELEASE_CANDIDATE_VERSION,
        "versions": {
            "scanner_version": SCANNER_VERSION,
            "base_rule_set_version": RULE_SET_VERSION,
            "formula_audit_rule_set_version": FORMULA_AUDIT_RULE_SET_VERSION,
        },
        "feature_flag_status": feature_flag_status,
        "pack_results": pack_results,
        "metrics": {
            "file_count": len(file_results),
            "expected_finding_count": len(expected_full_keys),
            "location_true_positive_count": len(
                expected_location_keys & actual_location_keys
            ),
            "location_false_negative_count": len(missing_locations),
            "top_rule_match_count": len(expected_top_rule_keys & actual_top_rule_keys),
            "top_rule_false_negative_count": len(missing_top_rules),
            "subtype_match_count": len(expected_full_keys & actual_full_keys),
            "subtype_mismatch_or_missing_count": len(missing_subtypes),
            "unexpected_candidate_count": len(unexpected_top_rules),
            "normal_exception_label_count": sum(
                item["normal_exception_label_count"] for item in pack_results
            ),
            "normal_exception_false_positive_count": len(normal_false_positive_keys),
            "source_manifest_target_normal_conflict_count": len(fixture_conflicts),
            "scan_failure_count": sum(
                item["audit_status"] != "COMPLETED" for item in file_results
            ),
            "clean_control_file_count": len(clean_control_results),
            "clean_control_candidate_count": sum(
                item["candidate_count"] for item in clean_control_results
            ),
            "finding_key_instability_count": sum(
                not item["finding_keys_stable"] for item in file_results
            ),
            "baseline_regression_count": sum(
                not item["baseline_unchanged"] for item in file_results
            ),
            "formula_text_exposure_count": sum(
                item["formula_text_exposure_count"] for item in file_results
            ),
            "raw_value_field_count": sum(
                item["raw_value_field_count"] for item in file_results
            ),
            "total_evaluation_ms": round((perf_counter() - started_at) * 1000, 3),
        },
        "clean_controls": clean_control_results,
        "source_manifest_conflicts": fixture_conflicts,
        "file_results": file_results,
        "validation_failures": validation_failures,
        "unwaived_validation_failures": unwaived_validation_failures,
        "product_owner_fixture_waiver": {
            "requested": accept_product_owner_fixture_waiver,
            "waiver_id": M4_C_PRODUCT_OWNER_WAIVER_ID
            if accept_product_owner_fixture_waiver
            else None,
            "applies": waiver_applies,
            "accepted_conflict_count": len(fixture_conflicts) if waiver_applies else 0,
            "note": (
                "The source labels remain contradictory and are not modified. "
                "Only the reviewed conflicts are waived."
                if waiver_applies
                else None
            ),
        },
        "completion_status": completion_status,
        "limits": [
            "Synthetic fixtures measure regression behaviour, not production accuracy.",
            "The audit is static only and does not calculate formulas or infer business correctness.",
            "No formula text or workbook cell values are stored in these artifacts.",
        ],
    }
    return result, 0 if not unwaived_validation_failures else 1


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIRECTORY,
        help="Directory for evaluation.json and evaluation.csv.",
    )
    parser.add_argument(
        "--accept-product-owner-fixture-waiver",
        action="store_true",
        help=(
            "Accept only the documented 2026-09-02 source-label conflict waiver. "
            "Any other failure still exits non-zero."
        ),
    )
    arguments = parser.parse_args()
    result, exit_code = evaluate(
        arguments.output_dir,
        accept_product_owner_fixture_waiver=arguments.accept_product_owner_fixture_waiver,
    )
    (arguments.output_dir / "evaluation.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    _write_csv(arguments.output_dir / "evaluation.csv", result["file_results"])
    print(
        f"M4-C final evaluation: {result['completion_status']} "
        f"({result['metrics']['location_true_positive_count']}/"
        f"{result['metrics']['expected_finding_count']} location targets, "
        f"{result['metrics']['top_rule_match_count']}/"
        f"{result['metrics']['expected_finding_count']} top-level rules)"
    )
    print(f"Artifacts: {arguments.output_dir}")
    for failure in result["unwaived_validation_failures"]:
        print(f"- {failure}")
    if result["product_owner_fixture_waiver"]["applies"]:
        print(
            "- product-owner fixture waiver applied: "
            f"{result['product_owner_fixture_waiver']['waiver_id']}"
        )
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
