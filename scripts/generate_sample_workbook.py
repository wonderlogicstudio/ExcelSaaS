"""Generate a synthetic workbook for local WorkbookCare testing.

The file contains no real customer or company data.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


def build_sample(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sales = workbook.active
    sales.title = "월별매출"
    sales.append(["월", "매출", "목표", "달성률", "외부예산"])

    for month, revenue, target in [
        ("1월", 12_500_000, 13_000_000),
        ("2월", 14_200_000, 13_500_000),
        ("3월", 13_100_000, 14_000_000),
        ("4월", 16_800_000, 15_000_000),
    ]:
        row = sales.max_row + 1
        sales.append(
            [
                month,
                revenue,
                target,
                f"=IF(B{row}/C{row}>1,1,B{row}/C{row})",
                f"='[Budget.xlsx]Plan'!B{row}",
            ]
        )

    sales.append(["5월", "15,900,000", 16_000_000, "=B6/C6", "='[Budget.xlsx]Plan'!B6"] )

    sales["B8"] = "=#REF!+100"
    sales["B9"] = '=INDIRECT("B2")'
    sales["B10"] = "=SUM(B:B)"
    sales.merge_cells("A12:C12")
    sales["A12"] = "병합 셀 예시"

    chart = BarChart()
    chart.title = "월별 매출"
    chart.add_data(Reference(sales, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(sales, min_col=1, min_row=2, max_row=5))
    sales.add_chart(chart, "G2")

    hidden = workbook.create_sheet("숨김기준정보")
    hidden.sheet_state = "hidden"
    hidden.append(["구분", "값"])
    hidden.append(["세율", 0.1])

    very_hidden = workbook.create_sheet("SystemCache")
    very_hidden.sheet_state = "veryHidden"
    very_hidden["A1"] = "Synthetic test only"

    workbook.save(path)
    workbook.close()


if __name__ == "__main__":
    destination = Path(__file__).resolve().parents[1] / "samples" / "demo-risky-workbook.xlsx"
    build_sample(destination)
    print(f"Created: {destination}")
