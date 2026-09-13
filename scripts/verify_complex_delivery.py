"""Execute real repair/comparison services against frozen synthetic expectations."""

from pathlib import Path
import argparse
import base64
import hashlib
from io import BytesIO
import json
import tempfile
import time

import verify_complex_validation as cases
from openpyxl import load_workbook

from app.comparison_service import (
    create_comparison,
    prepare_comparison,
    execute_comparison,
    download_comparison,
)
from app.delivery_execution import approve, execute, download
from app.delivery_hosted_rehearsal import grant
from app.delivery_inputs import inspect_input
from app.delivery_plan import build_plan
from app.delivery_store import DeliveryStore
from app.errors import WorkbookCareError


def hosted_settings():
    return cases.settings().model_copy(
        update={
            "app_env": "hosted_beta",
            "delivery_beta_enabled": True,
            "hosted_synthetic_delivery_enabled": True,
            "control_plane_hmac_secret": "synthetic-complex-proof-key-0000000000000",
        }
    )


def check_repaired(payload, case):
    with BytesIO(payload) as stream:
        book = load_workbook(stream, data_only=True)
        for address, expected in case["expected"].items():
            actual = book[case["sheet"]][address].value
            if expected["type"] == "blank":
                assert actual is None
            elif expected["type"] == "text":
                assert isinstance(actual, str) and actual == expected["value"]
            else:
                assert not isinstance(actual, str) and actual == expected["value"], (
                    address
                )
        book.close()
    with BytesIO(payload) as stream:
        book = load_workbook(stream, data_only=False)
        for address, formula in case["restored_formulas"].items():
            assert book[case["sheet"]][address].value == formula
        book.close()


def save_download(artifact, out):
    payload = base64.b64decode(artifact["file_base64"])
    name = Path(artifact["filename"]).name
    assert name == artifact["filename"]
    if out:
        out.mkdir(parents=True, exist_ok=True)
        target = out / name
        assert not target.exists()
        target.write_bytes(payload)
    return payload, {
        "filename": name,
        "sha256": hashlib.sha256(payload).hexdigest(),
        "bytes": len(payload),
    }


def repair(label, out=None):
    case = cases.oracle()["repair"][label]
    payload = (cases.PACK / case["file"]).read_bytes()
    settings = hosted_settings()
    with tempfile.TemporaryDirectory(prefix="workbookcare-complex-delivery-") as temp:
        store = DeliveryStore(Path(temp))
        job = store.create(
            "complex-owner",
            payload,
            inspect_input(case["file"], payload, settings),
            "complex-" + label,
        )
        policy = cases.policy(case)
        plan = build_plan(job, policy)
        job = store.update(
            job,
            job["revision"],
            {
                **job["state"],
                "policy": policy,
                "plan": plan,
                "status": "PREVIEW_VALIDATED",
            },
        )
        job = grant(
            store,
            job,
            {
                "revision": job["revision"],
                "acknowledge_no_payment": True,
                "acknowledge_limitations": True,
            },
            settings,
        )
        assert (
            not job["state"].get("approval")
            and job["state"]["payment"] == "NOT_PURCHASED"
        )
        try:
            execute(store, job, settings)
        except WorkbookCareError as error:
            assert error.code == "CURRENT_APPROVAL_REQUIRED"
        else:
            raise AssertionError(
                "Grant incorrectly allowed execution without exact approval"
            )
        job = approve(
            store,
            job,
            {
                "revision": job["revision"],
                "plan_digest": plan["digest"],
                "candidate_ids": [p["candidate_id"] for p in plan["patches"]],
                "acknowledge_exact_changes": True,
            },
            settings,
        )
        ready = execute(store, job, settings)
        assert ready["state"]["status"] == "READY"
        artifacts = []
        for kind in ["REPAIRED_XLSX", "CHANGES_XLSX", "VERIFICATION_HTML"]:
            artifact = download(store, ready, kind, settings)
            assert artifact == download(store, ready, kind, settings)
            data, record = save_download(artifact, out)
            if kind == "REPAIRED_XLSX":
                check_repaired(data, case)
            if kind == "CHANGES_XLSX":
                book = load_workbook(BytesIO(data), data_only=False)
                assert book["변경 셀"].max_row - 1 == case["patch_count"]
                assert {
                    row[1]
                    for row in book["변경 셀"].iter_rows(min_row=2, values_only=True)
                } == set(case["targets"])
                book.close()
            artifacts.append({"kind": kind, **record})
        assert store.load("complex-owner", job["id"])["source"] == payload
        assert not ready["state"].get("order_id")
        return {
            "case": label,
            "status": "PASS",
            "exact_cells": len(case["expected"]),
            "sum_after": case["sum_after"],
            "patches": case["patch_count"],
            "approval_required": True,
            "source_unchanged": True,
            "artifacts": artifacts,
        }


def comparison(out=None, capacity=False):
    expected = cases.oracle()["comparison"]
    if capacity:
        name = cases.oracle()["capacity"]["supported_file"]
        files = {"A": name, "B": name}
        controls = {"A": {"rows": 1000}, "B": {"rows": 1000}}
    else:
        files, controls = expected["files"], expected["controls"]
    settings = hosted_settings()
    body = {
        "request_key": "complex-comparison-fixed",
        "consent": True,
        "sources": {
            side: {
                "filename": name,
                "file_base64": base64.b64encode(
                    (cases.PACK / name).read_bytes()
                ).decode(),
            }
            for side, name in files.items()
        },
    }
    with tempfile.TemporaryDirectory(prefix="workbookcare-complex-comparison-") as temp:
        store = DeliveryStore(Path(temp))
        job = create_comparison(store, "complex-owner", body, settings)
        job = prepare_comparison(
            store,
            job,
            {
                "revision": job["revision"],
                "source_pair_hash": job["snapshot"]["source_pair_hash"],
                "spec": cases.comparison_spec(controls, capacity),
            },
            settings,
        )
        job = grant(
            store,
            job,
            {
                "revision": job["revision"],
                "acknowledge_no_payment": True,
                "acknowledge_limitations": True,
            },
            settings,
        )
        ready = execute_comparison(
            store, job, {"acknowledge_limitations": True}, settings
        )
        assert ready["state"]["status"] == "READY"
        summary = ready["state"]["comparison_result"]["summary"]
        if capacity:
            assert summary["counts"]["MATCHED"] == 1000
        else:
            assert (
                summary["counts"] == expected["counts"]
                and summary["group_count"] == 464
            )
        artifacts = []
        for kind in ["COMPARISON_REPORT_XLSX", "COMPARISON_VERIFICATION_HTML"]:
            data, record = save_download(
                download_comparison(store, ready, kind, settings), out
            )
            if kind == "COMPARISON_REPORT_XLSX":
                book = load_workbook(BytesIO(data), data_only=False)
                rows = [
                    row
                    for name in ["금액차이", "한쪽자료", "중복모호", "자료오류", "일치"]
                    for row in book[name].iter_rows(min_row=2, values_only=True)
                    if row[3]
                ]
                assert len(rows) == (2000 if capacity else 888)
                assert len({row[3] for row in rows}) == len(rows)
                if not capacity:
                    for side in ["A", "B"]:
                        for source in expected["physical_rows"][side]:
                            matches = [
                                r
                                for r in rows
                                if r[2] == side
                                and str(r[5]) == str(source["physical_row"])
                            ]
                            assert (
                                len(matches) == 1 and matches[0][1] == source["status"]
                            )
                    assert {r[10] for r in rows if "BIG" in r[6]} == {"1"}
                    for side, column in [("A", "C12"), ("B", "D12")]:
                        assert (
                            book["요약"][column].value
                            == expected["controls"][side]["known_amount"]
                        )
                book.close()
            artifacts.append({"kind": kind, **record})
        assert ready["state"]["payment"] == "NOT_PURCHASED" and not ready["state"].get(
            "order_id"
        )
        return {
            "case": "comparison-capacity" if capacity else "comparison",
            "status": "PASS",
            "summary": summary,
            "all_physical_rows_verified": True,
            "artifacts": artifacts,
        }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    assert not (args.out / "execution.json").exists()
    results = []
    for label in [
        "RP01_ALL",
        "RP01_SUBSET",
        "RP02_ALL",
        "comparison",
        "comparison-capacity",
    ]:
        start = time.monotonic()
        try:
            result = (
                comparison(args.out / label, label.endswith("capacity"))
                if label.startswith("comparison")
                else repair(label, args.out / label)
            )
        except Exception as error:
            result = {
                "case": label,
                "status": "FAIL",
                "type": type(error).__name__,
                "code": getattr(error, "code", None),
                "detail": str(error)[:800],
            }
        result["seconds"] = round(time.monotonic() - start, 3)
        results.append(result)
        print(
            json.dumps({k: result[k] for k in ["case", "status", "seconds"]}),
            flush=True,
        )
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "execution.json").write_text(
        json.dumps(
            {
                "kind": "ACTUAL_LOCAL_SERVICES_ENGINE_ARTIFACTS",
                "official_pg_verified": False,
                "results": results,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return 0 if all(r["status"] == "PASS" for r in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
