"""Generate the synthetic-only M4-A.5 Formula Audit Quality Gate corpus.

The corpus is intentionally made with simple, supported formulas. It is a
deterministic regression and quality gate for the narrow M4-A engine, not a
claim about formula correctness in arbitrary business workbooks.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
CORPUS_ROOT = ROOT / "samples" / "m4-evaluation"


def _save(workbook: Workbook, filename: str) -> None:
    CORPUS_ROOT.mkdir(parents=True, exist_ok=True)
    workbook.save(CORPUS_ROOT / filename)
    workbook.close()


def _add_formula_region(
    sheet,
    *,
    output_column: str,
    input_column: str,
    formula_for_row,
) -> None:
    sheet[f"{input_column}1"] = "Synthetic input"
    sheet[f"{output_column}1"] = "Synthetic formula"
    for row in range(2, 7):
        sheet[f"{input_column}{row}"] = row * 10
        sheet[f"{output_column}{row}"] = formula_for_row(row)


def _add_three_regions(sheet, formula_for_row, replacements: tuple[object, object, object]) -> None:
    for output_column, input_column, replacement in zip(
        ("C", "G", "K"), ("B", "F", "J"), replacements, strict=True
    ):
        _add_formula_region(
            sheet,
            output_column=output_column,
            input_column=input_column,
            formula_for_row=lambda row, input_column=input_column: f"=SUM({input_column}{row})",
        )
        sheet[f"{output_column}4"] = replacement


def build_detected_patterns() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    function = workbook.create_sheet("FunctionDrift")
    _add_three_regions(function, None, ("=AVERAGE(B4)", "=MAX(F4)", "=MIN(J4)"))

    sheet_reference = workbook.create_sheet("SheetDrift")
    for output_column, input_column, source_name in zip(
        ("C", "G", "K"), ("B", "F", "J"), ("SourceA", "SourceC", "SourceE"), strict=True
    ):
        _add_formula_region(
            sheet_reference,
            output_column=output_column,
            input_column=input_column,
            formula_for_row=lambda row, source_name=source_name: f"=SUM({source_name}!B{row})",
        )
    sheet_reference["C4"] = "=SUM(SourceB!B4)"
    sheet_reference["G4"] = "=SUM(SourceD!B4)"
    sheet_reference["K4"] = "=SUM(SourceF!B4)"

    relative = workbook.create_sheet("RelativeDrift")
    _add_three_regions(relative, None, ("=SUM(B3)", "=SUM(F5)", "=SUM(I4)"))

    absolute = workbook.create_sheet("AbsoluteDrift")
    for output_column, input_column in zip(("C", "G", "K"), ("B", "F", "J"), strict=True):
        _add_formula_region(
            absolute,
            output_column=output_column,
            input_column=input_column,
            formula_for_row=lambda row, input_column=input_column: f"=SUM(${input_column}{row})",
        )
    absolute["C4"] = "=SUM(B4)"
    absolute["G4"] = "=SUM(F4)"
    absolute["K4"] = "=SUM(J4)"

    range_end = workbook.create_sheet("RangeDrift")
    for output_column, input_column in zip(("C", "G", "K"), ("B", "F", "J"), strict=True):
        _add_formula_region(
            range_end,
            output_column=output_column,
            input_column=input_column,
            formula_for_row=lambda row, input_column=input_column: (
                f"=SUM({input_column}{row}:{chr(ord(input_column) + 2)}{row})"
            ),
        )
    range_end["C4"] = "=SUM(B4:C4)"
    range_end["G4"] = "=SUM(F4:G4)"
    range_end["K4"] = "=SUM(J4:K4)"

    constant = workbook.create_sheet("ConstantGaps")
    _add_three_regions(constant, None, (901, 902, 903))

    blank = workbook.create_sheet("BlankGaps")
    _add_three_regions(blank, None, (None, None, None))

    _save(workbook, "detected-patterns.xlsx")


def build_mixed_candidates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mixed"
    for output_column in ("C", "D", "E"):
        _add_formula_region(
            sheet,
            output_column=output_column,
            input_column="B",
            formula_for_row=lambda row: f"=SUM(B{row})",
        )
    sheet["C4"] = "=AVERAGE(B4)"
    sheet["D4"] = 777
    sheet["E4"] = None
    _save(workbook, "mixed-candidates.xlsx")


def _add_normal_region(sheet, *, outlier: str | None = "=AVERAGE(B4)") -> None:
    _add_formula_region(
        sheet,
        output_column="C",
        input_column="B",
        formula_for_row=lambda row: f"=SUM(B{row})",
    )
    if outlier is not None:
        sheet["C4"] = outlier


def build_normal_exceptions() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    subtotal = workbook.create_sheet("Subtotal")
    _add_normal_region(subtotal)
    subtotal["A4"] = "소계"

    total = workbook.create_sheet("Total")
    _add_normal_region(total)
    total["A4"] = "총계"

    header = workbook.create_sheet("Header")
    _add_normal_region(header, outlier=None)
    header["C1"] = "Different heading"

    divider = workbook.create_sheet("BlankDivider")
    _add_normal_region(divider, outlier=None)
    divider["A4"] = None
    divider["B4"] = None
    divider["C4"] = None

    boundary = workbook.create_sheet("Boundaries")
    _add_normal_region(boundary, outlier=None)
    boundary["C2"] = "=AVERAGE(B2)"
    boundary["C6"] = "=AVERAGE(B6)"

    manual = workbook.create_sheet("ManualRow")
    _add_normal_region(manual)
    manual["A4"] = "Manual adjustment"

    monthly = workbook.create_sheet("MovingMonthly")
    _add_formula_region(
        monthly,
        output_column="E",
        input_column="B",
        formula_for_row=lambda row: f"=SUM(B{row}:D{row})",
    )

    merged = workbook.create_sheet("Merged")
    _add_normal_region(merged)
    merged.merge_cells("C4:D4")

    table_sheet = workbook.create_sheet("Table")
    _add_normal_region(table_sheet)
    table = Table(displayName="SyntheticFormulaTable", ref="B1:C6")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table_sheet.add_table(table)

    multiple_regions = workbook.create_sheet("MultipleRegions")
    _add_formula_region(
        multiple_regions,
        output_column="C",
        input_column="B",
        formula_for_row=lambda row: f"=SUM(B{row})",
    )
    for row in range(8, 13):
        multiple_regions[f"B{row}"] = row * 10
        multiple_regions[f"C{row}"] = f"=AVERAGE(B{row})"

    clean = workbook.create_sheet("Clean")
    _add_normal_region(clean, outlier=None)
    _save(workbook, "normal-exceptions.xlsx")


def build_unsupported() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Unsupported"
    _add_formula_region(
        sheet,
        output_column="C",
        input_column="B",
        formula_for_row=lambda row: f"=SUM(B{row},1)",
    )
    sheet["C4"] = "=AVERAGE(B4,2)"

    external = workbook.create_sheet("ExternalWorkbook")
    _add_formula_region(
        external,
        output_column="C",
        input_column="B",
        formula_for_row=lambda row: f"=SUM('[Synthetic.xlsx]Data'!B{row})",
    )
    external["C4"] = "=AVERAGE('[Synthetic.xlsx]Data'!B4)"

    named = workbook.create_sheet("NamedRange")
    named["B2"] = 10
    workbook.defined_names.add(DefinedName("SyntheticNamed", attr_text="'NamedRange'!$B$2"))
    for row in range(2, 7):
        named[f"C{row}"] = "=SUM(SyntheticNamed)"
    named["C4"] = "=AVERAGE(SyntheticNamed)"
    _save(workbook, "unsupported-syntax.xlsx")


def build_truncated() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Truncated"
    _add_normal_region(sheet)
    _save(workbook, "scan-truncated.xlsx")


def build_performance(filename: str, formula_count: int) -> None:
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "Performance"
    sheet.append(["Input A", "Input B", "Calculated"])
    for row in range(2, formula_count + 2):
        sheet.append([row, row + 1, f"=SUM(A{row}:B{row})"])
    _save(workbook, filename)


def _case(
    case_id: str,
    workbook_id: str,
    sheet: str,
    cell: str,
    expected_detection: str,
    expected_pattern_type: str | None,
    expected_action: str,
    *,
    rule_code: str | None = None,
    pattern_subtype: str | None = None,
    normal_exception_reason: str | None = None,
    notes: str,
) -> dict[str, object]:
    return {
        "case_id": case_id,
        "workbook_id": workbook_id,
        "sheet": sheet,
        "cell": cell,
        "expected_detection": expected_detection,
        "expected_pattern_type": expected_pattern_type,
        "expected_pattern_subtype": pattern_subtype,
        "expected_rule_code": rule_code,
        "expected_action": expected_action,
        "normal_exception_reason": normal_exception_reason,
        "notes": notes,
    }


def build_manifest() -> None:
    detected_columns = ("C4", "G4", "K4")
    cases: list[dict[str, object]] = []
    for sheet, expected_subtype in (
        ("FunctionDrift", "FUNCTION_PATTERN_DRIFT"),
        ("SheetDrift", "REFERENCE_SHEET_DRIFT"),
        ("RelativeDrift", "RELATIVE_REFERENCE_DRIFT"),
        ("AbsoluteDrift", "ABSOLUTE_REFERENCE_DRIFT"),
        ("RangeDrift", "RANGE_BOUNDARY_DRIFT"),
    ):
        for index, cell in enumerate(detected_columns, start=1):
            cases.append(
                _case(
                    f"{sheet.lower()}-{index}",
                    "detected-patterns",
                    sheet,
                    cell,
                    "SHOULD_DETECT",
                    "DOMINANT_NORMALIZED_PATTERN_OUTLIER",
                    "EMIT_CANDIDATE",
                    rule_code="FORMULA_PATTERN_OUTLIER",
                    pattern_subtype=expected_subtype,
                    notes=f"Expected safe subtype: {expected_subtype}",
                )
            )
    for sheet, subtype in (("ConstantGaps", "CONSTANT_OVERRIDE_CANDIDATE"), ("BlankGaps", "BLANK_GAP_CANDIDATE")):
        for index, cell in enumerate(detected_columns, start=1):
            cases.append(
                _case(
                    f"{sheet.lower()}-{index}",
                    "detected-patterns",
                    sheet,
                    cell,
                    "SHOULD_DETECT",
                    "FORMULA_GAP_OR_CONSTANT",
                    "EMIT_CANDIDATE",
                    rule_code="FORMULA_PATTERN_GAP",
                    pattern_subtype=subtype,
                    notes=f"Expected safe subtype: {subtype}",
                )
            )
    cases.extend(
        [
            _case(
                "mixed-outlier", "mixed-candidates", "Mixed", "C4", "SHOULD_DETECT",
                "DOMINANT_NORMALIZED_PATTERN_OUTLIER", "EMIT_CANDIDATE",
                rule_code="FORMULA_PATTERN_OUTLIER", pattern_subtype="FUNCTION_PATTERN_DRIFT",
                notes="Mixed candidate workbook.",
            ),
            _case(
                "mixed-constant", "mixed-candidates", "Mixed", "D4", "SHOULD_DETECT",
                "FORMULA_GAP_OR_CONSTANT", "EMIT_CANDIDATE",
                rule_code="FORMULA_PATTERN_GAP", pattern_subtype="CONSTANT_OVERRIDE_CANDIDATE",
                notes="Mixed candidate workbook.",
            ),
            _case(
                "mixed-blank", "mixed-candidates", "Mixed", "E4", "SHOULD_DETECT",
                "FORMULA_GAP_OR_CONSTANT", "EMIT_CANDIDATE",
                rule_code="FORMULA_PATTERN_GAP", pattern_subtype="BLANK_GAP_CANDIDATE",
                notes="Mixed candidate workbook.",
            ),
        ]
    )

    normal_cases = (
        ("subtotal", "Subtotal", "C4", "Summary row is explicitly excluded."),
        ("total", "Total", "C4", "Summary row is explicitly excluded."),
        ("header", "Header", "C1", "Header is not a formula candidate."),
        ("divider", "BlankDivider", "C4", "Fully blank divider has no context."),
        ("boundary-first", "Boundaries", "C2", "First formula is not an internal candidate."),
        ("boundary-last", "Boundaries", "C6", "Last formula is not an internal candidate."),
        ("manual", "ManualRow", "C4", "Manual-labelled row is excluded."),
        ("monthly", "MovingMonthly", "E4", "Moving relative references normalize consistently."),
        ("merged", "Merged", "C4", "Merged range is excluded."),
        ("table", "Table", "C4", "Excel Table range is excluded."),
        ("multiple-regions", "MultipleRegions", "C10", "Separate formula regions are not compared."),
        ("clean", "Clean", "C4", "No pattern anomaly exists."),
    )
    for case_id, sheet, cell, reason in normal_cases:
        action = "CLEAN_PASS" if case_id in {"monthly", "multiple-regions", "clean", "header"} else "SUPPRESS_EXCLUDED"
        cases.append(
            _case(
                case_id,
                "normal-exceptions",
                sheet,
                cell,
                "SHOULD_NOT_DETECT",
                None,
                action,
                normal_exception_reason=reason,
                notes="Synthetic normal or explicit exclusion.",
            )
        )

    for case_id, sheet in (
        ("constant-formula", "Unsupported"),
        ("external-workbook", "ExternalWorkbook"),
        ("named-range", "NamedRange"),
    ):
        cases.append(
            _case(
                case_id,
                "unsupported-syntax",
                sheet,
                "C4",
                "UNSUPPORTED",
                None,
                "SKIP_UNSUPPORTED",
                normal_exception_reason="M4 intentionally skips this formula syntax.",
                notes="Unsupported structures must not become M4 candidates.",
            )
        )
    cases.append(
        _case(
            "truncated-scan",
            "scan-truncated",
            "Truncated",
            "C4",
            "SHOULD_NOT_DETECT",
            None,
            "SKIP_TRUNCATED",
            normal_exception_reason="M4-A is skipped for a truncated scan.",
            notes="This is not a clean negative and is excluded from accuracy denominators.",
        )
    )

    manifest = {
        "format_version": 1,
        "generator": "scripts/generate_m4_evaluation_corpus.py",
        "data_policy": "Synthetic workbook structures and values only. No customer or personal data.",
        "workbooks": [
            {"workbook_id": "detected-patterns", "filename": "detected-patterns.xlsx", "kind": "TARGET"},
            {"workbook_id": "mixed-candidates", "filename": "mixed-candidates.xlsx", "kind": "TARGET"},
            {"workbook_id": "normal-exceptions", "filename": "normal-exceptions.xlsx", "kind": "NORMAL"},
            {"workbook_id": "unsupported-syntax", "filename": "unsupported-syntax.xlsx", "kind": "UNSUPPORTED"},
            {
                "workbook_id": "scan-truncated",
                "filename": "scan-truncated.xlsx",
                "kind": "TRUNCATED",
                "scan_cell_limit": 2,
            },
        ],
        "cases": cases,
        "performance_workbooks": [
            {"workbook_id": "small", "filename": "performance-small.xlsx", "formula_cells": 1000},
            {"workbook_id": "medium", "filename": "performance-medium.xlsx", "formula_cells": 10000},
            {"workbook_id": "large", "filename": "performance-large.xlsx", "formula_cells": 30000},
        ],
    }
    (CORPUS_ROOT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def main() -> None:
    build_detected_patterns()
    build_mixed_candidates()
    build_normal_exceptions()
    build_unsupported()
    build_truncated()
    build_performance("performance-small.xlsx", 1000)
    build_performance("performance-medium.xlsx", 10000)
    build_performance("performance-large.xlsx", 30000)
    build_manifest()
    print(f"Generated synthetic M4-A.5 corpus in {CORPUS_ROOT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
