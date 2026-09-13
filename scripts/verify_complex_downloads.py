from pathlib import Path
import hashlib, json, sys, zipfile
from io import BytesIO
from lxml import html
from openpyxl import load_workbook

R = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(R / "scripts"))
import verify_complex_delivery as delivery

P = delivery.cases.PACK
O = R / "artifacts/verification/d08-complex"
D = O / "browser-downloads"
oracle = delivery.cases.oracle()
results = []
for label, case in oracle["repair"].items():
    directory = D / label
    files = list(directory.iterdir())
    assert len(files) == 3
    repaired = next(directory.glob("workbookcare_repaired_*.xlsx"))
    changes = next(directory.glob("changes_*.xlsx"))
    verification = next(directory.glob("verification_*.html"))
    delivery.check_repaired(repaired.read_bytes(), case)
    source = (P / case["file"]).read_bytes()
    with zipfile.ZipFile(BytesIO(source)) as a, zipfile.ZipFile(repaired) as b:
        assert set(a.namelist()) == set(b.namelist())
        allowed = {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
        changed = {n for n in a.namelist() if a.read(n) != b.read(n)}
        assert changed <= allowed, changed
        untouched = len(a.namelist()) - len(changed)
    book = load_workbook(changes, data_only=False)
    rows = list(book["변경 셀"].iter_rows(min_row=2, values_only=True))
    assert {r[1] for r in rows} == set(case["targets"])
    for row in rows:
        address = row[1]
        after = case["restored_formulas"].get(
            address, str(case["expected"][address]["value"])
        )
        assert row[5] == after, (address, "change-report after differs")
    book.close()
    tree = html.fromstring(verification.read_bytes())
    assert not tree.xpath("//script|//iframe|//form")
    assert tree.xpath('//*[@data-provenance="source_hash"]/text()') == [
        hashlib.sha256(source).hexdigest()
    ]
    assert tree.xpath('//*[@data-provenance="output_hash"]/text()') == [
        hashlib.sha256(repaired.read_bytes()).hexdigest()
    ]
    assert set(tree.xpath("//*[@data-change-cell]/@data-change-cell")) == set(
        case["targets"]
    )
    impacts = {k: v for k, v in case["impact"].items() if k not in case["targets"]}
    assert set(tree.xpath("//*[@data-result-cell]/@data-result-cell")) == set(impacts)
    for element in tree.xpath("//*[@data-result-cell]"):
        address = element.get("data-result-cell")
        assert element.xpath("./strong/text()") == [
            f"{case['expected'][address]['value']:,}"
        ]
    checks = tree.xpath("//*[@data-check]")
    assert len(checks) == 9 and all(
        e.xpath("./strong/text()") == ["PASS"] for e in checks
    )
    results.append(
        {
            "case": label,
            "status": "PASS",
            "cached_expected_values": len(case["expected"]),
            "change_rows": len(rows),
            "exact_html_impacts": len(impacts),
            "html_checks": 9,
            "unchanged_zip_members": untouched,
            "source_sha256": hashlib.sha256(source).hexdigest(),
        }
    )
for label in ["comparison", "comparison-capacity"]:
    directory = D / label
    files = list(directory.iterdir())
    assert len(files) == 2
    book = load_workbook(
        next(directory.glob("comparison_report_*.xlsx")), data_only=False
    )
    rows = [
        r
        for name in ["금액차이", "한쪽자료", "중복모호", "자료오류", "일치"]
        for r in book[name].iter_rows(min_row=2, values_only=True)
        if r[3]
    ]
    required = 888 if label == "comparison" else 2000
    assert len(rows) == required and len({r[3] for r in rows}) == required
    if label == "comparison":
        expected = oracle["comparison"]
        for side in ["A", "B"]:
            for source in expected["physical_rows"][side]:
                record = next(
                    r
                    for r in rows
                    if r[2] == side and r[5] == str(source["physical_row"])
                )
                assert record[1] == source["status"]
        assert {r[10] for r in rows if "BIG" in r[6]} == {"1"}
    for sheet in book:
        for row in sheet:
            assert all(c.data_type != "f" and c.hyperlink is None for c in row)
    book.close()
    tree = html.fromstring(
        next(directory.glob("comparison_verification_*.html")).read_bytes()
    )
    assert not tree.xpath("//script|//iframe|//form")
    elements = tree.xpath("//*[@data-source-row]")
    assert (
        len(elements) == required
        and len({e.get("data-source-row") for e in elements}) == required
    )
    labels = {
        "MATCHED": "일치",
        "AMOUNT_DIFF": "금액 차이",
        "ONLY_A": "A에만",
        "ONLY_B": "B에만",
        "AMBIGUOUS": "중복·모호",
        "INPUT_ERROR": "자료오류",
    }
    for side in ["A", "B"]:
        expected_total = (
            oracle["comparison"]["controls"][side]["known_amount"]
            if label == "comparison"
            else oracle["capacity"]["known_amount"]
        )
        assert tree.xpath(f'//*[@data-total="{side}"]/text()') == [expected_total]
        name = (
            oracle["comparison"]["files"][side]
            if label == "comparison"
            else oracle["capacity"]["supported_file"]
        )
        assert tree.xpath(f'//*[@data-info="source_{side}_hash"]/text()') == [
            hashlib.sha256((P / name).read_bytes()).hexdigest()
        ]
    for element in elements:
        side, physical = element.get("data-source-row").split(":")
        cells = [e.text_content() for e in element.xpath("./td")]
        if label == "comparison":
            expected_row = next(
                r
                for r in oracle["comparison"]["physical_rows"][side]
                if r["physical_row"] == int(physical)
            )
            assert cells[0] == labels[expected_row["status"]]
            if expected_row["key"] == ["정밀", "BIG"]:
                assert cells[-1] == "1"
        else:
            assert cells[0] == "일치" and cells[4] == str((int(physical) - 1) * 7)
    results.append(
        {
            "case": label,
            "status": "PASS",
            "all_source_rows": required,
            "report_is_literal_no_formulas_or_hyperlinks": True,
            "html_no_executable_content": True,
            "exact_html_source_classes_totals_and_hashes": True,
        }
    )
hashes = json.loads((P / "expected/source-hashes.json").read_text(encoding="utf-8"))
for name, digest in hashes.items():
    assert hashlib.sha256((P / name).read_bytes()).hexdigest() == digest
(O / "download-validation.json").write_text(
    json.dumps(
        {
            "kind": "ACTUAL_BETA_DOWNLOADED_FILES_VS_FROZEN_ORACLE",
            "status": "PASS",
            "results": results,
            "all_frozen_inputs_unchanged": True,
            "official_pg_verified": False,
        },
        ensure_ascii=False,
        indent=2,
    )
    + "\n",
    encoding="utf-8",
)
print(
    json.dumps({"status": "PASS", "downloaded_files": 13, "frozen_inputs": len(hashes)})
)
