from __future__ import annotations

import base64
import hashlib
import json
import os
import runpy
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


def _verify_packaged_manifest(directory: Path) -> None:
    manifest = json.loads((directory / "monthly-ux07-source-freeze.json").read_text(encoding="utf-8"))
    for item in manifest["overlay"]:
        path = Path("/") / item["image_path"]
        assert path.is_file(), item["image_path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"], item["image_path"]


def _rp03_verify(directory: Path) -> dict:
    _verify_packaged_manifest(directory)
    from app.config import Settings
    from app.delivery_execution import approve, compatibility_status, download, execute
    from app.delivery_inputs import PROFILE_3, inspect_input
    from app.delivery_plan import build_plan, plan_digest, reference_status
    from app.delivery_store import DeliveryStore
    from app.scanner import run_formula_audit

    source_path = directory / "monthly-rp03-supported-6sheet.xlsx"
    source = source_path.read_bytes()
    expected_source_sha = "c5e4589ca19ca892a005d92b71101aa1abbf7e1ffdb04d97b9b808cc1a617dab"
    assert hashlib.sha256(source).hexdigest() == expected_source_sha
    status = {"reference_status": reference_status(), "compatibility_status": compatibility_status()}
    assert status["reference_status"].get("status") == "PASS", status
    assert status["compatibility_status"].get("status") == "PASS", status
    settings = Settings(app_env="internal_beta")
    policy = {"profile": PROFILE_3, "sheet": "Budget", "targets": ["N18"], "before_formula": "=N15-N14", "confirmed": True}
    with tempfile.TemporaryDirectory(prefix="monthly-ux07-rp03-runtime-") as folder:
        store = DeliveryStore(Path(folder))
        snapshot = inspect_input("monthly-rp03-supported-6sheet.xlsx", source, settings)
        job = store.create("synthetic-runtime", source, snapshot, "monthly-ux07-rp03-runtime")
        plan = build_plan(job, policy)
        assert plan["profile_version"] == PROFILE_3
        assert plan["exact_targets"] == [["Budget", "N18"]]
        assert len(plan["patches"]) == 1
        assert plan["patches"][0]["after"]["value"] == "='M10'!B16-'M10'!B15"
        before = run_formula_audit("monthly-rp03-supported-6sheet.xlsx", source, settings)
        before_count = sum(1 for c in before.candidates if c.sheet == "Budget" and c.cell == "N18")
        assert before_count == 1
        job = store.update(job, job["revision"], {**job["state"], "policy": policy, "plan": plan, "status": "PREVIEW_VALIDATED", "internal_grant": {"kind": "INTERNAL_SYNTHETIC", "job_id": job["id"], "source_hash": job["snapshot"]["source_hash"], "expires_at": job["expires"]}})
        job = approve(store, job, {"revision": job["revision"], "plan_digest": plan_digest(plan), "candidate_ids": [p["candidate_id"] for p in plan["patches"]], "acknowledge_exact_changes": True}, settings)
        result = execute(store, job, settings)
        assert result["state"]["status"] == "READY"
        assert store.load("synthetic-runtime", job["id"])["source"] == source
        decoded = {}
        for kind in ["REPAIRED_XLSX", "CHANGES_XLSX", "VERIFICATION_HTML"]:
            artifact = download(store, result, kind, settings)
            data = base64.b64decode(artifact["file_base64"], validate=True)
            assert hashlib.sha256(data).hexdigest() == artifact["sha256"]
            decoded[kind] = {"bytes": len(data)}
            if kind == "REPAIRED_XLSX":
                repaired = data
        after = run_formula_audit("rp03-repaired.xlsx", repaired, settings)
        after_count = sum(1 for c in after.candidates if c.sheet == "Budget" and c.cell == "N18")
        assert after_count == 0
        formulas = load_workbook(BytesIO(repaired), data_only=False, read_only=True)
        values = load_workbook(BytesIO(repaired), data_only=True, read_only=True)
        try:
            assert formulas["Budget"]["N18"].value == "='M10'!B16-'M10'!B15"
            assert values["Budget"]["N18"].value == -5
        finally:
            formulas.close(); values.close()
    return {"event": "monthly_rp03_runtime_verified", "patch_count": 1, "expected_value_ok": True, "source_preserved": True, "monthly_detector_target_candidates": {"before": before_count, "after": after_count}, "decoded_artifact_count": len(decoded), "decoded_artifacts_nonempty": all(v["bytes"] > 0 for v in decoded.values()), "payment_tested": False}


def main() -> None:
    from app.delivery_runtime_check import verify
    verification = Path("/app/verification")
    try:
        print(json.dumps(verify(verification), separators=(",", ":")), flush=True)
        print(json.dumps(_rp03_verify(verification), separators=(",", ":")), flush=True)
    except Exception as error:
        print(json.dumps({"event": "monthly_ux07_startup_failed", "error_type": type(error).__name__, "safe_code": getattr(error, "code", "STARTUP_CHECK_FAILED")}, separators=(",", ":")), flush=True)
        raise SystemExit(1) from None
    os.environ["DELIVERY_RUNTIME_VERIFY"] = "false"
    runpy.run_module("app.runtime", run_name="__main__")


if __name__ == "__main__":
    main()
