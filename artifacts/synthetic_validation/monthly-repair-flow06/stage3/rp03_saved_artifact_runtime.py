from __future__ import annotations

import base64
import hashlib
import json
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.config import Settings
from app.delivery_execution import approve, compatibility_status, download, execute
from app.delivery_inputs import PROFILE_3, inspect_input
from app.delivery_plan import build_plan, plan_digest, reference_status
from app.delivery_store import DeliveryStore
from app.scanner import run_formula_audit

repo = Path(__file__).resolve().parents[4]
out_dir = Path(__file__).with_name("saved-artifacts03")
if out_dir.exists() and any(out_dir.iterdir()):
    raise SystemExit("saved-artifacts03 is not empty; choose a fresh isolated output directory")
out_dir.mkdir(parents=True, exist_ok=True)
source_path = Path(__file__).with_name("monthly-rp03-supported-6sheet.xlsx")
source = source_path.read_bytes()
source_sha256 = hashlib.sha256(source).hexdigest()
assert source_sha256 == "c5e4589ca19ca892a005d92b71101aa1abbf7e1ffdb04d97b9b808cc1a617dab"
settings = Settings(app_env="internal_beta")
status = {"reference_status": reference_status(), "compatibility_status": compatibility_status()}
assert status["reference_status"].get("status") == "PASS", status
assert status["compatibility_status"].get("status") == "PASS", status
policy = {
    "profile": PROFILE_3,
    "sheet": "Budget",
    "targets": ["N18"],
    "before_formula": "=N15-N14",
    "confirmed": True,
}

with tempfile.TemporaryDirectory(prefix="monthly-rp03-saved-artifact-runtime-") as folder:
    store = DeliveryStore(Path(folder))
    snapshot = inspect_input("monthly-rp03-supported-6sheet.xlsx", source, settings)
    assert "UNSUPPORTED_FORMULA" in snapshot["issues"]
    job = store.create("synthetic-runtime", source, snapshot, "rp03-saved-artifact-runtime")
    plan = build_plan(job, policy)
    assert plan["profile_version"] == PROFILE_3
    assert plan["calculation_mode"] == "monthly_sheet_internal"
    assert plan["exact_targets"] == [["Budget", "N18"]]
    assert len(plan["patches"]) == 1
    patch = plan["patches"][0]
    assert patch["sheet"] == "Budget" and patch["cell"] == "N18"
    assert patch["before"] == {"type": "formula", "value": "=N15-N14", "style": "0"}
    assert patch["after"] == {
        "type": "formula",
        "value": "='M10'!B16-'M10'!B15",
        "style": "0",
    }
    assert plan["expected_calculated_values"]["Budget"]["N18"] == {
        "type": "number",
        "value": -5.0,
        "provenance": "ENGINE_CALCULATED",
    }
    before_audit = run_formula_audit("monthly-rp03-supported-6sheet.xlsx", source, settings)
    before_candidates = [
        (candidate.sheet, candidate.cell, candidate.rule_code)
        for candidate in before_audit.candidates
        if candidate.sheet == "Budget" and candidate.cell == "N18"
    ]
    assert len(before_candidates) == 1, before_candidates
    job = store.update(
        job,
        job["revision"],
        {
            **job["state"],
            "policy": policy,
            "plan": plan,
            "status": "PREVIEW_VALIDATED",
            "internal_grant": {
                "kind": "INTERNAL_SYNTHETIC",
                "job_id": job["id"],
                "source_hash": job["snapshot"]["source_hash"],
                "expires_at": job["expires"],
            },
        },
    )
    job = approve(
        store,
        job,
        {
            "revision": job["revision"],
            "plan_digest": plan_digest(plan),
            "candidate_ids": [patch["candidate_id"] for patch in plan["patches"]],
            "acknowledge_exact_changes": True,
        },
        settings,
    )
    result = execute(store, job, settings)
    assert result["state"]["status"] == "READY"
    assert store.load("synthetic-runtime", job["id"])["source"] == source
    artifacts = {}
    for kind, suffix in [
        ("REPAIRED_XLSX", "repaired.xlsx"),
        ("CHANGES_XLSX", "changes.xlsx"),
        ("VERIFICATION_HTML", "verification.html"),
    ]:
        artifact = download(store, result, kind, settings)
        data = base64.b64decode(artifact["file_base64"], validate=True)
        path = out_dir / f"rp03-{suffix}"
        path.write_bytes(data)
        assert hashlib.sha256(data).hexdigest() == artifact["sha256"]
        artifacts[kind] = {
            "path": str(path.relative_to(repo)),
            "filename": artifact["filename"],
            "mime": artifact["mime"],
            "sha256": artifact["sha256"],
            "bytes": len(data),
        }
    repaired = (out_dir / "rp03-repaired.xlsx").read_bytes()
    after_audit = run_formula_audit("rp03-repaired.xlsx", repaired, settings)
    after_candidates = [
        (candidate.sheet, candidate.cell, candidate.rule_code)
        for candidate in after_audit.candidates
        if candidate.sheet == "Budget" and candidate.cell == "N18"
    ]
    assert after_candidates == []
    book = load_workbook(BytesIO(repaired), data_only=False, read_only=False)
    try:
        assert book["Budget"]["N18"].value == "='M10'!B16-'M10'!B15"
        assert book["M10"]["B15"].value == 1006
        assert book["M10"]["B16"].value == 1001
    finally:
        book.close()
    values_book = load_workbook(BytesIO(repaired), data_only=True, read_only=True)
    try:
        assert values_book["Budget"]["N18"].value == -5
    finally:
        values_book.close()

(plan_path := out_dir / "rp03-plan.json").write_text(
    json.dumps(plan, ensure_ascii=False, indent=2) + "\n",
    encoding="utf-8",
)
evidence = {
    "status": "PASS",
    "kind": "UNMOCKED_RP03_SAVED_ARTIFACT_RUNTIME_NO_REFERENCE_LOOP",
    "reference_and_compatibility": status,
    "source_path": str(source_path.relative_to(repo)),
    "source_sha256": source_sha256,
    "source_sha256_preserved": hashlib.sha256(source).hexdigest(),
    "profile": PROFILE_3,
    "patch_count": 1,
    "exact_targets": [["Budget", "N18"]],
    "expected_cell": "Budget!N18",
    "expected_value": -5,
    "downloaded_data_only_value": -5,
    "raw_formula": "='M10'!B16-'M10'!B15",
    "monthly_detector_target_candidates": {"before": 1, "after": 0},
    "delivery_status": result["state"]["status"],
    "artifacts": artifacts,
    "plan_path": str(plan_path.relative_to(repo)),
    "mocked": False,
    "note": "Skips the reference recalculation loop; requires registered real reference and compatibility PASS before execution.",
}
out_path = Path(__file__).with_name("rp03-saved-artifact-runtime-result.json")
out_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(evidence, ensure_ascii=False, sort_keys=True))
