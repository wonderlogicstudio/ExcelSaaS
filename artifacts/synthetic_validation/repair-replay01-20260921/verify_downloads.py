from pathlib import Path
import hashlib, json, sys, zipfile
from io import BytesIO
from lxml import html
from openpyxl import load_workbook

R = Path('C:\\Users\\JinwonLee\\project\\ExcelSaaS')
sys.path.insert(0, str(R / "scripts"))
import verify_complex_delivery as delivery

P = delivery.cases.PACK
O = Path(__file__).parent
D = O / "browser-downloads"
oracle = {"repair": {"COMBINED": json.loads((O / "combined-oracle.json").read_text(encoding="utf-8"))}}
assert hashlib.sha256((O/"combined-oracle.json").read_bytes()).hexdigest() == (O/"combined-oracle.sha256").read_text().strip()
results = []
for label in ["COMBINED"]:
    case = oracle["repair"][label]
    directory = D / (label + "_FINAL")
    files = list(directory.iterdir())
    assert len(files) == 3
    repaired = next(directory.glob("workbookcare_repaired_*.xlsx"))
    changes = next(directory.glob("changes_*.xlsx"))
    verification = next(directory.glob("verification_*.html"))
    delivery.check_repaired(repaired.read_bytes(), case)
    source = (P / case["file"]).read_bytes()
    with zipfile.ZipFile(BytesIO(source)) as a, zipfile.ZipFile(repaired) as b:
        assert set(a.namelist()) == set(b.namelist())
        allowed = {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
        changed = {n for n in a.namelist() if a.read(n) != b.read(n)}
        assert changed <= allowed, changed
        untouched = len(a.namelist()) - len(changed)
    book = load_workbook(changes, data_only=False)
    rows = list(book["변경 셀"].iter_rows(min_row=2, values_only=True))
    assert {r[1] for r in rows} == set(case["targets"])
    for row in rows:
        address = row[1]
        after = case["restored_formulas"].get(
            address, str(case["expected"][address]["value"])
        )
        assert row[5] == after, (address, "change-report after differs")
    book.close()
    tree = html.fromstring(verification.read_bytes())
    assert not tree.xpath("//script|//iframe|//form")
    assert tree.xpath('//*[@data-provenance="source_hash"]/text()') == [
        hashlib.sha256(source).hexdigest()
    ]
    assert tree.xpath('//*[@data-provenance="output_hash"]/text()') == [
        hashlib.sha256(repaired.read_bytes()).hexdigest()
    ]
    assert set(tree.xpath("//*[@data-change-cell]/@data-change-cell")) == set(
        case["targets"]
    )
    impacts = {k: v for k, v in case["impact"].items() if k not in case["targets"]}
    assert set(tree.xpath("//*[@data-result-cell]/@data-result-cell")) == set(impacts)
    for element in tree.xpath("//*[@data-result-cell]"):
        address = element.get("data-result-cell")
        assert element.xpath("./strong/text()") == [
            f"{case['expected'][address]['value']:,}"
        ]
    checks = tree.xpath("//*[@data-check]")
    codes = set(tree.xpath("//*[@data-check]/@data-check"))
    assert "FORMULA_AUDIT_COMPLETED" in codes
    assert "NO_NEW_FORMULA_CANDIDATES" in codes
    assert ("TARGET_STATIC_FINDINGS_RESOLVED" in codes) == True
    assert ("TARGET_FORMULA_CANDIDATES_RESOLVED" in codes) == True
    assert len(checks) >= 9 and all(
        e.xpath("./strong/text()") == ["PASS"] for e in checks
    )
    results.append(
        {
            "case": label,
            "status": "PASS",
            "cached_expected_values": len(case["expected"]),
            "change_rows": len(rows),
            "exact_html_impacts": len(impacts),
            "html_checks": len(checks),
            "unchanged_zip_members": untouched,
            "source_sha256": hashlib.sha256(source).hexdigest(),
        }
    )
hashes = json.loads((P / "expected/source-hashes.json").read_text(encoding="utf-8"))
for name, digest in hashes.items():
    assert hashlib.sha256((P / name).read_bytes()).hexdigest() == digest
(O / "download-validation.json").write_text(
    json.dumps(
        {
            "kind": "ACTUAL_BETA_DOWNLOADED_FILES_VS_FROZEN_ORACLE",
            "status": "PASS",
            "results": results,
            "all_frozen_inputs_unchanged": True,
            "official_pg_verified": False,
        },
        ensure_ascii=False,
        indent=2,
    )
    + "\n",
    encoding="utf-8",
)
print(
    json.dumps({"status": "PASS", "downloaded_files": 3, "frozen_inputs": len(hashes)})
)
