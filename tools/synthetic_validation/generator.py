from __future__ import annotations

import random
from copy import copy
from dataclasses import asdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .scenarios import (
    BOM_CATALOG_SHEET,
    BUDGET_MONTH_SHEETS,
    FAMILY_ORDER,
    INVENTORY_SUPPORT_SHEET,
    INVENTORY_TRANSACTIONS_SHEET,
    SCENARIOS,
    Scenario,
)
from .schema import (
    DATASET_VERSION,
    GENERATOR_VERSION,
    DatasetManifest,
    ExpectedFinding,
    FormulaCellSpec,
    MutationSnapshot,
    NormalRegion,
    WorkbookTruth,
    dataclass_to_dict,
    sha256_file,
    sha256_json,
    write_json,
)

MUTATION_TYPES = (
    "m4_relative_reference_drift",
    "m4_formula_gap",
    "m4_constant_override",
    "base_broken_ref_token",
    "base_visible_error_token",
)

M4_SUPPORTED_SCENARIOS = {
    "retail_sales_vertical",
    "inventory_cross_sheet",
    "payroll_overtime",
    "project_profit",
    "receivables_aging",
    "manufacturing_bom",
    "cashflow_calendar",
    "subscriptions_kpi",
}

PAYROLL_THRESHOLD_CELL = "C5"
PAYROLL_PREMIUM_CELL = "B5"
PAYROLL_OVERTIME_THRESHOLD = 160
PAYROLL_OVERTIME_PREMIUM = 0.5
INVENTORY_TRANSACTION_RANGE_END = 100
BOM_CATALOG_RANGE_END = 20


def _split_for_family(index: int) -> str:
    if index < 6:
        return "train"
    if index < 8:
        return "dev"
    return "holdout"


def _business_value(row: int, col: int, seed: int) -> int:
    return ((row * 17 + col * 11 + seed) % 91) + 9


def _m4_supported(scenario: Scenario) -> bool:
    return scenario.scenario_id in M4_SUPPORTED_SCENARIOS


def _inventory_transaction_total(wb: Any, sku: str, movement: str) -> float | int:
    transactions = wb[INVENTORY_TRANSACTIONS_SHEET]
    total = 0
    for row in range(2, INVENTORY_TRANSACTION_RANGE_END + 1):
        if transactions[f"A{row}"].value == sku and transactions[f"C{row}"].value == movement:
            total += transactions[f"D{row}"].value or 0
    return total


def _bom_catalog_unit_cost(wb: Any, part_id: str) -> float | int:
    catalog = wb[BOM_CATALOG_SHEET]
    for row in range(2, BOM_CATALOG_RANGE_END + 1):
        if catalog[f"A{row}"].value == part_id:
            return catalog[f"B{row}"].value
    raise ValueError(f"Part {part_id!r} missing from {BOM_CATALOG_SHEET}")


def _semantic_value(scenario: Scenario, row: int, ws: Any, prior: dict[int, float | int]) -> float | int:
    sid = scenario.scenario_id
    if sid == "retail_sales_vertical":
        return round(ws[f"E{row}"].value * ws[f"F{row}"].value * (1 - ws[f"H{row}"].value), 2)
    if sid == "purchase_orders_table":
        return ws[f"F{row}"].value * ws[f"G{row}"].value if ws[f"E{row}"].value > 0 else 0
    if sid == "inventory_cross_sheet":
        opening = ws.parent[INVENTORY_SUPPORT_SHEET][f"B{row}"].value
        if INVENTORY_TRANSACTIONS_SHEET in ws.parent.sheetnames:
            sku = ws[f"A{row}"].value
            receipts = _inventory_transaction_total(ws.parent, sku, "IN")
            issues = _inventory_transaction_total(ws.parent, sku, "OUT")
            return opening + receipts - issues
        return opening + ws[f"G{row}"].value - ws[f"H{row}"].value
    if sid == "budget_horizontal_months":
        return ws[f"L{row}"].value - ws[f"M{row}"].value
    if sid == "payroll_overtime":
        hours = ws[f"H{row}"].value
        rate = ws.parent["Rate Card"][f"B{row}"].value
        threshold = ws[PAYROLL_THRESHOLD_CELL].value
        premium = ws[PAYROLL_PREMIUM_CELL].value
        return hours * rate + max(hours - threshold, 0) * rate * premium
    if sid == "project_profit":
        uplift = 0.18 if row in scenario.normal_exception_rows else 0.12
        return round(ws[f"I{row}"].value - ws[f"J{row}"].value + ws[f"H{row}"].value * uplift, 2)
    if sid == "receivables_aging":
        amount = ws[f"F{row}"].value
        days = ws[f"G{row}"].value
        return round(amount * 0.97, 2) if days > 30 else amount
    if sid == "manufacturing_bom":
        unit_cost = _bom_catalog_unit_cost(ws.parent, ws[f"A{row}"].value) if BOM_CATALOG_SHEET in ws.parent.sheetnames else ws[f"G{row}"].value
        return round((unit_cost * ws[f"H{row}"].value) / (1 - ws[f"J{row}"].value), 2)
    if sid == "cashflow_calendar":
        net = ws[f"H{row}"].value - ws[f"I{row}"].value
        return net if row == scenario.start_row else prior[row - 1] + net
    if sid == "subscriptions_kpi":
        active = ws[f"J{row}"].value
        return round(ws[f"K{row}"].value / active, 4) if active else 0
    raise ValueError(scenario.scenario_id)


def _normal_exception_reason(scenario: Scenario) -> str:
    reasons = {
        "retail_sales_vertical": "Approved campaign discount line uses a documented higher discount rate.",
        "purchase_orders_table": "Blanket order line is reviewed by purchasing because minimum quantity policy applies.",
        "inventory_cross_sheet": "Opening stock reconciliation line is reviewed against the warehouse count sheet.",
        "budget_horizontal_months": "Forecast month includes a documented management override.",
        "payroll_overtime": "Union overtime rule requires payroll confirmation for this employee row.",
        "project_profit": "Approved change order uses a higher margin uplift than standard work.",
        "receivables_aging": "Customer dispute reserve is documented for manual credit review.",
        "manufacturing_bom": "Pilot batch yield loss is documented for production review.",
        "cashflow_calendar": "First collection date uses treasury timing judgment and needs user confirmation.",
        "subscriptions_kpi": "Enterprise cohort uses a documented retention adjustment.",
    }
    return reasons[scenario.scenario_id]


def _headers_for_scenario(scenario: Scenario) -> list[str]:
    headers_by_id = {
        "retail_sales_vertical": ["Line", "Channel", "Category", "Policy note", "Units", "Unit price", "Net sales", "Discount rate", "Tax code", "Region", "Campaign", "Reviewer"],
        "purchase_orders_table": ["Line", "Buyer", "Vendor", "Policy note", "Minimum qty", "Ordered qty", "Unit cost", "Line cost", "Lead days", "PO type", "Requester", "Reviewer"],
        "inventory_cross_sheet": ["SKU", "Warehouse", "Category", "Policy note", "Cycle", "Batch", "Receipts", "Issues", "Closing qty", "Planner", "Status", "Reviewer"],
        "budget_horizontal_months": ["Line", "Owner", "Category", "Policy note", "M01", "M02", "M03", "M04", "M05", "M06", "M07", "M08"],
        "payroll_overtime": ["Employee", "Team", "Category", "Policy note", "Period", "Shift", "Base code", "Hours", "Rate ref", "Gross pay", "Cost center", "Reviewer"],
        "project_profit": ["Line", "PM", "Category", "Policy note", "Phase", "Contract", "Risk", "Revenue", "Billings", "Cost", "Profit", "Reviewer"],
        "receivables_aging": ["Invoice", "Owner", "Customer", "Policy note", "Terms", "Amount", "Days past due", "Collectible amount", "Bucket", "Region", "Collector", "Reviewer"],
        "manufacturing_bom": ["Part", "Cell", "Category", "Policy note", "Batch", "Run", "Unit cost", "Quantity", "Extended cost", "Yield loss", "Station", "Reviewer"],
        "cashflow_calendar": ["Date", "Owner", "Category", "Policy note", "Scenario", "Bank", "Timing", "Inflows", "Outflows", "Cash balance", "Source", "Reviewer"],
        "subscriptions_kpi": ["Cohort", "Owner", "Segment", "Policy note", "Plan", "Region", "Channel", "Seats", "Status", "Active users", "Retained users", "Retention rate"],
    }
    return headers_by_id[scenario.scenario_id]


def _fill_inputs(ws: Any, scenario: Scenario, seed: int) -> None:
    for row in range(scenario.start_row, scenario.start_row + scenario.row_count):
        for col_index in range(5, 13):
            ws[f"{get_column_letter(col_index)}{row}"] = _business_value(row, col_index, seed)
        ws[f"B{row}"] = scenario.business_domain
        ws[f"C{row}"] = "standard"
        ws[f"D{row}"] = _normal_exception_reason(scenario) if row in scenario.normal_exception_rows else ""
        if scenario.scenario_id == "retail_sales_vertical":
            ws[f"H{row}"] = 0.05 if row not in scenario.normal_exception_rows else 0.1
        elif scenario.scenario_id == "receivables_aging":
            ws[f"G{row}"] = 15 + (row % 5) * 10
        elif scenario.scenario_id == "manufacturing_bom":
            ws[f"J{row}"] = 0.03 + (row % 3) * 0.02
        elif scenario.scenario_id == "subscriptions_kpi":
            ws[f"J{row}"] = 100 + row
            ws[f"K{row}"] = 80 + (row % 9)
    for row in range(scenario.start_row, scenario.start_row + scenario.row_count):
        if scenario.sheet_mode == "dates":
            ws[f"A{row}"] = datetime(2026, 10, 1) + timedelta(days=row - scenario.start_row)
        elif scenario.scenario_id == "inventory_cross_sheet":
            ws[f"A{row}"] = f"SKU-{row - scenario.start_row + 1}"
        elif scenario.scenario_id == "manufacturing_bom":
            ws[f"A{row}"] = f"PART-{row - scenario.start_row + 1:03d}"
        else:
            ws[f"A{row}"] = f"Line {row - scenario.start_row + 1}"


def _normal_formula(scenario: Scenario, row: int, *, diversity_features: bool = False) -> str:
    sid = scenario.scenario_id
    if sid == "retail_sales_vertical":
        return f"=E{row}*F{row}*(1-H{row})"
    if sid == "purchase_orders_table":
        return f"=IF(E{row}>0,F{row}*G{row},0)"
    if sid == "inventory_cross_sheet":
        if diversity_features:
            qty_range = f"'{INVENTORY_TRANSACTIONS_SHEET}'!$D$2:$D${INVENTORY_TRANSACTION_RANGE_END}"
            sku_range = f"'{INVENTORY_TRANSACTIONS_SHEET}'!$A$2:$A${INVENTORY_TRANSACTION_RANGE_END}"
            move_range = f"'{INVENTORY_TRANSACTIONS_SHEET}'!$C$2:$C${INVENTORY_TRANSACTION_RANGE_END}"
            return f"='{INVENTORY_SUPPORT_SHEET}'!B{row}+SUMIFS({qty_range},{sku_range},A{row},{move_range},\"IN\")-SUMIFS({qty_range},{sku_range},A{row},{move_range},\"OUT\")"
        return f"='{INVENTORY_SUPPORT_SHEET}'!B{row}+G{row}-H{row}"
    if sid == "budget_horizontal_months":
        return f"=L{row}-M{row}"
    if sid == "payroll_overtime":
        return f"=H{row}*'Rate Card'!$B{row}+IF(H{row}>C$5,(H{row}-C$5)*'Rate Card'!$B{row}*$B$5,0)"
    if sid == "project_profit":
        return f"=I{row}-J{row}+H{row}*0.12"
    if sid == "receivables_aging":
        return f"=IF(G{row}>30,F{row}*0.97,F{row})"
    if sid == "manufacturing_bom":
        if diversity_features:
            return f"=ROUND(VLOOKUP(A{row},'{BOM_CATALOG_SHEET}'!$A$2:$C${BOM_CATALOG_RANGE_END},2,FALSE)*H{row}/(1-J{row}),2)"
        return f"=ROUND(G{row}*H{row}/(1-J{row}),2)"
    if sid == "cashflow_calendar":
        return f"=H{row}-I{row}" if row == scenario.start_row else f"=J{row - 1}+H{row}-I{row}"
    if sid == "subscriptions_kpi":
        return f"=IF(J{row}>0,ROUND(K{row}/J{row},4),0)"
    return scenario.formula_template.format(row=row)


def _exception_formula(scenario: Scenario, row: int, *, diversity_features: bool = False) -> str:
    if scenario.scenario_id == "project_profit":
        return f"=I{row}-J{row}+H{row}*0.18"
    return _normal_formula(scenario, row, diversity_features=diversity_features)


def _style_workbook(ws: Any, scenario: Scenario) -> None:
    ws["A1"] = scenario.business_domain.title()
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Synthetic validation file - not customer data"
    if scenario.scenario_id == "payroll_overtime":
        ws["A5"] = "Overtime premium"
        ws[PAYROLL_PREMIUM_CELL] = PAYROLL_OVERTIME_PREMIUM
        ws["B5"].number_format = "0.0%"
        ws["C4"] = "Overtime threshold"
        ws[PAYROLL_THRESHOLD_CELL] = PAYROLL_OVERTIME_THRESHOLD
    for offset, header in enumerate(_headers_for_scenario(scenario), start=1):
        cell = ws.cell(row=scenario.start_row - 1, column=offset)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7")
    ws.freeze_panes = f"A{scenario.start_row}"


def _add_support_sheets(wb: Workbook, scenario: Scenario, seed: int, *, diversity_features: bool = False) -> None:
    if scenario.hidden_support_sheet:
        rates = wb.create_sheet("Rate Card")
        rates.sheet_state = "hidden"
        rates["A1"] = "Employee row"
        rates["B1"] = "Hourly rate"
        for row in range(scenario.start_row, scenario.start_row + scenario.row_count):
            rates[f"A{row}"] = row
            rates[f"B{row}"] = 20 + (row + seed) % 11
    if scenario.sheet_mode == "cross_sheet":
        lookup = wb.create_sheet(INVENTORY_SUPPORT_SHEET)
        lookup["A1"] = "SKU row"
        lookup["B1"] = "Opening quantity"
        for row in range(scenario.start_row, scenario.start_row + scenario.row_count):
            lookup[f"A{row}"] = wb[scenario.base_sheet][f"A{row}"].value
            lookup[f"B{row}"] = 100 + row + seed % 17
        if diversity_features:
            transactions = wb.create_sheet(INVENTORY_TRANSACTIONS_SHEET)
            transactions["A1"] = "SKU"
            transactions["B1"] = "Warehouse"
            transactions["C1"] = "Movement"
            transactions["D1"] = "Quantity"
            txn_row = 2
            detail = wb[scenario.base_sheet]
            for row in range(scenario.start_row, scenario.start_row + scenario.row_count):
                sku = detail[f"A{row}"].value
                receipts = int(detail[f"G{row}"].value)
                issues = int(detail[f"H{row}"].value)
                split_receipts = max(receipts - (row % 4 + 1), 0)
                for movement, quantity in (
                    ("IN", split_receipts),
                    ("IN", receipts - split_receipts),
                    ("OUT", issues),
                ):
                    transactions[f"A{txn_row}"] = sku
                    transactions[f"B{txn_row}"] = detail[f"B{row}"].value
                    transactions[f"C{txn_row}"] = movement
                    transactions[f"D{txn_row}"] = quantity
                    txn_row += 1
    if scenario.sheet_mode == "bom" and diversity_features:
        catalog = wb.create_sheet(BOM_CATALOG_SHEET)
        catalog["A1"] = "Part"
        catalog["B1"] = "Unit cost"
        catalog["C1"] = "Category"
        detail = wb[scenario.base_sheet]
        for offset, row in enumerate(range(scenario.start_row, scenario.start_row + scenario.row_count), start=2):
            catalog[f"A{offset}"] = detail[f"A{row}"].value
            catalog[f"B{offset}"] = detail[f"G{row}"].value
            catalog[f"C{offset}"] = detail[f"C{row}"].value


def _expected_normal_exceptions(workbook_id: str, scenario: Scenario, formula_cells: list[FormulaCellSpec]) -> list[ExpectedFinding]:
    rows: list[ExpectedFinding] = []
    for spec in formula_cells:
        if spec.role != "normal_exception":
            continue
        rows.append(
            ExpectedFinding(
                expected_id=f"{workbook_id}-normal-exception-{spec.cell}",
                kind="normal_exception" if spec.supported_by_current_m4 else "user_confirmation",
                sheet=spec.sheet,
                cell=spec.cell,
                rule_code=None,
                family_id=scenario.family_id,
                notes=_normal_exception_reason(scenario),
            )
        )
    return rows


def create_normal_workbook(*, scenario: Scenario, seed: int, workbook_id: str, split: str, output_dir: Path, diversity_features: bool = False) -> WorkbookTruth:
    wb = Workbook()
    ws = wb.active
    ws.title = scenario.base_sheet
    _style_workbook(ws, scenario)
    _fill_inputs(ws, scenario, seed)
    _add_support_sheets(wb, scenario, seed, diversity_features=diversity_features)

    formula_cells: list[FormulaCellSpec] = []
    if scenario.sheet_mode == "horizontal":
        ws["A15"] = "Budget"
        ws["A16"] = "Actual"
        ws["A18"] = "Variance"
        if diversity_features:
            for month_index, month_sheet in enumerate(BUDGET_MONTH_SHEETS, start=1):
                month_ws = wb.create_sheet(month_sheet)
                month_ws["A1"] = f"Budget actuals {month_sheet}"
                month_ws["A15"] = "Budget"
                month_ws["A16"] = "Actual"
                month_ws["B15"] = 900 + seed % 50 + month_index * 10
                month_ws["B16"] = 880 + seed % 45 + month_index * 12
        for month_index, col_index in enumerate(range(5, 17), start=1):
            col = get_column_letter(col_index)
            ws[f"{col}14"] = f"M{month_index:02d}"
            if diversity_features:
                month_sheet = BUDGET_MONTH_SHEETS[month_index - 1]
                ws[f"{col}15"] = f"='{month_sheet}'!B15"
                ws[f"{col}16"] = f"='{month_sheet}'!B16"
                formula = f"='{month_sheet}'!B16-'{month_sheet}'!B15"
                semantic = wb[month_sheet]["B16"].value - wb[month_sheet]["B15"].value
            else:
                ws[f"{col}15"] = 900 + seed % 50 + month_index * 10
                ws[f"{col}16"] = 880 + seed % 45 + month_index * 12
                formula = f"={col}16-{col}15"
                semantic = ws[f"{col}16"].value - ws[f"{col}15"].value
            ws[f"{col}18"] = formula
            formula_cells.append(FormulaCellSpec(scenario.base_sheet, f"{col}18", formula, semantic, "exploratory", False))
        filename = f"{workbook_id}.xlsx"
        path = output_dir / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        wb.close()
        return WorkbookTruth(
            workbook_id=workbook_id,
            parent_id=None,
            family_id=scenario.family_id,
            scenario_id=scenario.scenario_id,
            business_domain=scenario.business_domain,
            structural_family=scenario.structural_family,
            split=split,  # type: ignore[arg-type]
            kind="normal",
            seed=seed,
            filename=filename,
            relative_path=f"workbooks/{filename}",
            sha256=sha256_file(path),
            mutation_type=None,
            mutation_target=None,
            supported_scope=list(scenario.supported_scope),
            exploratory_scope=list(scenario.exploratory_scope),
            normal_regions=[NormalRegion(scenario.base_sheet, "E18:P18", "Horizontal repetition is generated as exploratory and is not M4-scored.")],
            formula_cells=formula_cells,
            expected_findings=[],
            generation_checks={"formula_string_checks": len(formula_cells), "semantic_arithmetic_checks": len(formula_cells), "answer_marker_cells": 0, "excel_recalculation": "NOT_RUN"},
        )

    supported = _m4_supported(scenario)
    normal_regions = [
        NormalRegion(
            sheet=scenario.base_sheet,
            range_ref=f"{scenario.formula_column}{scenario.start_row}:{scenario.formula_column}{scenario.start_row + scenario.row_count - 1}",
            notes="Detail formulas in this region are checked for unexpected findings when supported by the current engine.",
        )
    ]
    semantic_by_row: dict[int, float | int] = {}
    for row in range(scenario.start_row, scenario.start_row + scenario.row_count):
        cell = ws[f"{scenario.formula_column}{row}"]
        if row in scenario.normal_exception_rows:
            formula = _exception_formula(scenario, row, diversity_features=diversity_features)
            role = "normal_exception"
        else:
            formula = _normal_formula(scenario, row, diversity_features=diversity_features)
            role = "detail"
        semantic = _semantic_value(scenario, row, ws, semantic_by_row)
        semantic_by_row[row] = semantic
        cell.value = formula
        formula_cells.append(FormulaCellSpec(scenario.base_sheet, cell.coordinate, formula, semantic, role, supported))

    for row in scenario.summary_rows:
        label = ws[f"A{row}"]
        label.value = "Subtotal"
        target = ws[f"{scenario.formula_column}{row}"]
        start = scenario.start_row
        end = scenario.start_row + scenario.row_count - 1
        target.value = f"={scenario.formula_column}{end}" if scenario.scenario_id == "cashflow_calendar" else f"=SUM({scenario.formula_column}{start}:{scenario.formula_column}{end})"
        formula_cells.append(FormulaCellSpec(scenario.base_sheet, target.coordinate, target.value, None, "subtotal", False))

    if scenario.table:
        ref = f"A{scenario.start_row - 1}:H{scenario.start_row + scenario.row_count - 1}"
        table = Table(displayName=f"T_{workbook_id.replace('-', '_')}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)

    filename = f"{workbook_id}.xlsx"
    path = output_dir / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()

    return WorkbookTruth(
        workbook_id=workbook_id,
        parent_id=None,
        family_id=scenario.family_id,
        scenario_id=scenario.scenario_id,
        business_domain=scenario.business_domain,
        structural_family=scenario.structural_family,
        split=split,  # type: ignore[arg-type]
        kind="normal",
        seed=seed,
        filename=filename,
        relative_path=f"workbooks/{filename}",
        sha256=sha256_file(path),
        mutation_type=None,
        mutation_target=None,
        supported_scope=list(scenario.supported_scope),
        exploratory_scope=list(scenario.exploratory_scope),
        normal_regions=normal_regions,
        formula_cells=formula_cells,
        expected_findings=_expected_normal_exceptions(workbook_id, scenario, formula_cells),
        generation_checks={"formula_string_checks": len(formula_cells), "semantic_arithmetic_checks": len([spec for spec in formula_cells if spec.semantic_value is not None]), "answer_marker_cells": 0, "excel_recalculation": "NOT_RUN"},
    )


def _copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format


def _eligible_detail_rows(scenario: Scenario, supported: bool, mutation_type: str) -> list[int]:
    rows = [row for row in range(scenario.start_row + 1, scenario.start_row + scenario.row_count - 1) if row not in scenario.normal_exception_rows]
    if mutation_type.startswith("m4_") and supported:
        eligible_by_scenario = {
            "retail_sales_vertical": (13, 14, 18, 19),
            "inventory_cross_sheet": (11, 12, 16, 17),
            "payroll_overtime": (12, 13, 17, 18),
            "project_profit": (14, 15),
            "receivables_aging": (15, 16, 17),
            "manufacturing_bom": (10, 11, 12, 16, 17),
            "cashflow_calendar": (18, 19, 20, 21, 22),
            "subscriptions_kpi": (12, 13, 17, 18),
        }
        return [row for row in eligible_by_scenario[scenario.scenario_id] if row in rows]
    return rows


def _semantic_for_cell(truth: WorkbookTruth, sheet: str, cell: str) -> float | int | str | None:
    for spec in truth.formula_cells:
        if spec.sheet == sheet and spec.cell == cell:
            return spec.semantic_value
    return None


def _workbook_snapshot(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    try:
        cells: dict[str, tuple[Any, str, int, str]] = {}
        sheets: dict[str, Any] = {}
        for ws in wb.worksheets:
            sheets[ws.title] = {
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else "",
                "merged_ranges": sorted(str(item) for item in ws.merged_cells.ranges),
                "tables": sorted(f"{name}:{getattr(ws.tables[name], 'ref', ws.tables[name])}" for name in ws.tables),
                "row_heights": {str(key): value.height for key, value in ws.row_dimensions.items() if value.height is not None},
                "column_widths": {key: value.width for key, value in ws.column_dimensions.items() if value.width is not None},
            }
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[f"{ws.title}!{cell.coordinate}"] = (cell.value, cell.data_type, cell.style_id, cell.number_format)
        defined_names = sorted(str(name) for name in wb.defined_names)
        return {"cells": cells, "sheets": sheets, "defined_names": defined_names}
    finally:
        wb.close()


def _workbook_values(path: Path) -> dict[tuple[str, str], tuple[Any, str]]:
    wb = load_workbook(path, data_only=False)
    values: dict[tuple[str, str], tuple[Any, str]] = {}
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        values[(ws.title, cell.coordinate)] = (cell.value, cell.data_type)
    finally:
        wb.close()
    return values


def create_mutated_workbook(*, normal_truth: WorkbookTruth, mutation_index: int, seed: int, output_dir: Path) -> WorkbookTruth:
    scenario = next(item for item in SCENARIOS if item.family_id == normal_truth.family_id)
    source_path = output_dir / normal_truth.filename
    before_values = _workbook_values(source_path)
    before_snapshot = _workbook_snapshot(source_path)
    wb = load_workbook(source_path)
    ws = wb[scenario.base_sheet]
    mutation_type = MUTATION_TYPES[mutation_index % len(MUTATION_TYPES)]
    supported = _m4_supported(scenario)
    if scenario.sheet_mode == "horizontal":
        target_cell = ws["N18"]
        target_row = 18
    else:
        candidate_rows = _eligible_detail_rows(scenario, supported, mutation_type)
        target_row = candidate_rows[(seed + mutation_index) % len(candidate_rows)]
        target_cell = ws[f"{scenario.formula_column}{target_row}"]

    before_value = target_cell.value
    before_data_type = target_cell.data_type
    expected_rule: str | None
    expected_subtype: str | None = None
    expected_kind = "m4_candidate"

    if mutation_type == "m4_relative_reference_drift":
        drifted = str(before_value).replace(str(target_row), str(target_row - 1), 1)
        if drifted == before_value:
            col = target_cell.column_letter
            drifted = f"={col}{target_row - 3}-{col}{target_row - 4}"
        target_cell.value = drifted
        expected_rule = "FORMULA_PATTERN_OUTLIER"
        expected_subtype = "RELATIVE_REFERENCE_DRIFT"
    elif mutation_type == "m4_formula_gap":
        target_cell.value = None
        expected_rule = "FORMULA_PATTERN_GAP"
        expected_subtype = "BLANK_GAP_CANDIDATE"
    elif mutation_type == "m4_constant_override":
        target_cell.value = _semantic_for_cell(normal_truth, scenario.base_sheet, target_cell.coordinate) or 0
        expected_rule = "FORMULA_PATTERN_GAP"
        expected_subtype = "CONSTANT_OVERRIDE_CANDIDATE"
    elif mutation_type == "base_broken_ref_token":
        target_cell.value = "=SUM(#REF!)"
        expected_rule = "FORMULA_REF_ERROR"
        expected_kind = "confirmed_error"
    elif mutation_type == "base_visible_error_token":
        target_cell.value = "=#DIV/0!"
        expected_rule = "FORMULA_VISIBLE_ERROR_TOKEN"
        expected_kind = "confirmed_error"
    else:  # pragma: no cover
        raise ValueError(mutation_type)

    workbook_id = normal_truth.workbook_id.replace("-normal-", "-mutated-")
    filename = f"{workbook_id}.xlsx"
    path = output_dir / filename
    wb.save(path)
    after_value = target_cell.value
    after_data_type = target_cell.data_type
    wb.close()

    after_values = _workbook_values(path)
    after_snapshot = _workbook_snapshot(path)
    changed = [key for key, value in after_values.items() if before_values.get(key) != value]
    removed = [key for key in before_values if key not in after_values]
    intended = (scenario.base_sheet, target_cell.coordinate)
    preserved = sorted([f"{sheet}!{cell}" for sheet, cell in set(changed + removed) if (sheet, cell) != intended])
    before_cells = dict(before_snapshot["cells"])
    after_cells = dict(after_snapshot["cells"])
    before_cells.pop(f"{scenario.base_sheet}!{target_cell.coordinate}", None)
    after_cells.pop(f"{scenario.base_sheet}!{target_cell.coordinate}", None)
    snapshot_preserved = before_cells == after_cells and before_snapshot["sheets"] == after_snapshot["sheets"] and before_snapshot["defined_names"] == after_snapshot["defined_names"]
    if before_value == after_value:
        raise RuntimeError(f"Mutation {mutation_type} did not change {scenario.base_sheet}!{target_cell.coordinate}")

    expected: list[ExpectedFinding] = list(normal_truth.expected_findings)
    if supported or expected_kind == "confirmed_error":
        expected.append(
            ExpectedFinding(
                expected_id=f"{normal_truth.workbook_id}-mut-{mutation_index}",
                kind=expected_kind,  # type: ignore[arg-type]
                sheet=scenario.base_sheet,
                cell=target_cell.coordinate,
                rule_code=expected_rule,
                pattern_subtype=expected_subtype,
                family_id=scenario.family_id,
                mutation_type=mutation_type,
                notes=f"Before value: {before_value!r}; after value: {after_value!r}",
            )
        )
    else:
        expected.append(
            ExpectedFinding(
                expected_id=f"{normal_truth.workbook_id}-mut-{mutation_index}-outscope",
                kind="out_of_scope",
                sheet=scenario.base_sheet,
                cell=target_cell.coordinate,
                rule_code=expected_rule,
                pattern_subtype=expected_subtype,
                family_id=scenario.family_id,
                mutation_type=mutation_type,
                notes="Current engine support contract does not score this structural family for M4.",
            )
        )

    original_target_spec = next(
        spec for spec in normal_truth.formula_cells if spec.sheet == scenario.base_sheet and spec.cell == target_cell.coordinate
    )
    formula_cells = [spec for spec in normal_truth.formula_cells if not (spec.sheet == scenario.base_sheet and spec.cell == target_cell.coordinate)]
    formula_cells.append(
        FormulaCellSpec(
            sheet=scenario.base_sheet,
            cell=target_cell.coordinate,
            formula=after_value,
            semantic_value=None,
            role=original_target_spec.role,
            supported_by_current_m4=supported,
        )
    )

    return WorkbookTruth(
        workbook_id=workbook_id,
        parent_id=normal_truth.workbook_id,
        family_id=normal_truth.family_id,
        scenario_id=normal_truth.scenario_id,
        business_domain=normal_truth.business_domain,
        structural_family=normal_truth.structural_family,
        split=normal_truth.split,
        kind="mutated",
        seed=seed,
        filename=filename,
        relative_path=f"workbooks/{filename}",
        sha256=sha256_file(path),
        mutation_type=mutation_type,
        mutation_target=f"{scenario.base_sheet}!{target_cell.coordinate}",
        supported_scope=normal_truth.supported_scope,
        exploratory_scope=normal_truth.exploratory_scope,
        normal_regions=normal_truth.normal_regions,
        formula_cells=formula_cells,
        expected_findings=expected,
        generation_checks={
            "mutation_changed_target": before_value != after_value,
            "mutation_target_only_changed": not preserved and snapshot_preserved,
            "unexpected_changed_cells": preserved,
            "non_target_snapshot_preserved": snapshot_preserved,
            "formula_string_checks": len(formula_cells),
            "semantic_arithmetic_checks": len([spec for spec in normal_truth.formula_cells if spec.semantic_value is not None]),
            "answer_marker_cells": 0,
            "excel_recalculation": "NOT_RUN",
        },
        mutation_snapshot=MutationSnapshot(
            sheet=scenario.base_sheet,
            cell=target_cell.coordinate,
            before_value=before_value,
            before_data_type=before_data_type,
            after_value=after_value,
            after_data_type=after_data_type,
            before_formula=before_value if isinstance(before_value, str) and before_value.startswith("=") else None,
            after_formula=after_value if isinstance(after_value, str) and after_value.startswith("=") else None,
        ),
    )


def generate_dataset(*, run_id: str, pair_count: int, seed: int, dataset_root: Path) -> DatasetManifest:
    rng = random.Random(seed)
    workbooks_dir = dataset_root / "workbooks"
    workbooks_dir.mkdir(parents=True, exist_ok=True)
    family_split = {family_id: _split_for_family(index) for index, family_id in enumerate(FAMILY_ORDER)}
    family_occurrences = {family_id: 0 for family_id in FAMILY_ORDER}
    workbooks: list[WorkbookTruth] = []
    for index in range(pair_count):
        scenario = SCENARIOS[index % len(SCENARIOS)]
        split = family_split[scenario.family_id]
        normal_id = f"{index + 1:03d}-{scenario.scenario_id}-normal-{rng.randrange(1_000_000):06d}"
        normal = create_normal_workbook(scenario=scenario, seed=seed + index * 13, workbook_id=normal_id, split=split, output_dir=workbooks_dir, diversity_features=True)
        mutation_index = family_occurrences[scenario.family_id]
        family_occurrences[scenario.family_id] += 1
        mutated = create_mutated_workbook(normal_truth=normal, mutation_index=mutation_index, seed=seed + index * 17, output_dir=workbooks_dir)
        workbooks.extend([normal, mutated])

    workbooks.sort(key=lambda item: ({"train": 0, "dev": 1, "holdout": 2}[item.split], item.family_id, item.workbook_id))
    config = {"dataset_version": DATASET_VERSION, "generator_version": GENERATOR_VERSION, "pair_count": pair_count, "seed": seed, "family_order": FAMILY_ORDER, "split_by_family": family_split}
    truth_payload = [dataclass_to_dict(item) for item in workbooks]
    manifest = DatasetManifest(
        dataset_version=DATASET_VERSION,
        generator_version=GENERATOR_VERSION,
        run_id=run_id,
        seed=seed,
        pair_count=pair_count,
        family_order=FAMILY_ORDER,
        split_by_family=family_split,  # type: ignore[arg-type]
        workbooks=workbooks,
        config_hash=sha256_json(config),
        source_hash=sha256_json({"scenarios": [asdict(item) for item in SCENARIOS], "mutation_types": MUTATION_TYPES, "m4_supported_scenarios": sorted(M4_SUPPORTED_SCENARIOS)}),
        truth_hash=sha256_json(truth_payload),
    )
    write_json(dataset_root / "truth.json", truth_payload)
    write_json(dataset_root / "manifest.json", dataclass_to_dict(manifest))
    return manifest











