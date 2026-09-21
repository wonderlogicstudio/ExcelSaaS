from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from tools.synthetic_validation.generator import create_normal_workbook
from tools.synthetic_validation.scenarios import (
    BOM_CATALOG_SHEET,
    BUDGET_MONTH_SHEETS,
    INVENTORY_SUPPORT_SHEET,
    INVENTORY_TRANSACTIONS_SHEET,
    SCENARIOS,
)
from tools.synthetic_validation.verify_generated import (
    GenerationValidationError,
    validate_generated_workbook,
)


def _scenario(scenario_id: str):
    return next(scenario for scenario in SCENARIOS if scenario.scenario_id == scenario_id)


def _normal(dataset_root: Path, scenario_id: str):
    workbook_dir = dataset_root / "workbooks"
    return create_normal_workbook(
        scenario=_scenario(scenario_id),
        seed=31,
        workbook_id=f"unit-normal-{scenario_id}",
        split="train",
        output_dir=workbook_dir,
        diversity_features=True,
    )


class DiversityBatchTests(unittest.TestCase):
    def test_inventory_sumifs_uses_transaction_sheet_and_independent_arithmetic(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            truth = _normal(dataset_root, "inventory_cross_sheet")
            first = next(spec for spec in truth.formula_cells if spec.role == "detail")

            self.assertIn("SUMIFS", first.formula)
            self.assertIn(INVENTORY_TRANSACTIONS_SHEET, first.formula)
            wb = load_workbook(dataset_root / truth.relative_path, data_only=False)
            try:
                detail = wb[first.sheet]
                support = wb[INVENTORY_SUPPORT_SHEET]
                transactions = wb[INVENTORY_TRANSACTIONS_SHEET]
                row = int(first.cell[1:])
                sku = detail[f"A{row}"].value
                receipts = sum(
                    transactions[f"D{i}"].value or 0
                    for i in range(2, 101)
                    if transactions[f"A{i}"].value == sku and transactions[f"C{i}"].value == "IN"
                )
                issues = sum(
                    transactions[f"D{i}"].value or 0
                    for i in range(2, 101)
                    if transactions[f"A{i}"].value == sku and transactions[f"C{i}"].value == "OUT"
                )
                expected = support[f"B{row}"].value + receipts - issues
            finally:
                wb.close()

            self.assertEqual(first.semantic_value, expected)
            checks = validate_generated_workbook(dataset_root, truth)
            self.assertEqual(checks["excel_recalculation"], "NOT_RUN")

    def test_bom_vlookup_uses_catalog_and_independent_arithmetic(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            truth = _normal(dataset_root, "manufacturing_bom")
            first = next(spec for spec in truth.formula_cells if spec.role == "detail")

            self.assertIn("VLOOKUP", first.formula)
            self.assertIn(BOM_CATALOG_SHEET, first.formula)
            wb = load_workbook(dataset_root / truth.relative_path, data_only=False)
            try:
                detail = wb[first.sheet]
                catalog = wb[BOM_CATALOG_SHEET]
                row = int(first.cell[1:])
                part_id = detail[f"A{row}"].value
                unit_cost = next(catalog[f"B{i}"].value for i in range(2, 21) if catalog[f"A{i}"].value == part_id)
                expected = round((unit_cost * detail[f"H{row}"].value) / (1 - detail[f"J{row}"].value), 2)
            finally:
                wb.close()

            self.assertEqual(first.semantic_value, expected)
            checks = validate_generated_workbook(dataset_root, truth)
            self.assertEqual(checks["excel_recalculation"], "NOT_RUN")

    def test_budget_month_sheets_drive_independent_arithmetic(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            truth = _normal(dataset_root, "budget_horizontal_months")
            first = truth.formula_cells[0]

            self.assertIn(BUDGET_MONTH_SHEETS[0], first.formula)
            wb = load_workbook(dataset_root / truth.relative_path, data_only=False)
            try:
                month = wb[BUDGET_MONTH_SHEETS[0]]
                expected = month["B16"].value - month["B15"].value
            finally:
                wb.close()

            self.assertEqual(first.semantic_value, expected)
            checks = validate_generated_workbook(dataset_root, truth)
            self.assertEqual(checks["excel_recalculation"], "NOT_RUN")

    def test_missing_inventory_transaction_sheet_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            truth = _normal(dataset_root, "inventory_cross_sheet")
            path = dataset_root / truth.relative_path
            wb = load_workbook(path)
            try:
                del wb[INVENTORY_TRANSACTIONS_SHEET]
                wb.save(path)
            finally:
                wb.close()

            with self.assertRaisesRegex(GenerationValidationError, "missing transaction sheet|missing reference sheet"):
                validate_generated_workbook(dataset_root, truth)

    def test_missing_budget_month_reference_sheet_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            truth = _normal(dataset_root, "budget_horizontal_months")
            path = dataset_root / truth.relative_path
            wb = load_workbook(path)
            try:
                del wb[BUDGET_MONTH_SHEETS[0]]
                wb.save(path)
            finally:
                wb.close()

            with self.assertRaisesRegex(GenerationValidationError, "missing reference sheet"):
                validate_generated_workbook(dataset_root, truth)

    def test_missing_bom_catalog_lookup_key_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            truth = _normal(dataset_root, "manufacturing_bom")
            path = dataset_root / truth.relative_path
            first = next(spec for spec in truth.formula_cells if spec.role == "detail")
            wb = load_workbook(path)
            try:
                catalog = wb[BOM_CATALOG_SHEET]
                detail = wb[first.sheet]
                part_id = detail[f"A{int(first.cell[1:])}"].value
                for row in range(2, 21):
                    if catalog[f"A{row}"].value == part_id:
                        catalog[f"A{row}"] = "PART-MISSING"
                        break
                wb.save(path)
            finally:
                wb.close()

            with self.assertRaisesRegex(GenerationValidationError, "missing exact lookup key"):
                validate_generated_workbook(dataset_root, truth)


if __name__ == "__main__":
    unittest.main()
