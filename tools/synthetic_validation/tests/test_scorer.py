from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from tools.synthetic_validation.generator import generate_dataset
from tools.synthetic_validation.runner import (
    EngineRunConfig,
    _completion,
    _resume_ok,
    _settings_hash,
    _tool_hash,
    validate_frozen_contract,
    write_frozen_contract,
)
from tools.synthetic_validation.schema import DatasetManifest, ExpectedFinding, FormulaCellSpec, NormalRegion, WorkbookTruth, read_json, sha256_file, write_json
from tools.synthetic_validation.scorer import score_results
from tools.synthetic_validation.verify_generated import GenerationValidationError, validate_generated_workbook


def _truth(workbook_id: str, expected: list[ExpectedFinding] | None = None, *, extras: list[ExpectedFinding] | None = None) -> WorkbookTruth:
    formula_cells = [
        FormulaCellSpec("Sheet1", "A1", "=1+1", 2, "detail", True),
        FormulaCellSpec("Sheet1", "B3", "=2+2", 4, "detail", True),
        FormulaCellSpec("Sheet1", "C1", "=3+3", 6, "detail", True),
        FormulaCellSpec("Sheet1", "C2", "=4+4", 8, "normal_exception", True),
    ]
    return WorkbookTruth(
        workbook_id=workbook_id,
        parent_id=None,
        family_id="F-test",
        scenario_id="unit",
        business_domain="unit business",
        structural_family="unit structure",
        split="train",
        kind="mutated",
        seed=1,
        filename=f"{workbook_id}.xlsx",
        relative_path=f"workbooks/{workbook_id}.xlsx",
        sha256="unused",
        mutation_type="unit_mutation",
        mutation_target="Sheet1!A1",
        supported_scope=["unit"],
        exploratory_scope=[],
        normal_regions=[NormalRegion("Sheet1", "A1:C2", "unit")],
        formula_cells=formula_cells,
        expected_findings=(expected or []) + (extras or []),
    )


def _manifest(workbooks: list[WorkbookTruth]) -> DatasetManifest:
    return DatasetManifest("unit", "unit", "unit", 1, len(workbooks), ["F-test"], {"F-test": "train"}, workbooks, "c", "s", "t")


def _write_status(root: Path, raws: dict[str, dict], *, status: str = "COMPLETED") -> None:
    records = []
    for workbook_id, raw in raws.items():
        raw_path = root / "raw_engine" / f"{workbook_id}.json"
        write_json(raw_path, raw)
        records.append({"workbook_id": workbook_id, "status": raw.get("status", status), "elapsed_ms": 1, "raw_output": str(raw_path), "raw_sha256": sha256_file(raw_path)})
    write_json(root / "engine_status.json", {"records": records, "completion": {"workbooks_total": len(records), "completed": len(records)}})


class ScorerTests(unittest.TestCase):
    def test_precision_recall_example_tp9_fn1_fp2(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbooks: list[WorkbookTruth] = []
            raws: dict[str, dict] = {}
            for index in range(10):
                expected = ExpectedFinding(f"e{index}", "m4_candidate", "Sheet1", "A1", "FORMULA_PATTERN_GAP", "BLANK_GAP_CANDIDATE")
                truth = _truth(f"wb{index}", [expected])
                workbooks.append(truth)
                candidates = []
                if index < 9:
                    candidates.append({"sheet": "Sheet1", "cell": "A1", "rule_code": "FORMULA_PATTERN_GAP", "formula_pattern": {"pattern_subtype": "BLANK_GAP_CANDIDATE"}})
                if index == 0:
                    candidates.extend([
                        {"sheet": "Sheet1", "cell": "C1", "rule_code": "FORMULA_PATTERN_GAP", "formula_pattern": {"pattern_subtype": "BLANK_GAP_CANDIDATE"}},
                        {"sheet": "Sheet1", "cell": "C2", "rule_code": "FORMULA_PATTERN_OUTLIER", "formula_pattern": {"pattern_subtype": "RELATIVE_REFERENCE_DRIFT"}},
                    ])
                raws[truth.workbook_id] = {"status": "COMPLETED", "base": {"findings": []}, "audit": {"status": "COMPLETED", "candidates": candidates}}
            _write_status(root, raws)
            summary = score_results(manifest=_manifest(workbooks), dataset_root=root / "dataset", run_root=root)
            metrics = summary["metrics_full_contract"]
            self.assertEqual(metrics["tp"], 9)
            self.assertEqual(metrics["fn"], 1)
            self.assertEqual(metrics["fp"], 2)
            self.assertEqual(metrics["precision"], 0.8182)
            self.assertEqual(metrics["recall"], 0.9)

    def test_duplicate_actual_is_counted_not_false_positive_and_ranges_match_both_directions(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            range_expected = ExpectedFinding("range", "confirmed_error", "Sheet1", "B2:B4", "FORMULA_REF_ERROR")
            reverse_range = ExpectedFinding("reverse", "confirmed_error", "Sheet1", "B3", "FORMULA_REF_ERROR")
            duplicate_expected = ExpectedFinding("dup", "m4_candidate", "Sheet1", "A1", "FORMULA_PATTERN_GAP", "BLANK_GAP_CANDIDATE")
            workbooks = [_truth("range", [range_expected]), _truth("reverse", [reverse_range]), _truth("dup", [duplicate_expected])]
            raws = {
                "range": {"status": "COMPLETED", "base": {"findings": [{"sheet": "Sheet1", "cell": "B3", "rule_code": "FORMULA_REF_ERROR"}]}, "audit": {"status": "COMPLETED", "candidates": []}},
                "reverse": {"status": "COMPLETED", "base": {"findings": [{"sheet": "Sheet1", "cell": "B2:B4", "rule_code": "FORMULA_REF_ERROR"}]}, "audit": {"status": "COMPLETED", "candidates": []}},
                "dup": {"status": "COMPLETED", "base": {"findings": []}, "audit": {"status": "COMPLETED", "candidates": [
                    {"sheet": "Sheet1", "cell": "A1", "rule_code": "FORMULA_PATTERN_GAP", "formula_pattern": {"pattern_subtype": "BLANK_GAP_CANDIDATE"}},
                    {"sheet": "Sheet1", "cell": "A1", "rule_code": "FORMULA_PATTERN_GAP", "formula_pattern": {"pattern_subtype": "BLANK_GAP_CANDIDATE"}},
                ]}},
            }
            _write_status(root, raws)
            metrics = score_results(manifest=_manifest(workbooks), dataset_root=root / "dataset", run_root=root)["metrics_full_contract"]
            self.assertEqual(metrics["tp"], 3)
            self.assertEqual(metrics["fp"], 0)
            self.assertEqual(metrics["duplicates_not_fp"], 1)

    def test_unmatched_malformed_ambiguous_normalexception_and_partial_base_are_classified(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            confirmed = ExpectedFinding("base", "confirmed_error", "Sheet1", "B3", "FORMULA_REF_ERROR")
            ambiguous = ExpectedFinding("amb", "user_confirmation", "Sheet1", "C1", None)
            normal_exception = ExpectedFinding("norm", "normal_exception", "Sheet1", "C2", None)
            truth = _truth("mixed", [confirmed], extras=[ambiguous, normal_exception])
            raws = {
                "mixed": {"status": "PARTIAL", "base": {"findings": [
                    {"sheet": "Sheet1", "cell": "B3", "rule_code": "FORMULA_REF_ERROR"},
                    {"sheet": "Sheet1", "cell": "A:", "rule_code": "FORMULA_REF_ERROR"},
                ]}, "audit": {"status": "FAILED", "candidates": [
                    {"sheet": "Sheet1", "cell": "C1", "rule_code": "FORMULA_PATTERN_GAP", "formula_pattern": {"pattern_subtype": "BLANK_GAP_CANDIDATE"}},
                    {"sheet": "Sheet1", "cell": "C2", "rule_code": "FORMULA_PATTERN_GAP", "formula_pattern": {"pattern_subtype": "BLANK_GAP_CANDIDATE"}},
                ]}},
            }
            _write_status(root, raws, status="PARTIAL")
            summary = score_results(manifest=_manifest([truth]), dataset_root=root / "dataset", run_root=root)
            metrics = summary["metrics_full_contract"]
            self.assertEqual(metrics["tp"], 1)
            self.assertEqual(metrics["fp"], 1)
            self.assertEqual(metrics["excluded_ambiguous_or_out_of_scope_not_passes"], 1)
            self.assertEqual(metrics["unjudged_findings_not_passes"], 1)
            self.assertEqual(metrics["workbooks_partial"], 1)


class RunnerContractTests(unittest.TestCase):
    def test_contract_and_resume_detect_tamper_and_interrupted_totals(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = generate_dataset(run_id="unit", pair_count=1, seed=2, dataset_root=root / "dataset")
            config = EngineRunConfig()
            python_exe = Path("python")
            write_frozen_contract(manifest=manifest, dataset_root=root / "dataset", run_root=root, python_exe=python_exe, config=config)
            validate_frozen_contract(manifest=manifest, dataset_root=root / "dataset", run_root=root, python_exe=python_exe, config=config)
            original_truth = read_json(root / "dataset" / "truth.json")
            tampered = list(original_truth)
            tampered[0] = {**tampered[0], "workbook_id": "tampered"}
            write_json(root / "dataset" / "truth.json", tampered)
            with self.assertRaises(RuntimeError):
                validate_frozen_contract(manifest=manifest, dataset_root=root / "dataset", run_root=root, python_exe=python_exe, config=config)
            write_json(root / "dataset" / "truth.json", original_truth)

            status = {"config_hash": manifest.config_hash, "source_hash": manifest.source_hash, "truth_hash": manifest.truth_hash, "tool_hash": _tool_hash(), "engine_source_hash": read_json(root / "frozen_contract.json")["engine_source_hash"], "settings_hash": _settings_hash(python_exe, config), "engine_descriptor": read_json(root / "frozen_contract.json")["engine_descriptor"], "file_hashes": {item.workbook_id: sha256_file(root / "dataset" / item.relative_path) for item in manifest.workbooks}, "records": [{"workbook_id": manifest.workbooks[0].workbook_id, "status": "TIMEOUT", "raw_output": None, "raw_sha256": None}]}
            self.assertTrue(_resume_ok(status, manifest, root / "dataset", python_exe, config))
            status["tool_hash"] = "bad"
            self.assertFalse(_resume_ok(status, manifest, root / "dataset", python_exe, config))
            self.assertEqual(_completion(status["records"], len(manifest.workbooks)), {"workbooks_total": 2, "completed": 0, "partial": 0, "failed": 0, "timeouts": 1, "not_run": 1})

    def test_resume_rejects_raw_hash_tamper(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = generate_dataset(run_id="unit", pair_count=1, seed=4, dataset_root=root / "dataset")
            raw_path = root / "raw_engine" / "one.json"
            write_json(raw_path, {"status": "COMPLETED", "base": {"findings": []}, "audit": {"candidates": []}})
            config = EngineRunConfig()
            python_exe = Path("python")
            write_frozen_contract(manifest=manifest, dataset_root=root / "dataset", run_root=root, python_exe=python_exe, config=config)
            contract = read_json(root / "frozen_contract.json")
            status = {"config_hash": manifest.config_hash, "source_hash": manifest.source_hash, "truth_hash": manifest.truth_hash, "tool_hash": _tool_hash(), "engine_source_hash": contract["engine_source_hash"], "settings_hash": _settings_hash(python_exe, config), "engine_descriptor": contract["engine_descriptor"], "file_hashes": contract["file_hashes"], "records": [{"workbook_id": manifest.workbooks[0].workbook_id, "status": "COMPLETED", "raw_output": str(raw_path), "raw_sha256": "bad"}]}
            self.assertFalse(_resume_ok(status, manifest, root / "dataset", python_exe, config))


class GenerationTests(unittest.TestCase):
    def test_generation_50_pairs_covers_all_families_and_mutation_types_without_markers(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = generate_dataset(run_id="unit", pair_count=50, seed=3, dataset_root=root / "dataset")
            self.assertEqual(len({item.family_id for item in manifest.workbooks if item.kind == "normal"}), 10)
            by_family: dict[str, set[str]] = {}
            for truth in manifest.workbooks:
                if truth.kind == "mutated":
                    by_family.setdefault(truth.family_id, set()).add(truth.mutation_type or "")
                checks = validate_generated_workbook(root / "dataset", truth)
                self.assertEqual(checks["answer_marker_cells"], 0)
                self.assertEqual(checks["excel_recalculation"], "NOT_RUN")
            self.assertTrue(all(len(types) == 5 for types in by_family.values()))

    def test_generated_normal_tamper_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = generate_dataset(run_id="unit", pair_count=1, seed=5, dataset_root=root / "dataset")
            truth = manifest.workbooks[0]
            from openpyxl import load_workbook

            path = root / "dataset" / truth.relative_path
            wb = load_workbook(path)
            ws = wb[truth.formula_cells[0].sheet]
            ws[truth.formula_cells[0].cell] = "=A999"
            wb.save(path)
            wb.close()
            with self.assertRaises(GenerationValidationError):
                validate_generated_workbook(root / "dataset", truth)


if __name__ == "__main__":
    unittest.main()
