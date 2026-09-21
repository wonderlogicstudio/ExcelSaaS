from __future__ import annotations

import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from tools.synthetic_validation.generator import (
    create_mutated_workbook,
    create_normal_workbook,
)
from tools.synthetic_validation.scenarios import (
    INVENTORY_DETAIL_SHEET,
    INVENTORY_SUPPORT_SHEET,
    SCENARIOS,
)
from tools.synthetic_validation.schema import FormulaCellSpec, WorkbookTruth
from tools.synthetic_validation.verify_generated import (
    GenerationValidationError,
    _formula_references,
    validate_generated_workbook,
)


def _inventory_scenario():
    return next(scenario for scenario in SCENARIOS if scenario.scenario_id == "inventory_cross_sheet")


def _replace_spec(truth: WorkbookTruth, sheet: str, cell: str, formula: str) -> WorkbookTruth:
    specs: list[FormulaCellSpec] = []
    for spec in truth.formula_cells:
        specs.append(replace(spec, formula=formula) if spec.sheet == sheet and spec.cell == cell else spec)
    return replace(truth, formula_cells=specs)


def _write_formula(dataset_root: Path, truth: WorkbookTruth, sheet: str, cell: str, formula: str) -> None:
    path = dataset_root / truth.relative_path
    wb = load_workbook(path)
    try:
        wb[sheet][cell] = formula
        wb.save(path)
    finally:
        wb.close()


class KoreanSheetReferenceTests(unittest.TestCase):
    def test_inventory_cross_sheet_uses_quoted_korean_support_reference_and_independent_arithmetic(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            workbook_dir = dataset_root / "workbooks"
            scenario = _inventory_scenario()

            normal = create_normal_workbook(scenario=scenario, seed=17, workbook_id="unit-normal-korean", split="train", output_dir=workbook_dir)

            self.assertEqual(normal.normal_regions[0].sheet, INVENTORY_DETAIL_SHEET)
            first = next(spec for spec in normal.formula_cells if spec.role == "detail")
            self.assertEqual(first.sheet, INVENTORY_DETAIL_SHEET)
            self.assertEqual(first.formula, f"='{INVENTORY_SUPPORT_SHEET}'!B10+G10-H10")
            self.assertIn((INVENTORY_SUPPORT_SHEET, "B10"), _formula_references(first.formula))

            wb = load_workbook(dataset_root / normal.relative_path, data_only=False)
            try:
                self.assertIn(INVENTORY_DETAIL_SHEET, wb.sheetnames)
                self.assertIn(INVENTORY_SUPPORT_SHEET, wb.sheetnames)
                expected = wb[INVENTORY_SUPPORT_SHEET]["B10"].value + wb[INVENTORY_DETAIL_SHEET]["G10"].value - wb[INVENTORY_DETAIL_SHEET]["H10"].value
            finally:
                wb.close()
            self.assertEqual(first.semantic_value, expected)
            self.assertEqual(validate_generated_workbook(dataset_root, normal)["excel_recalculation"], "NOT_RUN")

    def test_inventory_reference_drift_mutation_targets_only_korean_detail_sheet_cell(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            workbook_dir = dataset_root / "workbooks"
            scenario = _inventory_scenario()
            normal = create_normal_workbook(scenario=scenario, seed=17, workbook_id="unit-normal-korean", split="train", output_dir=workbook_dir)

            mutated = create_mutated_workbook(normal_truth=normal, mutation_index=0, seed=19, output_dir=workbook_dir)

            self.assertTrue(mutated.mutation_target.startswith(f"{INVENTORY_DETAIL_SHEET}!"))
            self.assertEqual(mutated.generation_checks["mutation_target_only_changed"], True)
            expected = [finding for finding in mutated.expected_findings if finding.kind == "m4_candidate"]
            self.assertEqual(len(expected), 1)
            self.assertEqual(expected[0].sheet, INVENTORY_DETAIL_SHEET)
            self.assertEqual(expected[0].rule_code, "FORMULA_PATTERN_OUTLIER")
            self.assertEqual(expected[0].pattern_subtype, "RELATIVE_REFERENCE_DRIFT")
            self.assertEqual(validate_generated_workbook(dataset_root, mutated)["excel_recalculation"], "NOT_RUN")

    def test_missing_korean_support_sheet_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            workbook_dir = dataset_root / "workbooks"
            normal = create_normal_workbook(scenario=_inventory_scenario(), seed=17, workbook_id="unit-normal-korean", split="train", output_dir=workbook_dir)
            workbook_path = dataset_root / normal.relative_path
            wb = load_workbook(workbook_path)
            try:
                del wb[INVENTORY_SUPPORT_SHEET]
                wb.save(workbook_path)
            finally:
                wb.close()

            with self.assertRaisesRegex(GenerationValidationError, "missing reference sheet"):
                validate_generated_workbook(dataset_root, normal)

    def test_invalid_quoted_korean_reference_sheet_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            dataset_root = Path(temp) / "dataset"
            workbook_dir = dataset_root / "workbooks"
            normal = create_normal_workbook(scenario=_inventory_scenario(), seed=17, workbook_id="unit-normal-korean", split="train", output_dir=workbook_dir)
            first = next(spec for spec in normal.formula_cells if spec.role == "detail")
            bad_formula = "=\'없는 시트\'!B10+G10-H10"
            _write_formula(dataset_root, normal, first.sheet, first.cell, bad_formula)
            tampered = _replace_spec(normal, first.sheet, first.cell, bad_formula)

            with self.assertRaisesRegex(GenerationValidationError, "missing reference sheet"):
                validate_generated_workbook(dataset_root, tampered)


if __name__ == "__main__":
    unittest.main()
