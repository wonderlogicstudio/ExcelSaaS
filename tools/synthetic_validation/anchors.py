from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .generator import (
    PAYROLL_OVERTIME_PREMIUM,
    PAYROLL_OVERTIME_THRESHOLD,
    _workbook_snapshot,
    _workbook_values,
    create_normal_workbook,
)
from .scenarios import SCENARIOS
from .schema import (
    ExpectedFinding,
    FormulaCellSpec,
    MutationSnapshot,
    WorkbookTruth,
    sha256_file,
)

PAYROLL_ANCHOR_TARGET_ROW = 12
PAYROLL_ANCHOR_HOURS = 176
PAYROLL_ALT_PREMIUM_CELL = "B6"
PAYROLL_ALT_PREMIUM = 0.25


def payroll_anchor_scenario():
    return next(item for item in SCENARIOS if item.scenario_id == "payroll_overtime")


def _payroll_expected_value(wb: Any, row: int, premium_cell: str = "B5") -> float | int:
    ws = wb["Payroll"]
    rate_card = wb["Rate Card"]
    hours = ws[f"H{row}"].value
    rate = rate_card[f"B{row}"].value
    threshold = ws["C5"].value if ws["C5"].value is not None else PAYROLL_OVERTIME_THRESHOLD
    premium = ws[premium_cell].value if ws[premium_cell].value is not None else PAYROLL_OVERTIME_PREMIUM
    return hours * rate + max(hours - threshold, 0) * rate * premium


def prepare_payroll_anchor_normal(*, seed: int, workbook_id: str, split: str, output_dir: Path) -> tuple[WorkbookTruth, dict[str, Any]]:
    scenario = payroll_anchor_scenario()
    truth = create_normal_workbook(scenario=scenario, seed=seed, workbook_id=workbook_id, split=split, output_dir=output_dir)
    path = output_dir / truth.filename
    wb = load_workbook(path)
    try:
        ws = wb[scenario.base_sheet]
        ws[f"H{PAYROLL_ANCHOR_TARGET_ROW}"] = PAYROLL_ANCHOR_HOURS
        ws["A6"] = "Weekend premium"
        ws[PAYROLL_ALT_PREMIUM_CELL] = PAYROLL_ALT_PREMIUM
        ws[PAYROLL_ALT_PREMIUM_CELL].number_format = "0.0%"
        before_expected = _payroll_expected_value(wb, PAYROLL_ANCHOR_TARGET_ROW, "B5")
        after_expected = _payroll_expected_value(wb, PAYROLL_ANCHOR_TARGET_ROW, PAYROLL_ALT_PREMIUM_CELL)
        wb.save(path)
    finally:
        wb.close()

    wb = load_workbook(path, data_only=False)
    try:
        formula_cells: list[FormulaCellSpec] = []
        for spec in truth.formula_cells:
            if spec.sheet == scenario.base_sheet and spec.semantic_value is not None:
                formula_cells.append(replace(spec, semantic_value=_payroll_expected_value(wb, int(spec.cell[1:]), "B5")))
            else:
                formula_cells.append(spec)
    finally:
        wb.close()

    updated = replace(
        truth,
        sha256=sha256_file(path),
        formula_cells=formula_cells,
        generation_checks={
            **truth.generation_checks,
            "anchor_fixture_target_row": PAYROLL_ANCHOR_TARGET_ROW,
            "anchor_fixture_hours": PAYROLL_ANCHOR_HOURS,
            "anchor_fixture_before_expected": before_expected,
            "anchor_fixture_after_expected": after_expected,
            "excel_recalculation": "NOT_RUN",
        },
    )
    return updated, {
        "target_row": PAYROLL_ANCHOR_TARGET_ROW,
        "target_cell": f"{scenario.formula_column}{PAYROLL_ANCHOR_TARGET_ROW}",
        "hours": PAYROLL_ANCHOR_HOURS,
        "threshold": PAYROLL_OVERTIME_THRESHOLD,
        "normal_premium_cell": "B5",
        "normal_premium": PAYROLL_OVERTIME_PREMIUM,
        "mutated_premium_cell": PAYROLL_ALT_PREMIUM_CELL,
        "mutated_premium": PAYROLL_ALT_PREMIUM,
        "before_expected": before_expected,
        "after_expected": after_expected,
        "excel_recalculation": "NOT_RUN",
    }


def create_payroll_anchor_mutant(*, normal_truth: WorkbookTruth, output_dir: Path) -> tuple[WorkbookTruth, dict[str, Any]]:
    scenario = payroll_anchor_scenario()
    source_path = output_dir / normal_truth.filename
    before_values = _workbook_values(source_path)
    before_snapshot = _workbook_snapshot(source_path)
    wb = load_workbook(source_path)
    ws = wb[scenario.base_sheet]
    target = ws[f"{scenario.formula_column}{PAYROLL_ANCHOR_TARGET_ROW}"]
    before_value = target.value
    before_data_type = target.data_type
    if not isinstance(before_value, str) or "$B$5" not in before_value:
        wb.close()
        raise RuntimeError(f"Payroll anchor target does not contain $B$5: {before_value!r}")
    after_value = before_value.replace("$B$5", "$B$6")
    target.value = after_value
    before_expected = _payroll_expected_value(wb, PAYROLL_ANCHOR_TARGET_ROW, "B5")
    after_expected = _payroll_expected_value(wb, PAYROLL_ANCHOR_TARGET_ROW, PAYROLL_ALT_PREMIUM_CELL)

    workbook_id = normal_truth.workbook_id.replace("-normal-", "-mutated-")
    filename = f"{workbook_id}.xlsx"
    path = output_dir / filename
    wb.save(path)
    after_data_type = target.data_type
    wb.close()

    after_values = _workbook_values(path)
    after_snapshot = _workbook_snapshot(path)
    intended = (scenario.base_sheet, target.coordinate)
    changed = [key for key, value in after_values.items() if before_values.get(key) != value]
    removed = [key for key in before_values if key not in after_values]
    preserved = sorted([f"{sheet}!{cell}" for sheet, cell in set(changed + removed) if (sheet, cell) != intended])
    before_cells = dict(before_snapshot["cells"])
    after_cells = dict(after_snapshot["cells"])
    before_cells.pop(f"{scenario.base_sheet}!{target.coordinate}", None)
    after_cells.pop(f"{scenario.base_sheet}!{target.coordinate}", None)
    snapshot_preserved = before_cells == after_cells and before_snapshot["sheets"] == after_snapshot["sheets"] and before_snapshot["defined_names"] == after_snapshot["defined_names"]
    if before_expected == after_expected:
        raise RuntimeError("Anchor mutation did not change the independent expected value.")

    formula_cells = [
        spec
        for spec in normal_truth.formula_cells
        if not (spec.sheet == scenario.base_sheet and spec.cell == target.coordinate)
    ]
    original = next(spec for spec in normal_truth.formula_cells if spec.sheet == scenario.base_sheet and spec.cell == target.coordinate)
    formula_cells.append(replace(original, formula=after_value, semantic_value=None))
    expected = list(normal_truth.expected_findings)
    expected.append(
        ExpectedFinding(
            expected_id=f"{normal_truth.workbook_id}-anchor-mut",
            kind="m4_candidate",
            sheet=scenario.base_sheet,
            cell=target.coordinate,
            rule_code="FORMULA_PATTERN_OUTLIER",
            pattern_subtype="REFERENCE_CELL_DRIFT",
            family_id=scenario.family_id,
            mutation_type="m4_anchor_reference_shift",
            notes=f"Anchor changed from $B$5 to $B$6; independent expected value {before_expected!r} -> {after_expected!r}.",
        )
    )
    truth = replace(
        normal_truth,
        workbook_id=workbook_id,
        parent_id=normal_truth.workbook_id,
        filename=filename,
        relative_path=f"workbooks/{filename}",
        sha256=sha256_file(path),
        kind="mutated",
        mutation_type="m4_anchor_reference_shift",
        mutation_target=f"{scenario.base_sheet}!{target.coordinate}",
        formula_cells=formula_cells,
        expected_findings=expected,
        generation_checks={
            **normal_truth.generation_checks,
            "mutation_changed_target": before_value != after_value,
            "mutation_target_only_changed": not preserved and snapshot_preserved,
            "unexpected_changed_cells": preserved,
            "non_target_snapshot_preserved": snapshot_preserved,
            "anchor_fixture_before_expected": before_expected,
            "anchor_fixture_after_expected": after_expected,
            "excel_recalculation": "NOT_RUN",
        },
        mutation_snapshot=MutationSnapshot(
            sheet=scenario.base_sheet,
            cell=target.coordinate,
            before_value=before_value,
            before_data_type=before_data_type,
            after_value=after_value,
            after_data_type=after_data_type,
            before_formula=before_value,
            after_formula=after_value,
        ),
    )
    return truth, {
        "target_cell": target.coordinate,
        "before_formula": before_value,
        "after_formula": after_value,
        "before_expected": before_expected,
        "after_expected": after_expected,
        "mutation_target_only_changed": truth.generation_checks["mutation_target_only_changed"],
        "excel_recalculation": "NOT_RUN",
    }
