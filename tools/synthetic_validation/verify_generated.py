from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string, range_boundaries

from .scenarios import (
    BOM_CATALOG_SHEET,
    BUDGET_MONTH_SHEETS,
    INVENTORY_SUPPORT_SHEET,
    INVENTORY_TRANSACTIONS_SHEET,
    scenario_by_family,
)
from .schema import FormulaCellSpec, WorkbookTruth

_REF_PATTERN = re.compile(r"(?:(?:'(?P<quoted>[^']+)'|(?P<sheet>[A-Za-z0-9_ ]+))!)?(?P<col>\$?[A-Z]{1,3})(?P<row>\$?[1-9][0-9]*)")
PAYROLL_OVERTIME_THRESHOLD = 160
PAYROLL_OVERTIME_PREMIUM = 0.5
INVENTORY_TRANSACTION_RANGE_END = 100
BOM_CATALOG_RANGE_END = 20


class GenerationValidationError(AssertionError):
    pass


def _formula_references(formula: str) -> list[tuple[str | None, str]]:
    refs: list[tuple[str | None, str]] = []
    for match in _REF_PATTERN.finditer(formula):
        sheet = match.group("quoted") or match.group("sheet")
        col = match.group("col").replace("$", "")
        row = match.group("row").replace("$", "")
        refs.append((sheet, f"{col}{row}"))
    return refs


def _range_cells(range_ref: str) -> set[str]:
    min_col, min_row, max_col, max_row = range_boundaries(range_ref)
    return {f"{get_column_letter(col)}{row}" for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)}


def _same_number(left: Any, right: Any) -> bool:
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) < 0.000001
    return left == right


def _inventory_support_sheet(wb: Any) -> str:
    if INVENTORY_SUPPORT_SHEET in wb.sheetnames:
        return INVENTORY_SUPPORT_SHEET
    if "Opening" in wb.sheetnames:
        return "Opening"
    raise GenerationValidationError(f"inventory_cross_sheet: missing reference sheet {INVENTORY_SUPPORT_SHEET!r} or 'Opening'")


def _inventory_transaction_total(wb: Any, sku: str, movement: str) -> float | int:
    if INVENTORY_TRANSACTIONS_SHEET not in wb.sheetnames:
        raise GenerationValidationError(f"inventory_cross_sheet: missing transaction sheet {INVENTORY_TRANSACTIONS_SHEET!r}")
    transactions = wb[INVENTORY_TRANSACTIONS_SHEET]
    total = 0
    for row in range(2, INVENTORY_TRANSACTION_RANGE_END + 1):
        if transactions[f"A{row}"].value == sku and transactions[f"C{row}"].value == movement:
            total += transactions[f"D{row}"].value or 0
    return total


def _bom_catalog_unit_cost(wb: Any, part_id: str) -> float | int:
    if BOM_CATALOG_SHEET not in wb.sheetnames:
        raise GenerationValidationError(f"manufacturing_bom: missing catalog sheet {BOM_CATALOG_SHEET!r}")
    catalog = wb[BOM_CATALOG_SHEET]
    for row in range(2, BOM_CATALOG_RANGE_END + 1):
        if catalog[f"A{row}"].value == part_id:
            return catalog[f"B{row}"].value
    raise GenerationValidationError(f"manufacturing_bom: missing exact lookup key {part_id!r} in {BOM_CATALOG_SHEET!r}")


def _budget_month_value(wb: Any, month_index: int, ws: Any, col: str) -> float | int:
    if not 1 <= month_index <= len(BUDGET_MONTH_SHEETS):
        raise GenerationValidationError(f"budget_horizontal_months: invalid month index {month_index}")
    month_sheet = BUDGET_MONTH_SHEETS[month_index - 1]
    if month_sheet in wb.sheetnames:
        month_ws = wb[month_sheet]
        return month_ws["B16"].value - month_ws["B15"].value
    return ws[f"{col}16"].value - ws[f"{col}15"].value


def _business_value(scenario: Any, row: int, ws: Any, prior: dict[int, float | int]) -> float | int:
    sid = scenario.scenario_id
    if sid == "retail_sales_vertical":
        return round(ws[f"E{row}"].value * ws[f"F{row}"].value * (1 - ws[f"H{row}"].value), 2)
    if sid == "purchase_orders_table":
        return ws[f"F{row}"].value * ws[f"G{row}"].value if ws[f"E{row}"].value > 0 else 0
    if sid == "inventory_cross_sheet":
        opening = ws.parent[_inventory_support_sheet(ws.parent)][f"B{row}"].value
        if INVENTORY_TRANSACTIONS_SHEET in ws.parent.sheetnames:
            sku = ws[f"A{row}"].value
            receipts = _inventory_transaction_total(ws.parent, sku, "IN")
            issues = _inventory_transaction_total(ws.parent, sku, "OUT")
            return opening + receipts - issues
        return opening + ws[f"G{row}"].value - ws[f"H{row}"].value
    if sid == "budget_horizontal_months":
        return ws[f"L{row}"].value - ws[f"M{row}"].value
    if sid == "payroll_overtime":
        hours = ws[f"H{row}"].value
        rate = ws.parent["Rate Card"][f"B{row}"].value
        threshold = ws["C5"].value if ws["C5"].value is not None else PAYROLL_OVERTIME_THRESHOLD
        premium = ws["B5"].value if ws["B5"].value is not None else PAYROLL_OVERTIME_PREMIUM
        return hours * rate + max(hours - threshold, 0) * rate * premium
    if sid == "project_profit":
        uplift = 0.18 if row in scenario.normal_exception_rows else 0.12
        return round(ws[f"I{row}"].value - ws[f"J{row}"].value + ws[f"H{row}"].value * uplift, 2)
    if sid == "receivables_aging":
        amount = ws[f"F{row}"].value
        days = ws[f"G{row}"].value
        return round(amount * 0.97, 2) if days > 30 else amount
    if sid == "manufacturing_bom":
        unit_cost = _bom_catalog_unit_cost(ws.parent, ws[f"A{row}"].value) if BOM_CATALOG_SHEET in ws.parent.sheetnames else ws[f"G{row}"].value
        return round((unit_cost * ws[f"H{row}"].value) / (1 - ws[f"J{row}"].value), 2)
    if sid == "cashflow_calendar":
        net = ws[f"H{row}"].value - ws[f"I{row}"].value
        return net if row == scenario.start_row else prior[row - 1] + net
    if sid == "subscriptions_kpi":
        active = ws[f"J{row}"].value
        return round(ws[f"K{row}"].value / active, 4) if active else 0
    raise ValueError(sid)


def _payroll_allowed_formulas(row: int) -> set[str]:
    return {
        f"=H{row}*'Rate Card'!B{row}+IF(H{row}>160,(H{row}-160)*'Rate Card'!B{row}*0.5,0)",
        f"=H{row}*'Rate Card'!$B{row}+IF(H{row}>C$5,(H{row}-C$5)*'Rate Card'!$B{row}*$B$5,0)",
    }


def _spec_by_cell(specs: list[FormulaCellSpec]) -> dict[tuple[str, str], FormulaCellSpec]:
    return {(spec.sheet, spec.cell): spec for spec in specs}


def _declared_mutation_target(truth: WorkbookTruth, wb: Any, spec_lookup: dict[tuple[str, str], FormulaCellSpec]) -> tuple[str, str] | None:
    if truth.kind != "mutated":
        return None
    if not truth.mutation_target or truth.mutation_snapshot is None:
        raise GenerationValidationError(f"{truth.workbook_id}: mutated workbook missing mutation target metadata")
    if "!" not in truth.mutation_target:
        raise GenerationValidationError(f"{truth.workbook_id}: invalid mutation target {truth.mutation_target!r}")
    target_sheet, target_cell = truth.mutation_target.rsplit("!", 1)
    snapshot = truth.mutation_snapshot
    if snapshot.sheet != target_sheet or snapshot.cell != target_cell:
        raise GenerationValidationError(f"{truth.workbook_id}: mutation target does not match snapshot")
    if snapshot.before_value == snapshot.after_value and snapshot.before_data_type == snapshot.after_data_type:
        raise GenerationValidationError(f"{truth.workbook_id} {target_sheet}!{target_cell}: mutation snapshot did not change target")
    spec = spec_lookup.get((target_sheet, target_cell))
    if spec is None:
        raise GenerationValidationError(f"{truth.workbook_id} {target_sheet}!{target_cell}: mutation target missing formula spec")
    if spec.formula != snapshot.after_value:
        raise GenerationValidationError(f"{truth.workbook_id} {target_sheet}!{target_cell}: mutation spec does not match snapshot after value")
    if target_sheet not in wb.sheetnames:
        raise GenerationValidationError(f"{truth.workbook_id} {target_sheet}!{target_cell}: mutation target sheet missing")
    actual_cell = wb[target_sheet][target_cell]
    if actual_cell.value != snapshot.after_value or actual_cell.data_type != snapshot.after_data_type:
        raise GenerationValidationError(f"{truth.workbook_id} {target_sheet}!{target_cell}: mutation target does not match snapshot after state")
    if actual_cell.value == snapshot.before_value and actual_cell.data_type == snapshot.before_data_type:
        raise GenerationValidationError(f"{truth.workbook_id} {target_sheet}!{target_cell}: mutation target still matches snapshot before state")
    return (target_sheet, target_cell)


def _check_formula_references(truth: WorkbookTruth, wb: Any, sheet_name: str, cell: str, formula: str) -> int:
    checks = 0
    for ref_sheet, ref_cell in _formula_references(formula):
        target_sheet = ref_sheet or sheet_name
        if ref_cell == cell and target_sheet == sheet_name:
            raise GenerationValidationError(f"{truth.workbook_id} {sheet_name}!{cell}: formula references itself")
        if target_sheet not in wb.sheetnames:
            raise GenerationValidationError(f"{truth.workbook_id} {sheet_name}!{cell}: missing reference sheet {target_sheet!r}")
        try:
            wb[target_sheet][ref_cell]
        except Exception as exc:
            raise GenerationValidationError(f"{truth.workbook_id} {sheet_name}!{cell}: invalid reference {target_sheet}!{ref_cell}") from exc
        checks += 1
    return checks


def validate_generated_workbook(dataset_root: Path, truth: WorkbookTruth) -> dict[str, int | str]:
    workbook_path = dataset_root / truth.relative_path
    wb = load_workbook(workbook_path, data_only=False)
    scenario = scenario_by_family(truth.family_id)
    formula_checks = 0
    arithmetic_checks = 0
    boundary_checks = 0
    reference_checks = 0
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.column >= 20 and cell.value is not None:
                        raise GenerationValidationError(f"{truth.workbook_id} {ws.title}!{cell.coordinate}: forbidden answer-marker/helper cell")
        sequence_by_sheet: dict[str, dict[int, float | int]] = {}
        for sheet in {spec.sheet for spec in truth.formula_cells}:
            ws = wb[sheet]
            prior: dict[int, float | int] = {}
            if scenario.sheet_mode != "horizontal":
                for row_number in range(scenario.start_row, scenario.start_row + scenario.row_count):
                    prior[row_number] = _business_value(scenario, row_number, ws, prior)
            sequence_by_sheet[sheet] = prior
        spec_lookup = _spec_by_cell(truth.formula_cells)
        intentional_target = _declared_mutation_target(truth, wb, spec_lookup)
        for spec in truth.formula_cells:
            ws = wb[spec.sheet]
            actual = ws[spec.cell].value
            if actual != spec.formula:
                raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: expected cell value {spec.formula!r}, got {actual!r}")
            formula_checks += 1
            if scenario.scenario_id == "payroll_overtime" and spec.role in {"detail", "normal_exception"} and intentional_target != (spec.sheet, spec.cell):
                _col, row = coordinate_from_string(spec.cell)
                if actual not in _payroll_allowed_formulas(int(row)):
                    raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: unsupported payroll anchor formula {actual!r}")
            if intentional_target == (spec.sheet, spec.cell):
                continue
            if isinstance(actual, str) and actual.startswith("="):
                reference_checks += _check_formula_references(truth, wb, spec.sheet, spec.cell, actual)
            if spec.semantic_value is not None:
                col, row = coordinate_from_string(spec.cell)
                row_number = int(row)
                if scenario.sheet_mode == "horizontal":
                    month_index = column_index_from_string(col) - 4
                    independent = _budget_month_value(wb, month_index, ws, col)
                else:
                    independent = sequence_by_sheet[spec.sheet][row_number]
                if not _same_number(independent, spec.semantic_value):
                    raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: recomputed business value {independent!r} != truth {spec.semantic_value!r}")
                arithmetic_checks += 1
            if isinstance(actual, str) and actual.startswith("=SUM("):
                inner = actual.removeprefix("=SUM(").removesuffix(")")
                try:
                    cells = _range_cells(inner)
                except ValueError as exc:
                    raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: invalid subtotal range {inner!r}") from exc
                if spec.cell in cells:
                    raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: total includes itself")
                expected_members = {candidate.cell for candidate in truth.formula_cells if candidate.sheet == spec.sheet and candidate.role in {"detail", "normal_exception"}}
                if not cells.issubset(expected_members):
                    raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: subtotal range includes non-detail cells")
                total = sum(float(spec_lookup[(spec.sheet, coordinate)].semantic_value or 0) for coordinate in cells)
                if total == 0 and cells:
                    raise GenerationValidationError(f"{truth.workbook_id} {spec.sheet}!{spec.cell}: subtotal truth is empty")
                boundary_checks += 1
    finally:
        wb.close()
    return {"formula_checks": formula_checks, "arithmetic_checks": arithmetic_checks, "boundary_checks": boundary_checks, "reference_checks": reference_checks, "answer_marker_cells": 0, "excel_recalculation": "NOT_RUN"}

