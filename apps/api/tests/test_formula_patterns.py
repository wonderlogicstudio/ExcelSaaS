from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.config import Settings
from app.formula_patterns import normalize_formula
from app.recommendation_engine import add_finding_guidance
from app.scanner import run_formula_audit, scan_workbook


def _scan_pattern_workbook(workbook: Workbook):
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return run_formula_audit(
        "formula-patterns.xlsx",
        stream.getvalue(),
        Settings(app_env="internal_beta", formula_pattern_audit_enabled=True),
    )


def _reload_workbook(workbook: Workbook) -> Workbook:
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)

    from openpyxl import load_workbook

    return load_workbook(stream)


def _add_formula_region(sheet, *, column: str = "C", outlier_formula: str | None = None) -> None:
    sheet["A1"] = "Label"
    sheet["B1"] = "Input"
    sheet[f"{column}1"] = "Calculated"
    for row in range(2, 7):
        sheet[f"A{row}"] = f"Item {row}"
        sheet[f"B{row}"] = row
        sheet[f"{column}{row}"] = f"=SUM(B{row})"
    if outlier_formula is not None:
        sheet[f"{column}4"] = outlier_formula


def _add_horizontal_region(sheet, *, row: int = 8, outlier_formula: str | None = None) -> None:
    sheet[f"C{row - 2}"] = "Input A"
    sheet[f"C{row - 1}"] = "Input B"
    sheet[f"C{row}"] = "Difference"
    for index, column in enumerate(("D", "E", "F", "G", "H"), start=1):
        sheet[f"{column}{row - 2}"] = index * 10
        sheet[f"{column}{row - 1}"] = index
        sheet[f"{column}{row}"] = f"={column}{row - 2}-{column}{row - 1}"
    sheet[f"F{row}"] = outlier_formula or f"=F{row - 2}-F{row - 1}"


def _ensure_monthly_source_sheets(workbook: Workbook) -> None:
    for index in range(1, 13):
        title = f"M{index:02d}"
        sheet = workbook[title] if title in workbook.sheetnames else workbook.create_sheet(title)
        sheet["B15"] = 100 + index
        sheet["B16"] = 95 + index


def _add_monthly_budget_region(
    sheet,
    *,
    target_formula: str = "=N15-N14",
    header_row: int = 14,
    formula_row: int = 18,
    first_column: int = 5,
    month_count: int = 12,
) -> None:
    for offset in range(month_count):
        column = first_column + offset
        month = f"M{offset + 1:02d}"
        sheet.cell(row=header_row, column=column).value = month
        sheet.cell(row=15, column=column).value = f"='{month}'!B15"
        sheet.cell(row=16, column=column).value = f"='{month}'!B16"
        sheet.cell(row=formula_row, column=column).value = f"='{month}'!B16-'{month}'!B15"
    sheet.cell(row=formula_row, column=14).value = target_formula


def test_detects_supported_normalized_formula_pattern_outliers() -> None:
    workbook = Workbook()
    function_change = workbook.active
    function_change.title = "Function"
    _add_formula_region(function_change, outlier_formula="=AVERAGE(B4)")

    sheet_change = workbook.create_sheet("SheetReference")
    _add_formula_region(sheet_change)
    for row in range(2, 7):
        sheet_change[f"C{row}"] = f"=SUM(SourceA!B{row})"
    sheet_change["C4"] = "=SUM(SourceB!B4)"

    relative_change = workbook.create_sheet("RelativeReference")
    _add_formula_region(relative_change, outlier_formula="=SUM(B3)")

    range_change = workbook.create_sheet("RangeEnd")
    _add_formula_region(range_change)
    for row in range(2, 7):
        range_change[f"C{row}"] = f"=SUM(B{row}:D{row})"
    range_change["C4"] = "=SUM(B4:C4)"

    result = _scan_pattern_workbook(workbook)
    candidates = [
        finding
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ]

    assert {(finding.sheet, finding.cell) for finding in candidates} == {
        ("Function", "C4"),
        ("SheetReference", "C4"),
        ("RelativeReference", "C4"),
        ("RangeEnd", "C4"),
    }
    expected_subtypes = {
        "Function": "FUNCTION_PATTERN_DRIFT",
        "SheetReference": "REFERENCE_SHEET_DRIFT",
        "RelativeReference": "RELATIVE_REFERENCE_DRIFT",
        "RangeEnd": "RANGE_BOUNDARY_DRIFT",
    }
    for finding in candidates:
        assert finding.formula_pattern is not None
        assert finding.formula_pattern.pattern_type == "DOMINANT_NORMALIZED_PATTERN_OUTLIER"
        assert finding.formula_pattern.pattern_subtype == expected_subtypes[finding.sheet]
        assert finding.formula_pattern.neighbor_count >= 3
        assert finding.formula_pattern.evidence_locations
        assert finding.formula_pattern.comparison_locations
        assert finding.formula_pattern.evidence_summary
        assert finding.formula_pattern.dominant_pattern_summary
        assert finding.formula_pattern.current_pattern_summary
        assert finding.formula_pattern.normal_case_possibility
        assert finding.formula_pattern.current_pattern_id is not None
        assert finding.repair_class == "EXPERT_REVIEW"


def test_detects_same_row_unanchored_arithmetic_reference_drift_from_run_formula_audit() -> None:
    for formula in ("=E6-F7", "=F5-F7"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Horizontal"
        _add_horizontal_region(sheet, outlier_formula=formula)
        sheet["F5"] = 100

        result = _scan_pattern_workbook(workbook)
        candidates = [
            finding
            for finding in result.candidates
            if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
        ]

        assert [(finding.sheet, finding.cell) for finding in candidates] == [
            ("Horizontal", "F8")
        ]
        evidence = candidates[0].formula_pattern
        assert evidence is not None
        assert evidence.pattern_subtype == "RELATIVE_REFERENCE_DRIFT"
        assert evidence.formula_region == "D8:H8"
        assert evidence.neighbor_count == 4
        assert evidence.comparison_locations == ["D8", "E8", "G8", "H8"]
        assert evidence.evidence_locations == ["D8", "E8", "G8", "H8"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Horizontal"
    _add_horizontal_region(sheet)

    result = _scan_pattern_workbook(workbook)

    assert [
        finding
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []


def test_detects_monthly_sheet_reference_drift_from_run_formula_audit() -> None:
    workbook = Workbook()
    budget = workbook.active
    budget.title = "Budget"
    _ensure_monthly_source_sheets(workbook)
    _add_monthly_budget_region(budget, target_formula="=N15-N14")

    result = _scan_pattern_workbook(workbook)
    candidates = [
        finding
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ]

    assert [(finding.sheet, finding.cell) for finding in candidates] == [("Budget", "N18")]
    evidence = candidates[0].formula_pattern
    assert evidence is not None
    assert evidence.pattern_subtype == "REFERENCE_SHEET_DRIFT"
    assert evidence.formula_region == "E18:P18"
    assert evidence.neighbor_count >= 3
    assert evidence.comparison_locations[:2] == ["M18", "O18"]


def test_monthly_sheet_reference_drift_preserves_equivalent_local_normal() -> None:
    workbook = Workbook()
    budget = workbook.active
    budget.title = "Budget"
    _ensure_monthly_source_sheets(workbook)
    _add_monthly_budget_region(budget, target_formula="=N16-N15")

    result = _scan_pattern_workbook(workbook)

    assert [
        (finding.sheet, finding.cell)
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []

    missing_sheet_workbook = Workbook()
    missing_sheet_budget = missing_sheet_workbook.active
    missing_sheet_budget.title = "MissingSheet"
    _ensure_monthly_source_sheets(missing_sheet_workbook)
    missing_sheet_workbook.remove(missing_sheet_workbook["M10"])
    _add_monthly_budget_region(missing_sheet_budget, target_formula="=N15-N14")
    assert [
        (finding.sheet, finding.cell)
        for finding in _scan_pattern_workbook(missing_sheet_workbook).candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []


def test_monthly_sheet_reference_drift_excludes_header_and_support_negatives() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _ensure_monthly_source_sheets(workbook)

    missing_header = workbook.create_sheet("MissingHeader")
    _add_monthly_budget_region(missing_header, target_formula="=N15-N14")
    missing_header["N14"] = None

    duplicate_header = workbook.create_sheet("DuplicateHeader")
    _add_monthly_budget_region(duplicate_header, target_formula="=N15-N14")
    duplicate_header["N14"] = "M09"

    reversed_header = workbook.create_sheet("ReversedHeader")
    _add_monthly_budget_region(reversed_header, target_formula="=N15-N14")
    reversed_header["M14"] = "M10"
    reversed_header["N14"] = "M09"

    three = workbook.create_sheet("ThreeFormulas")
    _add_monthly_budget_region(three, month_count=3)
    three["F18"] = "=F15-F14"

    two_deviations = workbook.create_sheet("TwoDeviations")
    _add_monthly_budget_region(two_deviations, month_count=5)
    two_deviations["G18"] = "=G15-G14"
    two_deviations["H18"] = "=H15-H14"

    edge = workbook.create_sheet("EdgeDeviation")
    _add_monthly_budget_region(edge, target_formula="='M05'!B16-'M05'!B15", month_count=5)
    edge["E18"] = "=E15-E14"

    mismatch = workbook.create_sheet("LeftRightMismatch")
    _add_monthly_budget_region(mismatch, month_count=5)
    mismatch["G18"] = "=G15-G14"
    mismatch["H18"] = "='M05'!B16-'M05'!B15"

    result = _scan_pattern_workbook(workbook)

    assert [
        (finding.sheet, finding.cell)
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []



def test_monthly_sheet_reference_drift_excludes_out_of_scope_monthly_majorities() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _ensure_monthly_source_sheets(workbook)

    unresolved_majority = workbook.create_sheet("UnresolvedLocalMajority")
    _add_monthly_budget_region(unresolved_majority, month_count=5)
    for column in range(5, 10):
        unresolved_majority.cell(row=18, column=column).value = (
            f"={unresolved_majority.cell(row=18, column=column).coordinate[0]}15-"
            f"{unresolved_majority.cell(row=18, column=column).coordinate[0]}14"
        )
    unresolved_majority["G18"] = "='M03'!B16-'M03'!B15"

    fixed_other_month_majority = workbook.create_sheet("FixedOtherMonthMajority")
    _add_monthly_budget_region(fixed_other_month_majority, month_count=5)
    for column in range(5, 10):
        fixed_other_month_majority.cell(row=18, column=column).value = "='M12'!B16-'M12'!B15"
    fixed_other_month_majority["G18"] = "='M03'!B16-'M03'!B15"

    address_drift = workbook.create_sheet("AddressDrift")
    _add_monthly_budget_region(address_drift, month_count=5)
    address_drift["G18"] = "='M03'!B15-'M03'!B14"

    result = _scan_pattern_workbook(workbook)

    assert [
        (finding.sheet, finding.cell)
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []

def test_monthly_sheet_reference_drift_excludes_structural_and_grammar_negatives() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _ensure_monthly_source_sheets(workbook)

    hidden_group = workbook.create_sheet("HiddenGroup")
    _add_monthly_budget_region(hidden_group, target_formula="=N15-N14")
    hidden_group.column_dimensions.group("M", "O", hidden=True)

    hidden_header = workbook.create_sheet("HiddenHeader")
    _add_monthly_budget_region(hidden_header, target_formula="=N15-N14")
    hidden_header.row_dimensions[14].hidden = True

    table_sheet = workbook.create_sheet("MonthlyTable")
    _add_monthly_budget_region(table_sheet, target_formula="=N15-N14")
    table = Table(displayName="MonthlyTableRef", ref="E14:P18")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table_sheet.add_table(table)

    merged = workbook.create_sheet("MonthlyMerged")
    _add_monthly_budget_region(merged, target_formula="=N15-N14")
    merged.merge_cells("N18:O18")

    summary = workbook.create_sheet("MonthlySummary")
    _add_monthly_budget_region(summary, target_formula="=N15-N14")
    summary["D18"] = "total"

    blank = workbook.create_sheet("BlankTarget")
    _add_monthly_budget_region(blank)
    blank["N18"] = None

    constant = workbook.create_sheet("ConstantTarget")
    _add_monthly_budget_region(constant)
    constant["N18"] = 0

    unsupported_formulas = {
        "FunctionRange": "=SUM(N15:N16)",
        "Anchored": "=$N$16-$N$15",
        "External": "='[other.xlsx]M10'!B16-'M10'!B15",
        "NameRef": "=NamedAmount-N15",
        "MixedMonthOperands": "='M10'!B16-'M09'!B15",
    }
    for sheet_name, formula in unsupported_formulas.items():
        sheet = workbook.create_sheet(sheet_name)
        _add_monthly_budget_region(sheet, target_formula=formula)

    result = _scan_pattern_workbook(workbook)

    assert [
        (finding.sheet, finding.cell)
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []


def test_horizontal_reference_drift_excludes_unsupported_and_structural_cases() -> None:
    workbook = Workbook()

    three = workbook.active
    three.title = "ThreeFormulas"
    for column in ("D", "E", "F"):
        three[f"{column}8"] = f"={column}6-{column}7"
    three["E8"] = "=D6-E7"

    edge = workbook.create_sheet("Edge")
    _add_horizontal_region(edge)
    edge["D8"] = "=C6-D7"

    tie = workbook.create_sheet("Tie")
    for column in ("D", "E"):
        tie[f"{column}8"] = f"={column}6-{column}7"
    tie["F8"] = "=E6-F7"
    tie["G8"] = "=F6-G7"

    neighbor_mismatch = workbook.create_sheet("NeighborMismatch")
    _add_horizontal_region(neighbor_mismatch, outlier_formula="=E6-F7")
    neighbor_mismatch["G8"] = "=F6-G7"

    blank = workbook.create_sheet("Blank")
    _add_horizontal_region(blank)
    blank["F8"] = None

    constant = workbook.create_sheet("Constant")
    _add_horizontal_region(constant)
    constant["F8"] = 0

    hidden_column = workbook.create_sheet("HiddenColumn")
    _add_horizontal_region(hidden_column, outlier_formula="=E6-F7")
    hidden_column.column_dimensions["F"].hidden = True

    hidden_span = workbook.create_sheet("HiddenSpan")
    _add_horizontal_region(hidden_span, outlier_formula="=E6-F7")
    hidden_span.column_dimensions["H"].hidden = True

    hidden_row = workbook.create_sheet("HiddenRow")
    _add_horizontal_region(hidden_row, outlier_formula="=E6-F7")
    hidden_row.row_dimensions[8].hidden = True

    table_sheet = workbook.create_sheet("Table")
    _add_horizontal_region(table_sheet, outlier_formula="=E6-F7")
    for column in ("C", "D", "E", "F", "G", "H"):
        table_sheet[f"{column}5"] = f"Header {column}"
    table = Table(displayName="HorizontalTable", ref="C5:H8")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table_sheet.add_table(table)

    merged = workbook.create_sheet("Merged")
    _add_horizontal_region(merged, outlier_formula="=E6-F7")
    merged.merge_cells("F8:G8")

    summary = workbook.create_sheet("Summary")
    _add_horizontal_region(summary, outlier_formula="=E6-F7")
    summary["C8"] = "total"

    unsupported_formulas = {
        "UnsupportedFunctionRange": "=SUM(F6:F7)",
        "UnsupportedAnchor": "=F$6-F7",
        "UnsupportedSheet": "=Other!F6-F7",
        "UnsupportedName": "=NamedAmount-F7",
        "UnsupportedExternal": "='[other.xlsx]Sheet1'!F6-F7",
    }
    for sheet_name, formula in unsupported_formulas.items():
        unsupported = workbook.create_sheet(sheet_name)
        _add_horizontal_region(unsupported, outlier_formula=formula)

    result = _scan_pattern_workbook(workbook)

    assert [
        (finding.sheet, finding.cell)
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == []


def test_horizontal_reference_drift_respects_grouped_hidden_column_ranges_after_reload() -> None:
    workbook = Workbook()
    grouped_hidden = workbook.active
    grouped_hidden.title = "GroupedHidden"
    _add_horizontal_region(grouped_hidden, outlier_formula="=E6-F7")
    grouped_hidden.column_dimensions.group("C", "F", hidden=True)

    visible_control = workbook.create_sheet("VisibleControl")
    _add_horizontal_region(visible_control, outlier_formula="=E6-F7")

    outside_group_control = workbook.create_sheet("OutsideGroupControl")
    _add_horizontal_region(outside_group_control, outlier_formula="=E6-F7")
    outside_group_control.column_dimensions.group("A", "C", hidden=True)

    reloaded = _reload_workbook(workbook)
    result = _scan_pattern_workbook(reloaded)

    assert [
        (finding.sheet, finding.cell)
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == [
        ("VisibleControl", "F8"),
        ("OutsideGroupControl", "F8"),
    ]


def test_same_rule_cell_is_not_duplicated_between_vertical_and_horizontal_detection() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dedupe"
    _add_horizontal_region(sheet, outlier_formula="=E6-F7")
    for row in (6, 7, 9, 10):
        sheet[f"F{row}"] = f"=F{row - 2}-F{row - 1}"

    result = _scan_pattern_workbook(workbook)

    assert [
        (finding.sheet, finding.cell, finding.rule_code)
        for finding in result.candidates
        if finding.sheet == "Dedupe"
        and finding.cell == "F8"
        and finding.rule_code == "FORMULA_PATTERN_OUTLIER"
    ] == [("Dedupe", "F8", "FORMULA_PATTERN_OUTLIER")]


def test_detects_constant_and_blank_inside_supported_formula_patterns() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gaps"
    _add_formula_region(sheet, column="C")
    sheet["C4"] = 999
    _add_formula_region(sheet, column="D")
    sheet["D4"] = None
    sheet["A4"] = "Item 4"

    result = _scan_pattern_workbook(workbook)
    candidates = [
        finding
        for finding in result.candidates
        if finding.rule_code == "FORMULA_PATTERN_GAP"
    ]

    assert {(finding.sheet, finding.cell) for finding in candidates} == {
        ("Gaps", "C4"),
        ("Gaps", "D4"),
    }
    assert all(finding.formula_pattern is not None for finding in candidates)
    assert {
        finding.cell: finding.formula_pattern.pattern_subtype
        for finding in candidates
        if finding.formula_pattern is not None
    } == {
        "C4": "CONSTANT_OVERRIDE_CANDIDATE",
        "D4": "BLANK_GAP_CANDIDATE",
    }


def test_excludes_headers_boundaries_summary_rows_blank_dividers_merged_and_tables() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Boundary"

    _add_formula_region(sheet, column="C")
    sheet["C6"] = "=SUM(C2:C5)"  # total at a region boundary

    summary_sheet = workbook.create_sheet("Summary")
    _add_formula_region(summary_sheet, outlier_formula="=AVERAGE(B4)")
    summary_sheet["A4"] = "소계"

    divider_sheet = workbook.create_sheet("BlankDivider")
    _add_formula_region(divider_sheet)
    divider_sheet["C4"] = None
    divider_sheet["A4"] = None
    divider_sheet["B4"] = None  # a fully blank divider for this formula region

    merged_sheet = workbook.create_sheet("Merged")
    _add_formula_region(merged_sheet, outlier_formula="=AVERAGE(B4)")
    merged_sheet.merge_cells("C4:D4")

    table_sheet = workbook.create_sheet("Table")
    _add_formula_region(table_sheet, outlier_formula="=AVERAGE(B4)")
    table = Table(displayName="FormulaTable", ref="A1:C6")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table_sheet.add_table(table)

    result = _scan_pattern_workbook(workbook)
    pattern_findings = [
        finding for finding in result.candidates if finding.formula_pattern is not None
    ]

    assert pattern_findings == []


def test_does_not_treat_a_summary_marker_inside_a_formula_reference_as_a_summary_row() -> None:
    workbook = Workbook()
    source = workbook.active
    source.title = "총계정원장"
    for row in range(2, 7):
        source[f"A{row}"] = f"AC-{row}"
        source[f"B{row}"] = row * 100

    sheet = workbook.create_sheet("대사표")
    sheet["A1"] = "Account"
    sheet["B1"] = "Balance"
    for row in range(2, 7):
        sheet[f"A{row}"] = f"AC-{row}"
        sheet[f"B{row}"] = (
            "=SUMIFS('총계정원장'!$B$2:$B$6,'총계정원장'!$A$2:$A$6,"
            f"A{row})"
        )
    sheet["B4"] = "=SUMIFS('총계정원장'!$B$2:$B$6,'총계정원장'!$A$2:$A$6,A3)"

    result = _scan_pattern_workbook(workbook)

    assert {(finding.sheet, finding.cell, finding.rule_code) for finding in result.candidates} == {
        ("대사표", "B4", "FORMULA_PATTERN_OUTLIER")
    }


def test_pattern_audit_is_default_off_and_emits_value_free_stable_evidence() -> None:
    workbook = Workbook()
    sheet = workbook.active
    _add_formula_region(sheet, outlier_formula="=AVERAGE(B4)")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    payload = stream.getvalue()

    disabled = scan_workbook("patterns.xlsx", payload, Settings())
    still_isolated = scan_workbook(
        "patterns.xlsx",
        payload,
        Settings(app_env="internal_beta", formula_pattern_audit_enabled=True),
    )
    first = run_formula_audit(
        "patterns.xlsx",
        payload,
        Settings(app_env="internal_beta", formula_pattern_audit_enabled=True),
    )
    second = run_formula_audit(
        "patterns.xlsx",
        payload,
        Settings(app_env="internal_beta", formula_pattern_audit_enabled=True),
    )

    assert not any(finding.formula_pattern is not None for finding in disabled.findings)
    assert not any(finding.formula_pattern is not None for finding in still_isolated.findings)
    first_candidates = [finding for finding in first.candidates if finding.formula_pattern]
    second_candidates = [finding for finding in second.candidates if finding.formula_pattern]
    assert [finding.finding_key for finding in first_candidates] == [
        finding.finding_key for finding in second_candidates
    ]
    payload_json = first.model_dump_json()
    assert "=AVERAGE(B4)" not in payload_json
    evidence_json = "".join(
        finding.formula_pattern.model_dump_json()
        for finding in first_candidates
        if finding.formula_pattern is not None
    )
    # Finding IDs are UUIDs and can coincidentally contain a numeric literal.
    # Check value-free evidence, not the entire envelope, for literal leakage.
    assert "999" not in evidence_json

    enriched = add_finding_guidance(first.candidates)
    candidate = next(finding for finding in enriched if finding.formula_pattern)
    assert candidate.guidance is not None
    assert candidate.guidance.evidence_grade == "PATTERN_INFERENCE"
    assert candidate.guidance.action_category == "DEEP_VALIDATION_REQUIRED"
    assert candidate.guidance.repair_eligibility == "CURRENTLY_NOT_SUPPORTED"


def test_pattern_audit_skips_truncated_scans() -> None:
    workbook = Workbook()
    sheet = workbook.active
    _add_formula_region(sheet, outlier_formula="=AVERAGE(B4)")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    result = run_formula_audit(
        "truncated-patterns.xlsx",
        stream.getvalue(),
        Settings(app_env="internal_beta", formula_pattern_audit_enabled=True, scan_cell_limit=2),
    )

    assert result.status == "SKIPPED_TRUNCATED"
    assert result.candidates == []


def test_normalizes_supported_literals_without_retaining_values() -> None:
    workbook = Workbook()
    sheet = workbook.active
    first = sheet["C2"]
    second = sheet["C3"]
    first.value = '=IF(A2="internal-one",B2*0.1,0)'
    second.value = '=IF(A3="internal-two",B3*0.9,9)'

    first_signature = normalize_formula(first.value, first)
    second_signature = normalize_formula(second.value, second)

    assert first_signature is not None
    assert first_signature == second_signature
    assert "internal-one" not in first_signature
    assert "0.1" not in first_signature
    assert normalize_formula("=SUM(B2,1)", first) is None
