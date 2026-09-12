from __future__ import annotations

import base64
import csv
import json
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter

from app.comparison_engine import compare_sources, read_source
from app.comparison_process import run_comparison
from app.config import Settings
from app.errors import WorkbookCareError

ROOT = Path(__file__).resolve().parents[3]
REFERENCE = json.loads(
    (ROOT / "samples/delivery-v3_2/comparison-reference.json").read_text(encoding="utf-8")
)
CASES = REFERENCE["edges"]["cases"]
BASE = REFERENCE["baseline"]
EXPECTED = REFERENCE["expected"]


def source_bytes(rows, kind="csv"):
    count = max((len(row.get("key_parts", [row.get("key")])) for row in rows), default=1)
    headers = ["거래번호" + str(i + 1) for i in range(count)] + ["금액"]
    values = [headers] + [[*row.get("key_parts", [row.get("key")]), row["amount"]] for row in rows]
    if kind == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerows(values)
        return stream.getvalue().encode("utf-8-sig")
    book = Workbook()
    book.active.title = "명세"
    for row in values:
        book.active.append(row)
        for cell in book.active[book.active.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def spec_for(a, b, kind="csv", tolerance="0"):
    count = max(len(row.get("key_parts", [row.get("key")])) for row in a + b)

    def mapping(rows):
        return {
            "sheet": "CSV" if kind == "csv" else "명세",
            "range": f"A1:{get_column_letter(count + 1)}{len(rows) + 1}",
            "header_row": 1,
            "key_columns": [get_column_letter(i + 1) for i in range(count)],
            "amount_column": get_column_letter(count + 1),
            "exclusions": [],
            "criteria": {
                "period": "2026-09 synthetic",
                "amount_meaning": "총액",
                "currency": "KRW",
                "unit": "KRW_WON",
            },
        }

    return {
        "A": mapping(a),
        "B": mapping(b),
        "policy": {
            "confirmed": True,
            "currency": "KRW",
            "unit": "KRW_WON",
            "period": "2026-09 synthetic",
            "amount_meaning": "총액",
            "normalization": "NONE",
            "tolerance_krw": tolerance,
        },
    }


def actual(a, b, kind="csv", tolerance="0"):
    sources = {
        side: read_source("synthetic." + kind, source_bytes(rows, kind), Settings())
        for side, rows in [("A", a), ("B", b)]
    }
    return compare_sources(sources, spec_for(a, b, kind, tolerance))


@pytest.mark.parametrize("case", CASES, ids=[c["id"] for c in CASES])
def test_fixed_v31_edges_through_actual_csv_parser(case):
    result = actual(case["A"], case["B"], tolerance=case["policy"]["tolerance_krw"])
    projected = [
        {
            "status": r["status"],
            "A": [v["data_row"] for v in r["A"]],
            "B": [v["data_row"] for v in r["B"]],
            "delta": r["delta"],
        }
        for r in result["records"]
    ]

    def key(row):
        return json.dumps(row, sort_keys=True)

    assert sorted(projected, key=key) == sorted(case["expected"], key=key)
    assert result["eligibility"] == case["expected_preflight"]


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_v31_baseline_actual_parser_and_all_source_rows(kind):
    result = actual(BASE["A"], BASE["B"], kind)
    for field, value in EXPECTED["summary"].items():
        assert result["summary"][field] == value
    assert sum(len(r["A"]) + len(r["B"]) for r in result["records"]) == 14
    for expected in EXPECTED["records"]:
        r = next(r for r in result["records"] if r["key"] == expected["key"])
        for k in ["key", "status", "reason", "amount_A", "amount_B", "delta"]:
            assert r[k] == expected[k]
        for side in ["A", "B"]:
            assert [v["data_row"] for v in r[side]] == expected[side]


def test_real_bounded_process_and_timeout_never_returns_partial(tmp_path, monkeypatch):
    message = {
        "action": "compare",
        "sources": {
            side: {
                "filename": "synthetic.csv",
                "file_base64": base64.b64encode(source_bytes(BASE[side])).decode(),
            }
            for side in ["A", "B"]
        },
        "spec": spec_for(BASE["A"], BASE["B"]),
    }
    result = run_comparison(message)
    assert result["summary"]["known_amount_totals"] == {"A": "47900", "B": "51000"}
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))
    with pytest.raises(WorkbookCareError) as e:
        run_comparison(message, timeout=0.001)
    assert e.value.code == "COMPARISON_TIMEOUT" and list(tmp_path.iterdir()) == []


def test_numeric_identifier_date_boolean_and_formula_are_not_guessed():
    a = [{"key": 123, "amount": "10"}, {"key": "valid", "amount": "10"}]
    b = [{"key": "123", "amount": "10"}, {"key": "valid", "amount": "10"}]
    result = actual(a, b, "xlsx")
    assert result["summary"]["counts"]["INPUT_ERROR"] == 1
    book = Workbook()
    book.active.title = "명세"
    book.active.append(["거래번호1", "금액"])
    book.active.append(["K", "=1+1"])
    stream = BytesIO()
    book.save(stream)
    sources = {
        "A": read_source("s.xlsx", stream.getvalue(), Settings()),
        "B": read_source("s.xlsx", source_bytes([{"key": "K", "amount": "2"}], "xlsx"), Settings()),
    }
    with pytest.raises(WorkbookCareError) as e:
        compare_sources(
            sources, spec_for([{"key": "K", "amount": "2"}], [{"key": "K", "amount": "2"}], "xlsx")
        )
    assert e.value.code == "COMPARISON_SELECTED_FORMULA"


def test_swap_order_invariance_and_explicit_exclusions():
    a, b = BASE["A"], BASE["B"]
    normal = actual(a, b)
    swapped = actual(b, a)
    assert swapped["summary"]["known_amount_totals"] == {"A": "51000", "B": "47900"}
    assert next(r for r in swapped["records"] if r["key"] == "0002")["delta"] == "2000"
    shuffled = actual(list(reversed(a)), list(reversed(b)))
    assert shuffled["summary"]["counts"] == normal["summary"]["counts"]
    sources = {
        side: read_source("s.csv", source_bytes(rows), Settings())
        for side, rows in [("A", a), ("B", b)]
    }
    spec = spec_for(a, b)
    spec["A"]["exclusions"] = [{"row": 7, "reason": "고객이 명시한 범위 제외"}]
    result = compare_sources(sources, spec)
    assert result["summary"]["excluded_rows"]["A"] == 1
    assert result["summary"]["assigned_rows"]["A"] == 7
    assert result["summary"]["excluded_known_amount_totals"]["A"] == "900"
