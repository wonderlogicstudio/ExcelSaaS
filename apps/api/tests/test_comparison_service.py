from __future__ import annotations

import base64
import copy
import hashlib
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_comparison_engine import BASE, actual, spec_for
from test_delivery_inputs import make_client

from app.comparison_artifacts import make_comparison_artifacts, validate_comparison_artifacts
from app.comparison_service import (
    create_comparison,
    download_comparison,
    execute_comparison,
    grant_comparison,
    prepare_comparison,
    project_comparison,
)
from app.config import Settings
from app.delivery_execution import cancel, execute
from app.delivery_store import DeliveryStore
from app.errors import WorkbookCareError

ROOT = Path(__file__).resolve().parents[3]


def inputs(case="baseline", kind="csv", key="comparison-request-0001"):
    return {
        "request_key": key,
        "consent": True,
        "sources": {
            side: {
                "filename": "source." + kind,
                "file_base64": base64.b64encode(
                    (ROOT / f"samples/delivery-v3_2/comparison-{case}-{side}.{kind}").read_bytes()
                ).decode(),
            }
            for side in ["A", "B"]
        },
    }


def prepared(tmp_path):
    store = DeliveryStore(tmp_path)
    settings = Settings(app_env="internal_beta")
    job = create_comparison(store, "owner", inputs(), settings)
    job = prepare_comparison(
        store,
        job,
        {
            "revision": job["revision"],
            "source_pair_hash": job["snapshot"]["source_pair_hash"],
            "spec": spec_for(BASE["A"], BASE["B"]),
        },
        settings,
    )
    return store, job, settings


def test_actual_reports_preserve_all_rows_and_literal_amounts(tmp_path):
    store, job, settings = prepared(tmp_path)
    assert project_comparison(job, settings)["result"] is None
    assert not project_comparison(job, settings)["purchase_enabled"]
    with pytest.raises(WorkbookCareError):
        execute_comparison(store, job, {}, settings)
    job = grant_comparison(store, job, settings)
    with pytest.raises(WorkbookCareError):
        execute_comparison(store, job, {}, settings)
    ready = execute_comparison(store, job, {"acknowledge_limitations": True}, settings)
    assert ready["state"]["status"] == "READY"
    assert ready["state"]["payment"] == "NOT_PURCHASED"
    assert ready["state"]["comparison_result"]["summary"]["counts"] == {
        "MATCHED": 2,
        "AMOUNT_DIFF": 1,
        "ONLY_A": 1,
        "ONLY_B": 1,
        "AMBIGUOUS": 1,
        "INPUT_ERROR": 2,
    }
    artifact = download_comparison(store, ready, "COMPARISON_REPORT_XLSX", settings)
    assert artifact == download_comparison(store, ready, "COMPARISON_REPORT_XLSX", settings)
    book = load_workbook(BytesIO(base64.b64decode(artifact["file_base64"])), data_only=False)
    rows = [
        row
        for name in ["금액차이", "한쪽자료", "중복모호", "자료오류", "일치"]
        for row in book[name].iter_rows(min_row=2, values_only=True)
        if row[3]
    ]
    assert len(rows) == 14 and len({row[3] for row in rows}) == 14
    assert sum(1 for r in rows if r[9] == "확인 불가") == 1
    assert {r[10] for r in rows if r[1] == "AMOUNT_DIFF"} == {"-2000"}
    assert book["요약"]["C12"].value == "47900"
    assert book["요약"]["D12"].value == "51000"
    for sheet in book:
        for row in sheet:
            assert all(c.data_type != "f" and c.hyperlink is None for c in row)
    book.close()
    assert (
        execute_comparison(store, ready, {}, settings)["state"]["delivery"]
        == ready["state"]["delivery"]
    )
    assert store.load("owner", job["id"])["source"] == base64.b64decode(
        inputs()["sources"]["A"]["file_base64"]
    )
    with pytest.raises(WorkbookCareError):
        execute(store, ready, settings)
    with pytest.raises(WorkbookCareError):
        download_comparison(store, ready, "REPAIRED_XLSX", settings)
    with store.connection() as db:
        db.execute("DELETE FROM delivery_artifacts WHERE kind='COMPARISON_VERIFICATION_HTML'")
    with pytest.raises(WorkbookCareError):
        download_comparison(store, ready, "COMPARISON_REPORT_XLSX", settings)
    store.delete("owner", job["id"])
    with store.connection() as db:
        assert db.execute("SELECT COUNT(*) FROM secondary_sources").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM delivery_artifacts").fetchone()[0] == 0


def test_owner_pair_idempotency_tampering_and_expired_right(tmp_path):
    store, job, settings = prepared(tmp_path)
    assert create_comparison(store, "owner", inputs(), settings)["id"] == job["id"]
    different = inputs()
    different["sources"]["B"] = inputs("equal")["sources"]["B"]
    with pytest.raises(WorkbookCareError):
        create_comparison(store, "owner", different, settings)
    with pytest.raises(WorkbookCareError):
        store.load("wrong-owner", job["id"])
    job = grant_comparison(store, job, settings)
    expired = copy.deepcopy(job)
    expired["state"]["comparison_grant"]["expires_at"] = time.time() - 1
    with pytest.raises(WorkbookCareError):
        execute_comparison(store, expired, {}, settings)
    for mutation in ["sku", "spec", "model"]:
        bad = copy.deepcopy(job)
        if mutation == "sku":
            bad["product"] = "APPROVED_REPAIR"
        if mutation == "spec":
            bad["state"]["spec_hash"] = "0" * 64
        if mutation == "model":
            bad["state"]["comparison_result"]["summary"]["group_count"] = 99
        with pytest.raises(WorkbookCareError):
            execute_comparison(store, bad, {}, settings)
    with store.connection() as db:
        db.execute("UPDATE secondary_sources SET source=?", (b"tampered",))
    with pytest.raises(WorkbookCareError):
        store.load("owner", job["id"])


@pytest.mark.parametrize("case", ["equal", "large"])
def test_zero_difference_and_large_integer_actual_report(case):
    rows = (
        ([{"key": "0001", "amount": 0}, {"key": "0002", "amount": 100}],) * 2
        if case == "equal"
        else (
            [{"key": "0001", "amount": "10000000000000001"}],
            [{"key": "0001", "amount": "10000000000000000"}],
        )
    )
    model = actual(*rows)
    package = make_comparison_artifacts(
        model, {"id": "synthetic-report", "expires": time.time() + 60}
    )
    validate_comparison_artifacts(package, model)
    book = load_workbook(BytesIO(package["artifacts"]["COMPARISON_REPORT_XLSX"]["data"]))
    if case == "equal":
        assert model["summary"]["counts"]["MATCHED"] == 2
        assert model["eligibility"] == "ELIGIBLE"
        assert book["금액차이"]["A2"].value.startswith("0건")
    else:
        assert book["금액차이"]["J2"].value == "10000000000000001"
        assert book["금액차이"]["K2"].value == "1"
    book.close()


def test_missing_artifact_and_cancel_publish_fence(tmp_path, monkeypatch):
    store, job, settings = prepared(tmp_path)
    job = grant_comparison(store, job, settings)

    def missing(*args):
        package = make_comparison_artifacts(*args)
        package["artifacts"].pop("COMPARISON_VERIFICATION_HTML")
        return package

    monkeypatch.setattr("app.comparison_service.make_comparison_artifacts", missing)
    with pytest.raises(WorkbookCareError):
        execute_comparison(store, job, {"acknowledge_limitations": True}, settings)
    job = store.load("owner", job["id"])
    assert job["state"]["status"] == "QUARANTINED"
    waiting = threading.Event()
    release = threading.Event()

    def pause(*args):
        result = make_comparison_artifacts(*args)
        waiting.set()
        assert release.wait(15)
        return result

    monkeypatch.setattr("app.comparison_service.make_comparison_artifacts", pause)
    with ThreadPoolExecutor() as pool:
        future = pool.submit(
            execute_comparison, store, job, {"acknowledge_limitations": True}, settings
        )
        assert waiting.wait(15)
        current = store.load("owner", job["id"])
        assert execute_comparison(store, current, {}, settings)["state"]["status"] == "RUNNING"
        assert cancel(store, current)["state"]["status"] == "CANCEL_REQUESTED"
        release.set()
        with pytest.raises(WorkbookCareError):
            future.result(15)
    assert store.load("owner", job["id"])["state"]["status"] == "CANCELLED"
    with store.connection() as db:
        assert db.execute("SELECT COUNT(*) FROM delivery_artifacts").fetchone()[0] == 0


def test_actual_api_separate_sku_free_preview_and_server_criteria(tmp_path, monkeypatch):
    settings = Settings(
        app_env="internal_beta", delivery_beta_enabled=True, delivery_data_dir=str(tmp_path)
    )
    client = make_client(settings, monkeypatch)
    headers = {"Origin": "http://localhost:5173", "X-WorkbookCare-CSRF": "1"}

    def call(body):
        return client.post("/v1/delivery", json=body, headers=headers)

    response = call({"action": "create_comparison", **inputs()})
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    job = response.json()
    request = {
        "action": "prepare_comparison",
        "job_id": job["job_id"],
        "revision": job["revision"],
        "source_pair_hash": job["source_pair_hash"],
        "spec": spec_for(BASE["A"], BASE["B"]),
    }
    bad = copy.deepcopy(request)
    bad["spec"]["B"]["criteria"]["period"] = "different"
    assert call(bad).status_code == 400
    response = call(request)
    assert response.status_code == 200
    assert response.json()["result"] is None
    assert response.json()["preflight"]["comparable_pair_count"] == 3
    for action in ["approve_plan", "execute", "download", "prepare_plan"]:
        assert call({"action": action, "job_id": job["job_id"]}).status_code == 403
    assert call({"action": "execute_comparison", "job_id": job["job_id"]}).status_code == 403
    from app.delivery_api import get_store

    store = get_store()
    owner = hashlib.sha256(client.cookies["workbookcare_local_session"].encode()).hexdigest()
    grant_comparison(store, store.load(owner, job["job_id"]), settings)
    ready = call(
        {"action": "execute_comparison", "job_id": job["job_id"], "acknowledge_limitations": True}
    )
    assert ready.status_code == 200 and ready.json()["status"] == "READY"
    for kind in ["COMPARISON_REPORT_XLSX", "COMPARISON_VERIFICATION_HTML"]:
        result = call({"action": "comparison_download", "job_id": job["job_id"], "kind": kind})
        assert result.status_code == 200 and result.headers["cache-control"] == "no-store"
    other = make_client(settings, monkeypatch)
    assert (
        other.post(
            "/v1/delivery", json={"action": "get", "job_id": job["job_id"]}, headers=headers
        ).status_code
        == 404
    )


@pytest.mark.parametrize("mutation", ["html_row", "html_total", "xlsx_summary"])
def test_report_reopen_detects_semantic_tampering_even_with_new_hash(mutation):
    model = actual(BASE["A"], BASE["B"])
    package = make_comparison_artifacts(model, {"id": "synthetic", "expires": time.time() + 60})
    kind = (
        "COMPARISON_REPORT_XLSX" if mutation == "xlsx_summary" else "COMPARISON_VERIFICATION_HTML"
    )
    data = package["artifacts"][kind]["data"]
    if mutation == "xlsx_summary":
        book = load_workbook(BytesIO(data))
        book["요약"]["C12"] = "0"
        stream = BytesIO()
        book.save(stream)
        book.close()
        data = stream.getvalue()
    elif mutation == "html_row":
        data = data.replace(b">-2000</td>", b">-2001</td>", 1)
    else:
        data = data.replace(b">47900</strong>", b">0</strong>", 1)
    package["artifacts"][kind]["data"] = data
    package["manifest"]["files"][kind]["sha256"] = hashlib.sha256(data).hexdigest()
    with pytest.raises(WorkbookCareError):
        validate_comparison_artifacts(package, model)
