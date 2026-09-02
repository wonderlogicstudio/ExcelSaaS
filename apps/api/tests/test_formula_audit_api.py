from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.main as main
from app.config import Settings
from app.scanner import run_formula_audit

client = TestClient(main.app)
ROOT = Path(__file__).resolve().parents[3]
SUPPLIED_PACK_ROOT = (
    ROOT
    / "samples"
    / "WorkbookCare_M4C_Sample_Pack_2026-09-01"
    / "WorkbookCare_M4C_Sample_Pack"
)


def _workbook_bytes(*, unsupported: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pattern"
    sheet["A1"] = "Item"
    sheet["B1"] = "Input"
    sheet["C1"] = "Calculated"
    for row in range(2, 7):
        sheet[f"A{row}"] = f"Item {row}"
        sheet[f"B{row}"] = row
        sheet[f"C{row}"] = "=SUM(Table1[Amount])" if unsupported else f"=SUM(B{row})"
    if not unsupported:
        sheet["C4"] = "=AVERAGE(B4)"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _file(payload: bytes) -> dict[str, tuple[str, bytes, str]]:
    return {
        "file": (
            "pattern-check.xlsx",
            payload,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


def _baseline_contract(payload: dict[str, object]) -> dict[str, object]:
    return {
        "workbook": payload["workbook"],
        "summary": payload["summary"],
        "quote": payload["quote"],
        "limitations": payload["limitations"],
        "findings": [
            {key: value for key, value in finding.items() if key != "id"}
            for finding in payload["findings"]
        ],
    }


def test_formula_audit_endpoint_fails_closed_by_default() -> None:
    response = client.post("/v1/formula-audits", files=_file(_workbook_bytes()))

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "FORMULA_AUDIT_NOT_AVAILABLE"


def test_formula_audit_is_separate_from_free_scan_contract(monkeypatch) -> None:
    monkeypatch.setattr(
        main,
        "settings",
        Settings(app_env="internal_beta", formula_pattern_audit_enabled=True),
    )
    payload = _workbook_bytes()

    before = client.post("/v1/scans", files=_file(payload))
    audit = client.post("/v1/formula-audits", files=_file(payload))
    after = client.post("/v1/scans", files=_file(payload))

    assert before.status_code == after.status_code == audit.status_code == 200
    assert _baseline_contract(before.json()) == _baseline_contract(after.json())
    assert not any(
        finding["rule_code"].startswith("FORMULA_PATTERN_")
        for finding in after.json()["findings"]
    )
    audit_payload = audit.json()
    assert audit_payload["status"] == "COMPLETED"
    assert {candidate["rule_code"] for candidate in audit_payload["candidates"]} == {
        "FORMULA_PATTERN_OUTLIER"
    }
    assert audit_payload["candidates"][0]["guidance"]["evidence_grade"] == "PATTERN_INFERENCE"
    assert "=AVERAGE(" not in audit.text
    assert "Item 4" not in audit.text


def test_formula_audit_returns_safe_non_candidate_outcomes() -> None:
    settings = Settings(app_env="internal_beta", formula_pattern_audit_enabled=True)
    unsupported = run_formula_audit("unsupported.xlsx", _workbook_bytes(unsupported=True), settings)
    assert unsupported.status == "SKIPPED_UNSUPPORTED_STRUCTURE"
    assert unsupported.candidates == []

    limited = run_formula_audit(
        "limited.xlsx",
        _workbook_bytes(),
        Settings(
            app_env="internal_beta",
            formula_pattern_audit_enabled=True,
            formula_audit_max_formula_cells=4,
        ),
    )
    assert limited.status == "SKIPPED_FORMULA_LIMIT"
    assert limited.candidates == []

    failed = run_formula_audit("not-a-workbook.txt", b"not an xlsx", settings)
    assert failed.status == "FAILED"
    assert failed.candidates == []


def test_formula_audit_uses_configured_workbook_and_candidate_limits() -> None:
    workbook = Workbook()
    for index in range(2):
        if index:
            workbook.create_sheet(f"Sheet{index + 1}")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    sheet_limited = run_formula_audit(
        "sheet-limit.xlsx",
        stream.getvalue(),
        Settings(
            app_env="internal_beta",
            formula_pattern_audit_enabled=True,
            formula_audit_max_sheet_count=1,
        ),
    )
    assert sheet_limited.status == "SKIPPED_WORKBOOK_LIMIT"
    assert sheet_limited.candidates == []

    workbook = Workbook()
    sheet = workbook.active
    for column in ("C", "D"):
        for row in range(2, 7):
            sheet[f"A{row}"] = f"Item {row}"
            sheet[f"B{row}"] = row
            sheet[f"{column}{row}"] = f"=SUM(B{row})"
        sheet[f"{column}4"] = "=AVERAGE(B4)"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    candidate_limited = run_formula_audit(
        "candidate-limit.xlsx",
        stream.getvalue(),
        Settings(
            app_env="internal_beta",
            formula_pattern_audit_enabled=True,
            formula_audit_max_candidate_count=1,
        ),
    )
    assert candidate_limited.status == "SKIPPED_CANDIDATE_LIMIT"
    assert candidate_limited.candidates == []


def test_formula_audit_endpoint_matches_the_supplied_m4c_sample_pack(monkeypatch) -> None:
    """Exercise the endpoint payload that the internal-beta UI receives.

    The free scan deliberately reports a different rule set. This check uses
    the exact supplied pack rather than the convenient repository copy so an
    uploaded M4-C sample cannot silently drift away from the acceptance labels.
    """

    monkeypatch.setattr(
        main,
        "settings",
        Settings(
            app_env="internal_beta",
            formula_pattern_audit_enabled=True,
            scan_cell_limit=250_000,
            finding_limit=5_000,
        ),
    )
    manifest = json.loads(
        (SUPPLIED_PACK_ROOT / "expected" / "m4c_manifest.json").read_text(encoding="utf-8")
    )
    expected = {
        (
            item["file"],
            item["sheet"],
            item["cell"],
            item["expected_rule"],
            item["expected_subtype"],
        )
        for item in manifest["expected_findings"]
    }
    actual: set[tuple[str, str | None, str | None, str, str | None]] = set()

    for scenario in manifest["scenarios"]:
        filename = scenario["file"]
        payload = (SUPPLIED_PACK_ROOT / "samples" / filename).read_bytes()
        response = client.post("/v1/formula-audits", files=_file(payload))
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "COMPLETED"
        for candidate in body["candidates"]:
            actual.add(
                (
                    filename,
                    candidate["sheet"],
                    candidate["cell"],
                    candidate["rule_code"],
                    candidate["formula_pattern"]["pattern_subtype"],
                )
            )

    assert actual == expected
