from __future__ import annotations

import base64
import hashlib
import importlib.util
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.config import Settings
from app.control_plane import ControlPlaneHmacMiddleware, build_control_plane_signature
from app.delivery_api import router
from app.delivery_inputs import (
    PROFILE_1,
    PROFILE_2,
    PROFILE_3,
    PROFILE_COMBINED,
    inspect_input,
    numeric_text,
    preflight,
)
from app.delivery_store import DeliveryStore
from app.errors import WorkbookCareError

ROOT = Path(__file__).resolve().parents[3]
spec = importlib.util.spec_from_file_location(
    "synthetic_delivery", ROOT / "scripts/generate_delivery_fixtures.py"
)
fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture)


def policy(profile=PROFILE_1, targets=None, **changes):
    return {
        "profile": profile,
        "sheet": "검증",
        "targets": targets or ["B2", "B3"],
        "role": "AMOUNT",
        "confirmed": True,
        "anchor": "F2",
        "anchor_formula": "=ROUND(C2*D2*(1-E2),0)",
        **changes,
    }


def monthly_workbook_bytes(*, target_formula="=N15-N14", extra=None):
    workbook = Workbook()
    budget = workbook.active
    budget.title = "Budget"
    for month in range(8, 13):
        sheet = workbook.create_sheet(f"M{month:02d}")
        sheet["B15"] = 1000 + month
        sheet["B16"] = 995 + month
    workbook["M10"]["B15"] = 1006
    workbook["M10"]["B16"] = 1001
    for offset, column in enumerate(range(12, 17), start=8):
        month = f"M{offset:02d}"
        budget.cell(row=14, column=column).value = month
        budget.cell(row=15, column=column).value = f"='{month}'!B15"
        budget.cell(row=16, column=column).value = f"='{month}'!B16"
        budget.cell(row=18, column=column).value = f"='{month}'!B16-'{month}'!B15"
    budget["N18"] = target_formula
    if extra:
        extra(workbook)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    raw = stream.getvalue()
    cleaned = BytesIO()
    with ZipFile(BytesIO(raw), "r") as source, ZipFile(
        cleaned, "w", compression=ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/workbook.xml":
                data = data.replace(b"<workbookProtection/>", b"")
            target.writestr(info, data)
    return cleaned.getvalue()


def monthly_policy(**changes):
    return {
        "profile": PROFILE_3,
        "sheet": "Budget",
        "targets": ["N18"],
        "before_formula": "=N15-N14",
        "confirmed": True,
        **changes,
    }


@pytest.mark.parametrize(
    "value",
    [
        "00123",
        "비용12원",
        "2026-09-13",
        "",
        " ",
        None,
        "1234567890123456",
        "-0",
        "1.2",
        "1e3",
        "12%",
        " 1200",
        "1,00",
        "+0",
    ],
)
def test_numeric_text_rejects_ambiguous_values(value):
    assert numeric_text(value) is None


def test_preflight_is_not_quote_or_approval():
    raw = fixture.workbook_bytes()
    s = inspect_input("synthetic.xlsx", raw, Settings())
    result = preflight(s, policy())
    assert result["eligible_count"] == 2 and result["status"] == "PRELIMINARY_ONLY"
    assert not result["quote_enabled"] and not result["purchase_enabled"]
    assert result["source_hash"] == hashlib.sha256(raw).hexdigest()
    for cell in ["A2", "B4", "B5", "B6", "B7", "B8", "B9", "B10", "B11", "B12"]:
        assert preflight(s, policy(targets=[cell]))["eligible_count"] == 0
    assert preflight(s, policy(role="ID"))["status"] == "UNSUPPORTED"
    assert preflight(s, policy(confirmed=False))["status"] == "UNSUPPORTED"


def test_true_blank_and_unsupported_function_are_distinct():
    s = inspect_input("synthetic.xlsx", fixture.workbook_bytes(), Settings())
    assert preflight(s, policy(PROFILE_2, ["F3"]))["eligible_count"] == 1
    for cell in ["G3", "H3", "I3", "J3"]:
        assert preflight(s, policy(PROFILE_2, [cell]))["eligible_count"] == 0
    unsupported = inspect_input(
        "synthetic.xlsx", fixture.workbook_bytes(unsupported=True), Settings()
    )
    assert "UNSUPPORTED_FORMULA" in preflight(unsupported, policy())["reason_codes"]


def test_monthly_profile_inspection_is_profile_specific_and_single_target():
    raw = monthly_workbook_bytes()
    default_snapshot = inspect_input("monthly.xlsx", raw, Settings())
    assert "UNSUPPORTED_FORMULA" in default_snapshot["issues"]

    monthly_snapshot = inspect_input(
        "monthly.xlsx", raw, Settings(), profile_context=PROFILE_3
    )
    assert "UNSUPPORTED_FORMULA" not in monthly_snapshot["issues"]
    result = preflight(monthly_snapshot, monthly_policy())
    assert result["status"] == "PRELIMINARY_ONLY"
    assert result["eligible_count"] == 1
    assert result["targets"][0]["change_kind"] == "MONTHLY_FORMULA_REPLACEMENT"

    assert preflight(monthly_snapshot, monthly_policy(targets=["N18", "O18"]))[
        "status"
    ] == "UNSUPPORTED"
    with pytest.raises(WorkbookCareError):
        preflight(
            monthly_snapshot,
            {"profile": PROFILE_COMBINED, "items": [monthly_policy()]},
        )


def test_monthly_profile_inspection_keeps_unsupported_formula_rejection():
    external = inspect_input(
        "monthly.xlsx",
        monthly_workbook_bytes(target_formula="='[other.xlsx]M10'!B16-'M10'!B15"),
        Settings(),
        profile_context=PROFILE_3,
    )
    assert "UNSUPPORTED_FORMULA" in external["issues"]

    unsupported_function = inspect_input(
        "monthly.xlsx",
        monthly_workbook_bytes(target_formula="=OFFSET(N15,0,0)"),
        Settings(),
        profile_context=PROFILE_3,
    )
    assert "UNSUPPORTED_FORMULA" in unsupported_function["issues"]


def test_owner_integrity_expiry_and_concurrent_updates(tmp_path):
    store = DeliveryStore(tmp_path)
    raw = fixture.workbook_bytes()
    snap = inspect_input("synthetic.xlsx", raw, Settings())
    job = store.create("owner-a", raw, snap, "request-0000000001")
    assert store.create("owner-a", raw, snap, "request-0000000001")["id"] == job["id"]
    with pytest.raises(WorkbookCareError, match="작업이 없거나"):
        store.load("owner-b", job["id"])

    def update(_):
        try:
            return store.update(job, 1, {**job["state"], "status": "UPDATED"})["revision"]
        except WorkbookCareError as error:
            return error.code

    with ThreadPoolExecutor(2) as pool:
        assert sorted(map(str, pool.map(update, range(2)))) == ["2", "STALE_JOB"]
    with store.connection() as db:
        db.execute("UPDATE jobs SET source=? WHERE id=?", (b"changed", job["id"]))
    with pytest.raises(WorkbookCareError, match="무결성"):
        store.load("owner-a", job["id"])
    assert store.cleanup(now=job["expires"] + 1) == 1


def make_client(settings, monkeypatch):
    monkeypatch.setattr("app.delivery_api.get_settings", lambda: settings)
    app = FastAPI()
    app.include_router(router)
    app.add_middleware(ControlPlaneHmacMiddleware, settings=settings)
    from app.main import workbookcare_error_handler

    app.add_exception_handler(WorkbookCareError, workbookcare_error_handler)
    return TestClient(app)


def test_actual_api_owner_cookie_preflight_and_negative_version(tmp_path, monkeypatch):
    settings = Settings(delivery_beta_enabled=True, delivery_data_dir=str(tmp_path))
    client = make_client(settings, monkeypatch)
    headers = {"Origin": "http://localhost:5173", "X-WorkbookCare-CSRF": "1"}
    assert client.post("/v1/delivery", json={"action": "capabilities"}).status_code == 403
    result = client.post(
        "/v1/delivery",
        headers=headers,
        json={
            "action": "create_input",
            "filename": "synthetic.xlsx",
            "file_base64": base64.b64encode(fixture.workbook_bytes()).decode(),
            "consent": True,
            "request_key": "fixture-request-0001",
        },
    )
    assert result.status_code == 200 and result.headers["cache-control"] == "no-store"
    job = result.json()
    assert "cells" not in job and "owner" not in job
    request = {
        "action": "preflight",
        "job_id": job["job_id"],
        "revision": 1,
        "source_hash": job["source_hash"],
        "policy": policy(),
    }
    checked = client.post("/v1/delivery", headers=headers, json=request)
    assert checked.status_code == 200 and checked.json()["preflight"]["eligible_count"] == 2
    assert client.post("/v1/delivery", headers=headers, json=request).status_code == 409
    other = make_client(settings, monkeypatch)
    receipt_request = {"action": "verify_approval_receipt", "job_id": job["job_id"]}
    assert other.post("/v1/delivery", headers=headers, json=receipt_request).status_code == 404
    assert client.post("/v1/delivery", json=receipt_request).status_code == 403
    assert client.post("/v1/delivery", headers=headers, json=receipt_request).status_code == 409
    assert (
        other.post(
            "/v1/delivery", headers=headers, json={"action": "get", "job_id": job["job_id"]}
        ).status_code
        == 404
    )


def test_api_monthly_preflight_reinspects_default_upload_with_profile_context(
    tmp_path, monkeypatch
):
    settings = Settings(delivery_beta_enabled=True, delivery_data_dir=str(tmp_path))
    client = make_client(settings, monkeypatch)
    headers = {"Origin": "http://localhost:5173", "X-WorkbookCare-CSRF": "1"}
    raw = monthly_workbook_bytes()
    created = client.post(
        "/v1/delivery",
        headers=headers,
        json={
            "action": "create_input",
            "filename": "monthly.xlsx",
            "file_base64": base64.b64encode(raw).decode(),
            "consent": True,
            "request_key": "monthly-api-request-0001",
        },
    )
    assert created.status_code == 200
    job = created.json()
    checked = client.post(
        "/v1/delivery",
        headers=headers,
        json={
            "action": "preflight",
            "job_id": job["job_id"],
            "revision": 1,
            "source_hash": job["source_hash"],
            "policy": monthly_policy(),
        },
    )

    assert checked.status_code == 200
    payload = checked.json()
    assert payload["preflight"]["status"] == "PRELIMINARY_ONLY"
    assert payload["preflight"]["eligible_count"] == 1


def test_hosted_owner_is_bound_to_body_signature(tmp_path, monkeypatch):
    import json
    import time

    secret = "synthetic-only-proof-secret-00000000000000"
    settings = Settings(
        app_env="hosted_beta",
        cors_origins=["https://beta.example"],
        control_plane_hmac_secret=secret,
        delivery_beta_enabled=True,
        delivery_data_dir=str(tmp_path),
    )
    client = make_client(settings, monkeypatch)
    body = json.dumps({"action": "capabilities"}).encode()
    now = int(time.time())
    owner = "a" * 64
    headers = {
        "Content-Type": "application/json",
        "Origin": "https://beta.example",
        "X-WorkbookCare-CSRF": "1",
        "x-workbookcare-owner": owner,
        "x-workbookcare-timestamp": str(now),
    }
    headers["x-workbookcare-signature"] = build_control_plane_signature(
        secret, now, "POST", "/v1/delivery", body, owner
    )
    assert client.post("/v1/delivery", content=body, headers=headers).status_code == 200
    headers["x-workbookcare-owner"] = "b" * 64
    assert client.post("/v1/delivery", content=body, headers=headers).status_code == 403
    headers["x-workbookcare-signature"] = build_control_plane_signature(
        secret, now, "POST", "/v1/delivery", body
    )
    assert client.post("/v1/delivery", content=body, headers=headers).status_code == 403


def test_mixed_selection_is_not_a_valid_whole_plan():
    snapshot = inspect_input("synthetic.xlsx", fixture.workbook_bytes(), Settings())
    result = preflight(snapshot, policy(targets=["B2", "A2"]))
    assert result["eligible_count"] == 1
    assert result["status"] == "UNSUPPORTED"
    assert "INELIGIBLE_TARGETS" in result["reason_codes"]


def test_delivery_origin_and_environment_fail_closed(tmp_path, monkeypatch):
    settings = Settings(delivery_beta_enabled=True, delivery_data_dir=str(tmp_path))
    client = make_client(settings, monkeypatch)
    assert client.post("/v1/delivery", json={"action": "capabilities"}).status_code == 403
    settings = settings.model_copy(update={"app_env": "production"})
    # Exercise the route's production gate independently; the hosted HMAC gate
    # has its own signed-owner tests and also rejects unsigned requests.
    monkeypatch.setattr("app.delivery_api.get_settings", lambda: settings)
    application = FastAPI()
    application.include_router(router)
    from app.main import workbookcare_error_handler
    application.add_exception_handler(WorkbookCareError, workbookcare_error_handler)
    client = TestClient(application)
    assert client.post("/v1/delivery", json={"action": "capabilities"}).status_code == 404
