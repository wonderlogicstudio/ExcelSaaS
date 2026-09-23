from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from app.repair_rules.monthly_formula import (
    PROFILE_MONTHLY,
    monthly_formula_replacement,
)

REPO_ROOT = Path(__file__).resolve().parents[3]
ARTIFACT_DIR = REPO_ROOT / "artifacts/synthetic_validation/monthly-repair-flow06/stage1"


def _monthly_control(
    *,
    target_formula: str = "=N15-N14",
    target_cell: str = "N18",
    target_cached: str | None = None,
    hide_month_columns: bool = False,
    first_source_type: str = "number",
    omit_first_source: bool = False,
) -> tuple[Workbook, dict]:
    workbook = Workbook()
    budget = workbook.active
    budget.title = "Budget"
    for month in range(8, 13):
        title = f"M{month:02d}"
        sheet = workbook.create_sheet(title)
        sheet["B15"] = 1000 + month
        sheet["B16"] = 995 + month
    workbook["M10"]["B15"] = 1006
    workbook["M10"]["B16"] = 1001

    for offset, column in enumerate(range(12, 17), start=8):
        month = f"M{offset:02d}"
        budget.cell(row=14, column=column).value = month
        budget.cell(row=15, column=column).value = f"='{month}'!B15"
        budget.cell(row=16, column=column).value = f"='{month}'!B16"
        budget.cell(row=18, column=column).value = f"='{month}'!B16-'{month}'!B15"
    budget[target_cell] = target_formula
    if hide_month_columns:
        budget.column_dimensions.group("M", "O", hidden=True)

    snapshot_cells = {
        "Budget": {
            target_cell: {"type": "formula", "value": target_formula},
        }
    }
    if target_cached is not None:
        snapshot_cells["Budget"][target_cell]["cached"] = target_cached
    for month in range(8, 13):
        title = f"M{month:02d}"
        snapshot_cells[title] = {
            "B15": {"type": "number", "value": str(workbook[title]["B15"].value)},
            "B16": {"type": "number", "value": str(workbook[title]["B16"].value)},
        }
    if first_source_type == "text":
        snapshot_cells["M10"]["B16"] = {"type": "text", "value": "1001"}
    elif first_source_type == "nonfinite":
        snapshot_cells["M10"]["B16"] = {"type": "number", "value": "Infinity"}
    if omit_first_source:
        del snapshot_cells["M10"]["B16"]
    return workbook, {"cells": snapshot_cells}


def _before_value(value: str = "#VALUE!") -> dict:
    return {"type": "error", "value": value, "provenance": "ENGINE_CALCULATED"}


def _before_number(value: int | str = 0) -> dict:
    return {"type": "number", "value": value, "provenance": "ENGINE_CALCULATED"}


def _replacement(workbook: Workbook, snapshot: dict, before_result: dict | None = None):
    return monthly_formula_replacement(
        workbook,
        snapshot,
        before_result or _before_value(),
        sheet="Budget",
        cell="N18",
    )


def test_monthly_formula_replacement_derives_exact_candidate_from_neighbors() -> None:
    workbook, snapshot = _monthly_control()
    candidate = _replacement(workbook, snapshot)

    assert candidate is not None
    assert candidate.profile == PROFILE_MONTHLY
    assert candidate.sheet == "Budget"
    assert candidate.cell == "N18"
    assert candidate.before_formula == "=N15-N14"
    assert candidate.before_value == {"type": "error", "value": "#VALUE!"}
    assert candidate.after_formula == "='M10'!B16-'M10'!B15"
    assert candidate.after_value == {"type": "number", "value": -5}
    assert candidate.evidence_cells == ("M18", "O18")
    assert candidate.approval_required is True
    assert candidate.automatic_approval is False
    assert len(workbook.sheetnames) == 6


def test_monthly_formula_replacement_uses_engine_before_result_not_cache_truth() -> None:
    for cached in (None, "#DIV/0!", "123"):
        workbook, snapshot = _monthly_control(target_cached=cached)
        assert _replacement(workbook, snapshot) is not None

    for before_result in [
        _before_number(-5),
        _before_value("#DIV/0!"),
        {"type": "error", "value": "#VALUE!", "provenance": "SOURCE_VALUE"},
    ]:
        workbook, snapshot = _monthly_control(target_cached="#VALUE!")
        assert _replacement(workbook, snapshot, before_result) is None


def test_monthly_formula_replacement_preserves_normal_equivalent_hidden() -> None:
    for workbook, snapshot, cell in [
        _monthly_control(target_formula="='M10'!B16-'M10'!B15") + ("N18",),
        _monthly_control(target_formula="=N16-N15") + ("N18",),
        _monthly_control(hide_month_columns=True) + ("N18",),
    ]:
        assert (
            monthly_formula_replacement(
                workbook, snapshot, _before_value(), sheet="Budget", cell=cell
            )
            is None
        )

    workbook, snapshot = _monthly_control()
    workbook["Budget"]["M18"] = "=M15-M14"
    assert _replacement(workbook, snapshot) is None

    edge_workbook, edge_snapshot = _monthly_control(target_formula="=L15-L14", target_cell="L18")
    assert (
        monthly_formula_replacement(
            edge_workbook, edge_snapshot, _before_value(), sheet="Budget", cell="L18"
        )
        is None
    )


def test_monthly_formula_replacement_requires_strict_monthly_run_not_generic_drift() -> None:
    workbook, snapshot = _monthly_control()
    budget = workbook["Budget"]
    budget["L18"] = None
    budget["P18"] = None
    budget["N17"] = "='M10'!B15-'M10'!B14"
    budget["N19"] = "='M10'!B17-'M10'!B16"
    budget["N20"] = "='M10'!B18-'M10'!B17"
    for cell, value in {"B14": 1007, "B17": 1000, "B18": 999}.items():
        workbook["M10"][cell] = value
        snapshot["cells"]["M10"][cell] = {"type": "number", "value": str(value)}

    assert _replacement(workbook, snapshot) is None

    duplicate_workbook, duplicate_snapshot = _monthly_control()
    duplicate_workbook["Budget"]["O14"] = "M10"
    duplicate_workbook["Budget"]["O18"] = "='M10'!B16-'M10'!B15"
    assert _replacement(duplicate_workbook, duplicate_snapshot) is None


def test_monthly_formula_replacement_rejects_bad_formulas_and_inputs() -> None:
    for target_formula in [
        "='M10'!B16-'M09'!B15",
        "='[other.xlsx]M10'!B16-'M10'!B15",
        "=$N$15-$N$14",
        "=SUM(N15:N16)",
    ]:
        workbook, snapshot = _monthly_control(target_formula=target_formula)
        assert _replacement(workbook, snapshot) is None

    for workbook, snapshot in [
        _monthly_control(first_source_type="text"),
        _monthly_control(first_source_type="nonfinite"),
        _monthly_control(omit_first_source=True),
    ]:
        assert _replacement(workbook, snapshot) is None

    workbook, snapshot = _monthly_control()
    snapshot["cells"]["M10"]["B16"] = {"type": "number", "value": "1e308"}
    snapshot["cells"]["M10"]["B15"] = {"type": "number", "value": "-1e308"}
    assert _replacement(workbook, snapshot) is None


def test_stage1_expected_control_fixture_is_frozen() -> None:
    expected = json.loads((ARTIFACT_DIR / "expected-control.json").read_text(encoding="utf-8"))

    assert expected["status"] == "FROZEN_INDEPENDENT_EXPECTATION"
    assert expected["sheet_count"] == 6
    assert expected["expected_candidate"] == {
        "profile": PROFILE_MONTHLY,
        "sheet": "Budget",
        "cell": "N18",
        "before_formula": "=N15-N14",
        "before_value": {"type": "error", "value": "#VALUE!"},
        "after_formula": "='M10'!B16-'M10'!B15",
        "after_value": {"type": "number", "value": -5},
        "derivation": "M10!B16=1001 minus M10!B15=1006",
        "approval_required": True,
        "automatic_approval": False,
    }
