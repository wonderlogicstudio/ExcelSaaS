from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook


@pytest.fixture
def risky_workbook_bytes() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = 100
    summary["A2"] = "=#REF!+1"
    summary["A3"] = "=SUM(A:A)"
    summary["A4"] = '=INDIRECT("A1")'
    summary["A5"] = "=IF(A1>0,IF(A1>1,IF(A1>2,IF(A1>3,IF(A1>4,1,0),0),0),0),0)"
    summary["A6"] = "='[Budget.xlsx]Sheet1'!A1"
    summary["B1"] = "Amount"
    summary["B2"] = 1000
    summary["B3"] = 1200
    summary["B4"] = 1400
    summary["B5"] = 1600
    summary["B6"] = "1,800"

    hidden = workbook.create_sheet("HiddenInput")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "internal"

    very_hidden = workbook.create_sheet("SystemCache")
    very_hidden.sheet_state = "veryHidden"
    very_hidden["A1"] = "internal"

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()
