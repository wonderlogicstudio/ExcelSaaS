from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from time import perf_counter
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from .config import Settings
from .errors import WorkbookCareError
from .formula_patterns import inspect_worksheet_formula_patterns
from .models import (
    DiagnosisSummary,
    Finding,
    FormulaAuditResult,
    FormulaPatternEvidence,
    ScanResult,
    WorkbookSummary,
)
from .pricing import build_quote
from .security import validate_ooxml

SCANNER_VERSION = "0.1.3"
RULE_SET_VERSION = "2026.09.4"
# M4 remains a separate optional audit. Advancing its rule-set identifier must
# not make a same-session M2.5 re-validation look like the base scanner rules
# changed when the base scan itself is unchanged.
FORMULA_AUDIT_RULE_SET_VERSION = "2026.09.5"

ERROR_TOKENS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NULL!")
VOLATILE_PATTERN = re.compile(
    r"\b(?:INDIRECT|OFFSET|NOW|TODAY|RAND|RANDBETWEEN|CELL|INFO)\s*\(",
    re.IGNORECASE,
)
WHOLE_COLUMN_PATTERN = re.compile(r"(?<![A-Z0-9_])\$?[A-Z]{1,3}:\$?[A-Z]{1,3}(?![A-Z0-9_])")
EXTERNAL_FORMULA_PATTERN = re.compile(
    r"\[[^\]]+\.(?:xlsx?|xlsm|xlsb|xlam|xltx|xltm|csv)\]",
    re.IGNORECASE,
)
NUMERIC_TEXT_PATTERN = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
FORMULA_AUDIT_ALLOWED_RULE_CODES = {
    "FORMULA_PATTERN_OUTLIER",
    "FORMULA_PATTERN_GAP",
}


@dataclass(slots=True)
class FindingCollector:
    limit: int
    findings: list[Finding]
    severity_counts: Counter[str]
    repair_counts: Counter[str]
    total: int = 0

    @classmethod
    def create(cls, limit: int) -> FindingCollector:
        return cls(limit=limit, findings=[], severity_counts=Counter(), repair_counts=Counter())

    def add(
        self,
        *,
        rule_code: str,
        severity: str,
        title: str,
        description: str,
        repair_class: str,
        sheet: str | None = None,
        cell: str | None = None,
        confidence: float = 1.0,
        formula_pattern: FormulaPatternEvidence | None = None,
    ) -> None:
        self.total += 1
        self.severity_counts[severity] += 1
        self.repair_counts[repair_class] += 1
        if len(self.findings) >= self.limit:
            return
        self.findings.append(
            Finding(
                id=str(uuid4()),
                finding_key=_finding_key(rule_code, sheet, cell),
                rule_code=rule_code,
                severity=severity,  # type: ignore[arg-type]
                title=title,
                description=description,
                sheet=sheet,
                cell=cell,
                confidence=confidence,
                repair_class=repair_class,  # type: ignore[arg-type]
                formula_pattern=formula_pattern,
            )
        )


def _finding_key(rule_code: str, sheet: str | None, cell: str | None) -> str:
    """Create a stable, value-free key for same-session re-validation.

    A key intentionally contains only the rule and normalized location. It is
    not a workbook fingerprint and does not try to identify moved cells or
    renamed sheets across different workbook versions.
    """

    location = "|".join(
        value.strip().casefold() for value in (sheet, cell) if value and value.strip()
    )
    return f"{rule_code}|{location}" if location else rule_code


def _formula_text(cell: Cell) -> str | None:
    value = cell.value
    if cell.data_type == "f" and isinstance(value, str):
        return value
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _risk_band(score: int) -> str:
    if score < 20:
        return "low"
    if score < 45:
        return "moderate"
    if score < 70:
        return "high"
    return "critical"


def _complexity_band(score: int) -> str:
    if score <= 3:
        return "basic"
    if score <= 7:
        return "standard"
    if score <= 11:
        return "advanced"
    return "expert"


def _calculate_risk(collector: FindingCollector, workbook: WorkbookSummary) -> int:
    critical = min(collector.severity_counts["critical"] * 16, 48)
    warning = min(collector.severity_counts["warning"] * 6, 36)
    info = min(collector.severity_counts["info"], 10)
    modifiers = 0
    if workbook.has_macros:
        modifiers += 6
    if workbook.external_link_count:
        modifiers += min(6, workbook.external_link_count * 2)
    if workbook.scan_truncated:
        modifiers += 4
    return min(100, critical + warning + info + modifiers)


def _calculate_complexity(
    *,
    file_size_bytes: int,
    workbook: WorkbookSummary,
    collector: FindingCollector,
) -> int:
    score = 0
    if file_size_bytes > 5 * 1024 * 1024:
        score += 1
    if workbook.sheet_count > 10:
        score += 1
    if workbook.sheet_count > 25:
        score += 1
    if workbook.formula_count > 5_000:
        score += 1
    if workbook.formula_count > 20_000:
        score += 2
    if workbook.external_link_count:
        score += 2
    if workbook.has_macros:
        score += 3
    if workbook.drawing_part_count:
        score += 1
    if workbook.merged_range_count > 100:
        score += 1
    if collector.severity_counts["critical"]:
        score += 2
    if collector.total > 20:
        score += 1
    if workbook.scan_truncated:
        score += 2
    return score


def _defined_name_count(workbook: Any, collector: FindingCollector) -> int:
    count = 0
    try:
        names = list(workbook.defined_names.values())
    except (AttributeError, TypeError):
        return 0

    for name in names:
        count += 1
        attr_text = getattr(name, "attr_text", None)
        if isinstance(attr_text, str) and "#REF!" in attr_text.upper():
            collector.add(
                rule_code="DEFINED_NAME_BROKEN_REF",
                severity="critical",
                title="정의된 이름에 깨진 참조가 있습니다.",
                description=(
                    "이름 범위가 삭제된 셀이나 시트를 가리켜 "
                    "수식 결과에 영향을 줄 수 있습니다."
                ),
                repair_class="EXPERT_REVIEW",
                cell=getattr(name, "name", None),
            )
    return count


def scan_workbook(filename: str, payload: bytes, settings: Settings) -> ScanResult:
    envelope = validate_ooxml(filename, payload, settings)
    collector = FindingCollector.create(settings.finding_limit)

    source = BytesIO(payload)
    try:
        workbook = load_workbook(
            source,
            read_only=False,
            data_only=False,
            keep_links=True,
            # Static diagnosis never writes a workbook, so retaining a VBA
            # archive is unnecessary and creates an avoidable temporary ZIP.
            keep_vba=False,
        )
    except Exception as exc:  # openpyxl exposes multiple parser exceptions
        source.close()
        raise WorkbookCareError(
            "WORKBOOK_PARSE_FAILED",
            (
                "통합문서 구조를 읽지 못했습니다. 파일이 손상되었거나 "
                "지원되지 않는 요소가 있을 수 있습니다."
            ),
            status_code=422,
        ) from exc

    if envelope.has_macros:
        collector.add(
            rule_code="FILE_MACRO_ENABLED",
            severity="warning",
            title="매크로가 포함된 통합문서입니다.",
            description=(
                "보안을 위해 매크로는 실행하거나 수정하지 않으며 "
                "자동 수정 범위에서 제외합니다."
            ),
            repair_class="EXPERT_REVIEW",
            confidence=1.0,
        )

    if envelope.drawing_part_count:
        collector.add(
            rule_code="FILE_DRAWING_PARTS",
            severity="info",
            title="차트·도형·이미지 요소가 포함되어 있습니다.",
            description=(
                "파일 보존 위험 때문에 단순한 라이브러리 저장 방식으로 "
                "자동 수정하지 않습니다."
            ),
            repair_class="EXPERT_REVIEW",
            confidence=1.0,
        )

    hidden_count = 0
    very_hidden_count = 0
    merged_count = 0
    formula_count = 0
    scanned_cell_count = 0
    scan_truncated = False
    formula_external_locations: set[tuple[str, str]] = set()
    numeric_cells_by_column: Counter[tuple[str, int]] = Counter()
    nonempty_cells_by_column: Counter[tuple[str, int]] = Counter()
    numeric_text_candidates: list[tuple[str, str, int, str]] = []

    for worksheet in workbook.worksheets:
        if worksheet.sheet_state == "hidden":
            hidden_count += 1
            collector.add(
                rule_code="SHEET_HIDDEN",
                severity="info",
                title="숨겨진 시트가 있습니다.",
                description="의도적으로 숨긴 시트인지 확인하는 것이 좋습니다.",
                repair_class="CONFIRMATION_REQUIRED",
                sheet=worksheet.title,
            )
        elif worksheet.sheet_state == "veryHidden":
            very_hidden_count += 1
            collector.add(
                rule_code="SHEET_VERY_HIDDEN",
                severity="warning",
                title="일반 메뉴에서 표시할 수 없는 VeryHidden 시트가 있습니다.",
                description="VBA 또는 속성 편집으로만 표시되는 시트이므로 용도 확인이 필요합니다.",
                repair_class="CONFIRMATION_REQUIRED",
                sheet=worksheet.title,
            )

        merged_count += len(worksheet.merged_cells.ranges)

        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                scanned_cell_count += 1
                if scanned_cell_count > settings.scan_cell_limit:
                    scan_truncated = True
                    break

                column_key = (worksheet.title, cell.column)

                formula = _formula_text(cell)
                if formula is None:
                    nonempty_cells_by_column[column_key] += 1
                    value = cell.value
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric_cells_by_column[column_key] += 1
                    elif isinstance(value, str):
                        stripped = value.strip()
                        has_ambiguous_leading_zero = (
                            len(stripped) > 1
                            and stripped[0] in {"0", "+", "-"}
                            and stripped.lstrip("+-").startswith("0")
                        )
                        if (
                            cell.row > 1
                            and value == stripped
                            and not has_ambiguous_leading_zero
                            and NUMERIC_TEXT_PATTERN.fullmatch(stripped)
                        ):
                            numeric_text_candidates.append(
                                (worksheet.title, cell.coordinate, cell.column, stripped)
                            )
                    continue

                formula_count += 1
                upper_formula = formula.upper()

                if "#REF!" in upper_formula:
                    collector.add(
                        rule_code="FORMULA_REF_ERROR",
                        severity="critical",
                        title="삭제되거나 이동된 참조가 수식에 포함되어 있습니다.",
                description=(
                    "수식 문자열에서 #REF! 참조 토큰을 확인했습니다. "
                    "계산 결과와 수식 의도는 이번 검사에서 확인하지 않았습니다."
                        ),
                        repair_class="EXPERT_REVIEW",
                        sheet=worksheet.title,
                        cell=cell.coordinate,
                    )
                elif any(token in upper_formula for token in ERROR_TOKENS):
                    collector.add(
                        rule_code="FORMULA_VISIBLE_ERROR_TOKEN",
                        severity="critical",
                        title="오류 토큰이 포함된 수식입니다.",
                description=(
                    "수식 문자열에서 Excel 오류 토큰을 확인했습니다. "
                    "계산 결과와 수식 의도는 이번 검사에서 확인하지 않았습니다."
                        ),
                        repair_class="EXPERT_REVIEW",
                        sheet=worksheet.title,
                        cell=cell.coordinate,
                    )

                if VOLATILE_PATTERN.search(formula):
                    collector.add(
                        rule_code="FORMULA_VOLATILE",
                        severity="info",
                        title="재계산이 잦은 휘발성 함수가 사용되었습니다.",
                        description="대규모 통합문서에서는 계산 지연의 원인이 될 수 있습니다.",
                        repair_class="INFORMATION_ONLY",
                        sheet=worksheet.title,
                        cell=cell.coordinate,
                        confidence=0.95,
                    )

                if upper_formula.count("IF(") >= 5:
                    collector.add(
                        rule_code="FORMULA_DEEP_NESTING",
                        severity="warning",
                        title="중첩 IF가 깊어 유지보수가 어렵습니다.",
                        description=(
                            "업무 로직 변경 시 실수 가능성이 높아 "
                            "구조 개선을 검토할 수 있습니다."
                        ),
                        repair_class="INFORMATION_ONLY",
                        sheet=worksheet.title,
                        cell=cell.coordinate,
                        confidence=0.93,
                    )

                if WHOLE_COLUMN_PATTERN.search(upper_formula):
                    collector.add(
                        rule_code="FORMULA_WHOLE_COLUMN_REFERENCE",
                        severity="info",
                        title="전체 열을 참조하는 수식입니다.",
                        description="수식 수가 많으면 계산량과 파일 성능에 영향을 줄 수 있습니다.",
                        repair_class="INFORMATION_ONLY",
                        sheet=worksheet.title,
                        cell=cell.coordinate,
                        confidence=0.9,
                    )

                if EXTERNAL_FORMULA_PATTERN.search(formula):
                    key = (worksheet.title, cell.coordinate)
                    if key not in formula_external_locations:
                        formula_external_locations.add(key)
                        collector.add(
                            rule_code="FORMULA_EXTERNAL_REFERENCE",
                            severity="warning",
                            title="외부 통합문서 참조가 발견됨",
                            description=(
                                "원본 파일이 이동되거나 갱신되지 않으면 "
                                "오래된 값이 남을 수 있습니다."
                            ),
                            repair_class="CONFIRMATION_REQUIRED",
                            sheet=worksheet.title,
                            cell=cell.coordinate,
                            confidence=0.99,
                        )
            if scan_truncated:
                break
        if scan_truncated:
            break

    for sheet_name, coordinate, column_index, _raw_value in numeric_text_candidates:
        key = (sheet_name, column_index)
        numeric_peers = numeric_cells_by_column[key]
        nonempty_peers = nonempty_cells_by_column[key]
        if numeric_peers >= 3 and numeric_peers / max(nonempty_peers, 1) >= 0.6:
            collector.add(
                rule_code="NUMBER_STORED_AS_TEXT",
                severity="warning",
                title="숫자가 텍스트로 저장된 것으로 보입니다.",
                description=(
                    "같은 열의 숫자 패턴과 다르게 숫자 형태의 텍스트가 저장되어 있습니다. "
                    "합계 또는 조건부 집계에서 누락될 가능성이 있습니다. "
                    "실제 집계 결과와 변환 필요 여부는 이번 검사에서 확인하지 않았습니다."
                ),
                repair_class="SAFE_CANDIDATE",
                sheet=sheet_name,
                cell=coordinate,
                confidence=0.97,
            )

    if merged_count > 100:
        collector.add(
            rule_code="MERGED_CELL_HEAVY",
            severity="info",
            title="병합된 셀이 많이 사용되었습니다.",
            description="정렬, 필터, 복사와 자동화 작업의 안정성을 낮출 수 있습니다.",
            repair_class="INFORMATION_ONLY",
            confidence=0.95,
        )

    if formula_count > 50_000:
        collector.add(
            rule_code="FORMULA_COUNT_HIGH",
            severity="info",
            title="수식 수가 매우 많습니다.",
            description="계산과 파일 열기 속도에 영향을 줄 수 있어 성능 점검이 필요합니다.",
            repair_class="INFORMATION_ONLY",
            confidence=1.0,
        )

    if scan_truncated:
        collector.add(
            rule_code="SCAN_CELL_LIMIT_REACHED",
            severity="warning",
            title="안전 제한 때문에 일부 셀만 검사했습니다.",
            description="더 큰 처리 환경이나 전문가 검토가 필요할 수 있습니다.",
            repair_class="EXPERT_REVIEW",
            confidence=1.0,
        )

    defined_name_count = _defined_name_count(workbook, collector)
    external_link_count = max(
        envelope.external_link_part_count,
        len(getattr(workbook, "_external_links", []) or []),
        len(formula_external_locations),
    )

    workbook_summary = WorkbookSummary(
        sheet_count=len(workbook.worksheets),
        hidden_sheet_count=hidden_count,
        very_hidden_sheet_count=very_hidden_count,
        formula_count=formula_count,
        merged_range_count=merged_count,
        external_link_count=external_link_count,
        defined_name_count=defined_name_count,
        drawing_part_count=envelope.drawing_part_count,
        has_macros=envelope.has_macros,
        scanned_cell_count=min(scanned_cell_count, settings.scan_cell_limit),
        scan_truncated=scan_truncated,
    )

    risk_score = _calculate_risk(collector, workbook_summary)
    complexity_score = _calculate_complexity(
        file_size_bytes=len(payload),
        workbook=workbook_summary,
        collector=collector,
    )

    diagnosis_summary = DiagnosisSummary(
        risk_score=risk_score,
        risk_band=_risk_band(risk_score),  # type: ignore[arg-type]
        complexity_score=complexity_score,
        complexity_band=_complexity_band(complexity_score),  # type: ignore[arg-type]
        issue_count=collector.total,
        critical_count=collector.severity_counts["critical"],
        warning_count=collector.severity_counts["warning"],
        info_count=collector.severity_counts["info"],
        safe_candidate_count=collector.repair_counts["SAFE_CANDIDATE"],
        confirmation_required_count=collector.repair_counts["CONFIRMATION_REQUIRED"],
        expert_review_count=collector.repair_counts["EXPERT_REVIEW"],
        information_only_count=collector.repair_counts["INFORMATION_ONLY"],
        repair_review_candidate_count=(
            collector.repair_counts["SAFE_CANDIDATE"]
            + collector.repair_counts["CONFIRMATION_REQUIRED"]
            + collector.repair_counts["EXPERT_REVIEW"]
        ),
    )

    quote = build_quote(diagnosis_summary, workbook_summary, collector.findings)
    safe_filename = Path(filename).name

    result = ScanResult(
        analysis_id=str(uuid4()),
        filename=safe_filename,
        file_size_bytes=len(payload),
        scanned_at=datetime.now(UTC),
        scanner_version=SCANNER_VERSION,
        rule_set_version=RULE_SET_VERSION,
        workbook=workbook_summary,
        summary=diagnosis_summary,
        findings=collector.findings,
        quote=quote,
        limitations=[
            "정적 구조 분석 결과이며 Excel 계산 엔진을 실행하지 않았습니다.",
            "업무 의미상 잘못된 수식까지 모두 발견한다고 보장하지 않습니다.",
            "현재 버전은 파일을 수정하거나 다시 저장하지 않습니다.",
        ],
    )
    workbook.close()
    source.close()
    return result


def _formula_audit_result(
    *,
    status: str,
    started_at: float,
    formula_cell_count: int = 0,
    audited_sheet_count: int = 0,
    audited_formula_region_count: int = 0,
    candidates: list[Finding] | None = None,
    limitations: list[str] | None = None,
) -> FormulaAuditResult:
    return FormulaAuditResult(
        status=status,  # type: ignore[arg-type]
        formula_cell_count=formula_cell_count,
        audited_sheet_count=audited_sheet_count,
        audited_formula_region_count=audited_formula_region_count,
        candidates=candidates or [],
        limitations=limitations or [],
        elapsed_ms=max(0, round((perf_counter() - started_at) * 1000)),
        scanner_version=SCANNER_VERSION,
        rule_set_version=FORMULA_AUDIT_RULE_SET_VERSION,
    )


def _formula_audit_finding(
    *,
    rule_code: str,
    title: str,
    description: str,
    sheet: str,
    cell: str,
    evidence: FormulaPatternEvidence,
) -> Finding:
    """Create an isolated candidate without changing free-scan aggregates."""

    return Finding(
        id=str(uuid4()),
        finding_key=_finding_key(rule_code, sheet, cell),
        rule_code=rule_code,
        severity="warning",
        title=title,
        description=description,
        sheet=sheet,
        cell=cell,
        confidence=0.0,
        repair_class="EXPERT_REVIEW",
        formula_pattern=evidence,
    )


def run_formula_audit(
    filename: str,
    payload: bytes,
    settings: Settings,
) -> FormulaAuditResult:
    """Run the opt-in M4 audit in a result envelope separate from ``ScanResult``.

    This is static inspection only. It validates and reopens the browser's
    re-uploaded file for the audit request, never stores it, and never invokes
    the free-diagnosis collector, risk score, quote, or CSV path.
    """

    started_at = perf_counter()
    source: BytesIO | None = None
    workbook: Any | None = None
    formula_cell_count = 0
    audited_sheets: set[str] = set()

    try:
        validate_ooxml(filename, payload, settings)
        source = BytesIO(payload)
        workbook = load_workbook(
            source,
            read_only=False,
            data_only=False,
            keep_links=True,
            keep_vba=False,
        )

        if len(workbook.worksheets) > settings.formula_audit_max_sheet_count:
            return _formula_audit_result(
                status="SKIPPED_WORKBOOK_LIMIT",
                started_at=started_at,
                audited_sheet_count=len(workbook.worksheets),
                limitations=[
                    "통합문서 시트 수가 내부 베타 정밀검사 한도를 초과했습니다.",
                    "기본 무료 진단 결과는 그대로 사용할 수 있습니다.",
                ],
            )

        scanned_cell_count = 0
        scan_truncated = False
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    scanned_cell_count += 1
                    if scanned_cell_count > settings.scan_cell_limit:
                        scan_truncated = True
                        break
                    formula = _formula_text(cell)
                    if formula is None:
                        continue
                    formula_cell_count += 1
                    audited_sheets.add(worksheet.title)
                if scan_truncated:
                    break
            if scan_truncated:
                break

        if scan_truncated:
            return _formula_audit_result(
                status="SKIPPED_TRUNCATED",
                started_at=started_at,
                formula_cell_count=formula_cell_count,
                audited_sheet_count=len(audited_sheets),
                limitations=[
                    "안전 제한 때문에 검사 범위가 잘려 수식 패턴 정밀검사를 실행하지 않았습니다.",
                    "기본 무료 진단 결과는 그대로 유지됩니다.",
                ],
            )

        if formula_cell_count > settings.formula_audit_max_formula_cells:
            return _formula_audit_result(
                status="SKIPPED_FORMULA_LIMIT",
                started_at=started_at,
                formula_cell_count=formula_cell_count,
                audited_sheet_count=len(audited_sheets),
                limitations=[
                    "내부 베타의 검증된 수식 셀 한도(30,000개)를 넘어 "
                    "이번 정밀검사를 실행하지 않았습니다.",
                    "기본 무료 진단 결과는 그대로 유지됩니다.",
                ],
            )

        if formula_cell_count == 0:
            return _formula_audit_result(
                status="ABSTAINED_INSUFFICIENT_EVIDENCE",
                started_at=started_at,
                formula_cell_count=0,
                audited_sheet_count=0,
                limitations=[
                    "수식 셀이 없어 주변 수식 패턴을 비교하지 않았습니다.",
                    "후보가 없다는 뜻은 파일의 수식이나 계산 결과가 정확하다는 뜻이 아닙니다.",
                ],
            )

        worksheet_audits = [
            inspect_worksheet_formula_patterns(worksheet)
            for worksheet in workbook.worksheets
        ]
        supported_formula_count = sum(
            audit.supported_formula_count for audit in worksheet_audits
        )

        if supported_formula_count == 0:
            return _formula_audit_result(
                status="SKIPPED_UNSUPPORTED_STRUCTURE",
                started_at=started_at,
                formula_cell_count=formula_cell_count,
                audited_sheet_count=len(audited_sheets),
                limitations=[
                    "현재 내부 베타가 비교할 수 없는 수식 구조라 이번 정밀검사를 생략했습니다.",
                    "지원하지 않는 문법·외부 통합문서·이름 범위·배열·Table 구조는 "
                    "후보화하지 않습니다.",
                ],
            )

        if supported_formula_count < 4:
            return _formula_audit_result(
                status="ABSTAINED_INSUFFICIENT_EVIDENCE",
                started_at=started_at,
                formula_cell_count=formula_cell_count,
                audited_sheet_count=len(audited_sheets),
                limitations=[
                    "비교할 수 있는 주변 수식 패턴이 충분하지 않아 판단을 유보했습니다.",
                    "후보가 없다는 뜻은 파일의 수식이나 계산 결과가 정확하다는 뜻이 아닙니다.",
                ],
            )

        raw_candidates = [
            candidate
            for audit in worksheet_audits
            for candidate in audit.candidates
            if candidate.rule_code in FORMULA_AUDIT_ALLOWED_RULE_CODES
            and candidate.evidence.pattern_subtype != "GENERIC_PATTERN_DRIFT"
        ]
        if len(raw_candidates) > settings.formula_audit_max_candidate_count:
            return _formula_audit_result(
                status="SKIPPED_CANDIDATE_LIMIT",
                started_at=started_at,
                formula_cell_count=formula_cell_count,
                audited_sheet_count=len(audited_sheets),
                audited_formula_region_count=len(
                    {
                        (candidate.sheet, candidate.evidence.formula_region)
                        for candidate in raw_candidates
                    }
                ),
                limitations=[
                    "이상 후보 수가 내부 베타 정밀검사 한도를 초과했습니다.",
                    "부분 결과를 제공하지 않았습니다.",
                    "기본 무료 진단 결과는 그대로 사용할 수 있습니다.",
                ],
            )
        candidates = [
            _formula_audit_finding(
                rule_code=candidate.rule_code,
                title=(
                    "주변 수식과 다른 패턴 후보"
                    if candidate.rule_code == "FORMULA_PATTERN_OUTLIER"
                    else "반복 수식 영역의 패턴 공백 후보"
                ),
                description=(
                    "주변 반복 수식과 다른 구조를 확인한 후보입니다. "
                    "업무적으로 정상일 수 있으므로 사용자가 직접 확인해야 합니다."
                ),
                sheet=candidate.sheet,
                cell=candidate.cell,
                evidence=candidate.evidence,
            )
            for candidate in raw_candidates
        ]
        regions = {
            (candidate.sheet, candidate.evidence.formula_region)
            for candidate in raw_candidates
        }
        return _formula_audit_result(
            status="COMPLETED",
            started_at=started_at,
            formula_cell_count=formula_cell_count,
            audited_sheet_count=len(audited_sheets),
            audited_formula_region_count=len(regions),
            candidates=candidates,
            limitations=[
                "수식 계산 결과, 업무 규칙, 의도된 예외 계산은 확인하지 않았습니다.",
                "지원하는 A1 참조 수식의 제한된 주변 패턴만 비교했습니다.",
                "수정 수식 생성·적용이나 자동 수정 가능 여부는 제공하지 않습니다.",
            ],
        )
    except Exception:
        return _formula_audit_result(
            status="FAILED",
            started_at=started_at,
            formula_cell_count=formula_cell_count,
            audited_sheet_count=len(audited_sheets),
            limitations=[
                "수식 패턴 정밀검사를 완료하지 못했습니다. 기본 무료 진단 결과는 유지됩니다.",
                "파일을 다시 선택해 무료 진단 결과를 기준으로 확인하세요.",
            ],
        )
    finally:
        if workbook is not None:
            workbook.close()
        if source is not None:
            source.close()
