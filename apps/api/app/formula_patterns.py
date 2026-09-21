"""Conservative, value-free formula-pattern candidate detection for M4-A.

This module does not evaluate formulas, infer business intent, produce a
replacement formula, or alter a workbook. It recognizes only a narrow subset
of A1-style formulas that can be normalized without retaining formula text.
Any unsupported syntax or likely structural exception is excluded.
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from hashlib import sha256

from openpyxl.cell.cell import Cell
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from .models import FormulaPatternEvidence

_CELL_REFERENCE = re.compile(
    r"^(?P<column_absolute>\$?)(?P<column>[A-Za-z]{1,3})"
    r"(?P<row_absolute>\$?)(?P<row>[1-9]\d*)$"
)
_SUMMARY_MARKERS = (
    "subtotal",
    "total",
    "grand total",
    "소계",
    "합계",
    "총계",
    "조정",
    "보정",
    "adjustment",
    "manual",
)
_STRUCTURAL_TOKEN_TYPES = {
    "FUNC",
    "PAREN",
    "SEP",
    # openpyxl Tokenizer uses the long names below. The abbreviated names are
    # retained for compatibility with older tokenizer variants.
    "OPERATOR-INFIX",
    "OPERATOR-PREFIX",
    "OPERATOR-POSTFIX",
    "OP_IN",
    "OP_PRE",
    "OP_POST",
}
_FUNCTION_SIGNATURE = re.compile(r"(?P<name>[A-Z][A-Z0-9_.]*)\(")
_REFERENCE_SIGNATURE = re.compile(
    r"REF\((?:(?:SHEET\((?P<sheet>[^)]*)\)!)?)(?P<coordinates>[^)]*)\)"
)
_MONTH_HEADER = re.compile(r"^M(?P<number>0[1-9]|1[0-2])$")


@dataclass(frozen=True, slots=True)
class FormulaPatternCandidate:
    rule_code: str
    title: str
    description: str
    sheet: str
    cell: str
    evidence: FormulaPatternEvidence


@dataclass(frozen=True, slots=True)
class FormulaPatternWorksheetAudit:
    """Private execution metadata plus candidates for one worksheet.

    The support count is used only to choose a conservative execution outcome;
    neither formula text nor cell values are retained in a result payload.
    """

    candidates: list[FormulaPatternCandidate]
    supported_formula_count: int


@dataclass(frozen=True, slots=True)
class _NormalizedFormula:
    cell: Cell
    signature: str
    pattern_id: str


@dataclass(frozen=True, slots=True)
class _MonthlyFormula:
    cell: Cell
    header_row: int
    month_number: int
    signature: str
    pattern_id: str


@dataclass(frozen=True, slots=True)
class _MonthlyReference:
    sheet: str | None
    column: str
    row: int
    relative_column: int
    relative_row: int


@dataclass(frozen=True, slots=True)
class _PatternProfile:
    """Private comparison metadata. It is never returned in an API payload."""

    functions: tuple[str, ...]
    sheets: tuple[str, ...]
    absolute_styles: tuple[str, ...]
    absolute_coordinates: tuple[str, ...]
    range_shapes: tuple[int, ...]
    relative_shapes: tuple[str, ...]


def _formula_text(cell: Cell) -> str | None:
    value = cell.value
    if cell.data_type == "f" and isinstance(value, str):
        return value
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _pattern_id(signature: str) -> str:
    return sha256(signature.encode("utf-8")).hexdigest()[:16]


def _normalise_reference(reference: str, *, origin_row: int, origin_column: int) -> str | None:
    """Represent a supported A1 reference relative to its formula cell."""
    if "[" in reference or "]" in reference:
        return None

    sheet_prefix = ""
    coordinate_text = reference
    if "!" in reference:
        sheet_name, coordinate_text = reference.rsplit("!", 1)
        if not sheet_name:
            return None
        sheet_prefix = f"SHEET({sheet_name.casefold()})!"

    parts = coordinate_text.split(":")
    if len(parts) not in {1, 2} or not all(parts):
        return None

    normalised_parts: list[str] = []
    for part in parts:
        match = _CELL_REFERENCE.fullmatch(part)
        if match is None:
            return None
        try:
            column_index = column_index_from_string(match.group("column").upper())
        except ValueError:
            return None
        row_index = int(match.group("row"))

        if match.group("column_absolute"):
            column_signature = f"C${column_index}"
        else:
            column_signature = f"C[{column_index - origin_column}]"
        if match.group("row_absolute"):
            row_signature = f"R${row_index}"
        else:
            row_signature = f"R[{row_index - origin_row}]"
        normalised_parts.append(f"{column_signature}{row_signature}")

    return sheet_prefix + ":".join(normalised_parts)


def normalize_formula(formula: str, origin: Cell) -> str | None:
    """Return a safe normalized signature, or None when parsing is uncertain.

    Literal operands are represented by their value-free token category so
    repeated business formulas such as ``IF`` and ``IFERROR`` remain
    comparable without retaining a literal value. Array, table,
    external-workbook, and named-range operands remain unsupported. The
    returned signature is used only to create a one-way pattern ID and is
    never emitted.
    """
    if not formula.startswith("=") or "[" in formula or "]" in formula:
        return None
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return None

    canonical: list[str] = []
    has_reference = False
    function_stack: list[str] = []
    for token in tokens:
        if token.type == "WSPACE":
            continue
        if token.type == "FUNC":
            if token.subtype == "OPEN":
                function_stack.append(token.value[:-1].upper())
            elif token.subtype == "CLOSE" and function_stack:
                function_stack.pop()
            canonical.append(token.value.upper())
            continue
        if token.type == "OPERAND":
            if token.subtype == "RANGE":
                normalized_reference = _normalise_reference(
                    token.value,
                    origin_row=origin.row,
                    origin_column=origin.column,
                )
                if normalized_reference is None:
                    return None
                canonical.append(f"REF({normalized_reference})")
                has_reference = True
                continue
            literal_category = {
                "NUMBER": "NUMBER",
                "TEXT": "TEXT",
                "LOGICAL": "LOGICAL",
                "ERROR": "ERROR",
            }.get(token.subtype)
            if literal_category is None:
                return None
            if function_stack and function_stack[-1] in {"SUM", "AVERAGE"}:
                # A static aggregate operand can be an intentional adjustment.
                # Keep this narrow legacy syntax outside the M4-C comparison
                # surface until it has its own target/normal quality gate.
                return None
            canonical.append(f"LITERAL({literal_category})")
            continue
        if token.type in _STRUCTURAL_TOKEN_TYPES:
            canonical.append(token.value.upper())
            continue
        return None

    if not has_reference or not canonical:
        return None
    return "".join(canonical)


def _normalize_horizontal_arithmetic_formula(formula: str, origin: Cell) -> str | None:
    """Normalize only same-sheet, unanchored A1 arithmetic formulas."""
    unsupported_markers = ("[", "]", "!", "$", ":")
    if not formula.startswith("=") or any(marker in formula for marker in unsupported_markers):
        return None
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return None

    canonical: list[str] = []
    has_reference = False
    for token in tokens:
        if token.type == "WSPACE":
            continue
        if token.type == "OPERAND":
            if token.subtype != "RANGE":
                return None
            normalized_reference = _normalise_reference(
                token.value,
                origin_row=origin.row,
                origin_column=origin.column,
            )
            if normalized_reference is None:
                return None
            canonical.append(f"REF({normalized_reference})")
            has_reference = True
            continue
        if token.type == "OPERATOR-INFIX" and token.value in {"+", "-", "*", "/", "^"}:
            canonical.append(token.value)
            continue
        if token.type == "PAREN":
            canonical.append(token.value)
            continue
        return None

    if not has_reference or not canonical:
        return None
    return "".join(canonical)


def _sheet_name_token(sheet_name: str) -> str:
    return sheet_name.strip("'").casefold()


def _parse_binary_subtraction_references(
    formula: str, origin: Cell
) -> tuple[_MonthlyReference, _MonthlyReference] | None:
    if not formula.startswith("=") or any(marker in formula for marker in ("[", "]", "$", ":")):
        return None
    try:
        tokens = [token for token in Tokenizer(formula).items if token.type != "WSPACE"]
    except Exception:
        return None
    if len(tokens) != 3 or tokens[1].type != "OPERATOR-INFIX" or tokens[1].value != "-":
        return None
    references: list[_MonthlyReference] = []
    for token in (tokens[0], tokens[2]):
        if token.type != "OPERAND" or token.subtype != "RANGE":
            return None
        sheet_name: str | None = None
        coordinate_text = token.value
        if "!" in coordinate_text:
            raw_sheet, coordinate_text = coordinate_text.rsplit("!", 1)
            if not raw_sheet:
                return None
            sheet_name = _sheet_name_token(raw_sheet)
        match = _CELL_REFERENCE.fullmatch(coordinate_text)
        if match is None:
            return None
        try:
            column_index = column_index_from_string(match.group("column").upper())
        except ValueError:
            return None
        references.append(
            _MonthlyReference(
                sheet=sheet_name,
                column=match.group("column").upper(),
                row=int(match.group("row")),
                relative_column=column_index - origin.column,
                relative_row=int(match.group("row")) - origin.row,
            )
        )
    return references[0], references[1]


def _month_header_for_cell(worksheet: Worksheet, cell: Cell) -> tuple[int, int, str] | None:
    matches: list[tuple[int, int, str]] = []
    workbook_sheet_names = {_sheet_name_token(name) for name in worksheet.parent.sheetnames}
    for header_row in range(max(1, cell.row - 4), cell.row):
        if worksheet.row_dimensions[header_row].hidden:
            continue
        header_cell = worksheet.cell(row=header_row, column=cell.column)
        if _is_hidden_column(worksheet, cell.column) or _is_merged(worksheet, header_cell):
            continue
        value = header_cell.value
        if not isinstance(value, str):
            continue
        match = _MONTH_HEADER.fullmatch(value.strip().upper())
        if match is None:
            continue
        sheet_token = value.strip().casefold()
        if sheet_token not in workbook_sheet_names:
            continue
        matches.append((header_row, int(match.group("number")), sheet_token))
    if len(matches) != 1:
        return None
    return matches[0]


def _direct_monthly_signature(
    references: tuple[_MonthlyReference, _MonthlyReference],
    expected_sheet: str,
    worksheet: Worksheet,
) -> str | None:
    first, second = references
    if first.sheet is None or second.sheet is None:
        return None
    if first.sheet != second.sheet:
        return None
    workbook_sheet_names = {_sheet_name_token(name) for name in worksheet.parent.sheetnames}
    if first.sheet not in workbook_sheet_names:
        return None
    coordinate_signature = f"{first.column}{first.row},{second.column}{second.row}"
    if first.sheet == expected_sheet:
        return f"MONTHLY_SUB({coordinate_signature})"
    return f"MONTHLY_SHEET({first.sheet})_SUB({coordinate_signature})"


def _parse_single_monthly_reference(formula: str, origin: Cell) -> _MonthlyReference | None:
    if not formula.startswith("=") or any(marker in formula for marker in ("[", "]", "$", ":")):
        return None
    try:
        tokens = [token for token in Tokenizer(formula).items if token.type != "WSPACE"]
    except Exception:
        return None
    if len(tokens) != 1 or tokens[0].type != "OPERAND" or tokens[0].subtype != "RANGE":
        return None
    value = tokens[0].value
    sheet_name: str | None = None
    if "!" in value:
        raw_sheet, value = value.rsplit("!", 1)
        if not raw_sheet:
            return None
        sheet_name = _sheet_name_token(raw_sheet)
    match = _CELL_REFERENCE.fullmatch(value)
    if match is None:
        return None
    try:
        column_index = column_index_from_string(match.group("column").upper())
    except ValueError:
        return None
    return _MonthlyReference(
        sheet=sheet_name,
        column=match.group("column").upper(),
        row=int(match.group("row")),
        relative_column=column_index - origin.column,
        relative_row=int(match.group("row")) - origin.row,
    )


def _resolve_local_monthly_reference(
    worksheet: Worksheet,
    origin: Cell,
    reference: _MonthlyReference,
    expected_sheet: str,
) -> tuple[str, int] | None:
    if reference.sheet is not None or reference.relative_column != 0:
        return None
    source_cell = worksheet.cell(row=reference.row, column=origin.column)
    source_formula = _formula_text(source_cell)
    if source_formula is None:
        return None
    direct = _parse_single_monthly_reference(source_formula, source_cell)
    if direct is None or direct.sheet != expected_sheet:
        return None
    return direct.column, direct.row


def _normalize_monthly_subtraction_formula(
    formula: str,
    origin: Cell,
    worksheet: Worksheet,
) -> _MonthlyFormula | None:
    header = _month_header_for_cell(worksheet, origin)
    if header is None:
        return None
    header_row, month_number, expected_sheet = header
    references = _parse_binary_subtraction_references(formula, origin)
    if references is None:
        return None
    direct_signature = _direct_monthly_signature(references, expected_sheet, worksheet)
    if direct_signature is not None:
        return _MonthlyFormula(
            cell=origin,
            header_row=header_row,
            month_number=month_number,
            signature=direct_signature,
            pattern_id=_pattern_id(direct_signature),
        )
    if any(reference.sheet is not None for reference in references):
        return None
    resolved: list[tuple[str, int]] = []
    for reference in references:
        resolved_reference = _resolve_local_monthly_reference(
            worksheet,
            origin,
            reference,
            expected_sheet,
        )
        if resolved_reference is None:
            unresolved_signature = (
                "MONTHLY_LOCAL_UNRESOLVED("
                f"C[{references[0].relative_column}]R[{references[0].relative_row}]-"
                f"C[{references[1].relative_column}]R[{references[1].relative_row}])"
            )
            return _MonthlyFormula(
                cell=origin,
                header_row=header_row,
                month_number=month_number,
                signature=unresolved_signature,
                pattern_id=_pattern_id(unresolved_signature),
            )
        resolved.append(resolved_reference)
    signature = f"MONTHLY_SUB({resolved[0][0]}{resolved[0][1]},{resolved[1][0]}{resolved[1][1]})"
    return _MonthlyFormula(
        cell=origin,
        header_row=header_row,
        month_number=month_number,
        signature=signature,
        pattern_id=_pattern_id(signature),
    )


def _table_bounds(worksheet: Worksheet) -> list[tuple[int, int, int, int]]:
    bounds: list[tuple[int, int, int, int]] = []
    for table in worksheet.tables.values():
        try:
            bounds.append(range_boundaries(table.ref))
        except (TypeError, ValueError):
            continue
    return bounds


def _is_inside_table(
    *, row: int, column: int, table_bounds: list[tuple[int, int, int, int]]
) -> bool:
    return any(
        min_column <= column <= max_column and min_row <= row <= max_row
        for min_column, min_row, max_column, max_row in table_bounds
    )


def _is_merged(worksheet: Worksheet, cell: Cell) -> bool:
    return any(cell.coordinate in merged_range for merged_range in worksheet.merged_cells.ranges)


def _is_summary_or_manual_row(worksheet: Worksheet, row: int, cells: tuple[Cell, ...]) -> bool:
    if row <= 1 or worksheet.row_dimensions[row].hidden:
        return True
    for cell in cells:
        # A formula may legitimately refer to a worksheet named "총계..." or
        # "Adjustment...". Summary markers describe row labels, not formula
        # source text; inspecting formulas here would exclude every data row
        # in that pattern region.
        if _formula_text(cell) is not None:
            continue
        if isinstance(cell.value, str):
            text = cell.value.strip().casefold()
            if text and any(marker in text for marker in _SUMMARY_MARKERS):
                return True
    return False


def _summary_rows(worksheet: Worksheet) -> set[int]:
    """Find known summary/manual rows in one worksheet traversal.

    ``worksheet[row]`` performs repeated slice construction on large sheets.
    Using the already materialized row from ``iter_rows`` preserves the exact
    exclusion rules while keeping the audit linear in worksheet size.
    """
    return {
        row
        for row, cells in enumerate(worksheet.iter_rows(), start=1)
        if _is_summary_or_manual_row(worksheet, row, cells)
    }


def _is_excluded_location(
    worksheet: Worksheet,
    cell: Cell,
    table_bounds: list[tuple[int, int, int, int]],
    summary_rows: set[int],
) -> bool:
    return (
        cell.row in summary_rows
        or _is_merged(worksheet, cell)
        or _is_inside_table(row=cell.row, column=cell.column, table_bounds=table_bounds)
    )


def _is_hidden_column(worksheet: Worksheet, column: int) -> bool:
    letter = get_column_letter(column)
    direct_dimension = worksheet.column_dimensions[letter]
    if direct_dimension.hidden:
        return True
    for dimension in worksheet.column_dimensions.values():
        if not dimension.hidden:
            continue
        min_column = getattr(dimension, "min", None)
        max_column = getattr(dimension, "max", None)
        if (
            isinstance(min_column, int)
            and isinstance(max_column, int)
            and min_column <= column <= max_column
        ):
            return True
    return False


def _has_hidden_column_between(worksheet: Worksheet, columns: list[int]) -> bool:
    return any(
        _is_hidden_column(worksheet, column)
        for column in range(min(columns), max(columns) + 1)
    )


def _blank_row_has_context(worksheet: Worksheet, *, row: int, target_column: int) -> bool:
    """Avoid treating a fully empty separator row as a missing formula."""
    return any(
        cell.column != target_column and cell.value is not None
        for cell in worksheet[row]
    )


def _contiguous_formula_runs(items: list[_NormalizedFormula]) -> list[list[_NormalizedFormula]]:
    if not items:
        return []
    runs: list[list[_NormalizedFormula]] = [[items[0]]]
    for item in items[1:]:
        if item.cell.row == runs[-1][-1].cell.row + 1:
            runs[-1].append(item)
        else:
            runs.append([item])
    return runs


def _contiguous_horizontal_formula_runs(
    items: list[_NormalizedFormula],
) -> list[list[_NormalizedFormula]]:
    if not items:
        return []
    runs: list[list[_NormalizedFormula]] = [[items[0]]]
    for item in items[1:]:
        if item.cell.column == runs[-1][-1].cell.column + 1:
            runs[-1].append(item)
        else:
            runs.append([item])
    return runs


def _supporting_items(
    items: list[_NormalizedFormula],
    dominant_signature: str,
    candidate_row: int,
    worksheet: Worksheet,
    table_bounds: list[tuple[int, int, int, int]],
    summary_rows: set[int],
) -> list[_NormalizedFormula]:
    return [
        item
        for item in items
        if item.cell.row != candidate_row
        and item.signature == dominant_signature
        and not _is_excluded_location(worksheet, item.cell, table_bounds, summary_rows)
    ]


def _region(column: int, rows: list[int]) -> str:
    return f"{get_column_letter(column)}{min(rows)}:{get_column_letter(column)}{max(rows)}"


def _horizontal_region(row: int, columns: list[int]) -> str:
    return f"{get_column_letter(min(columns))}{row}:{get_column_letter(max(columns))}{row}"


def _monthly_headers_are_sequential(items: list[_MonthlyFormula]) -> bool:
    if len({item.header_row for item in items}) != 1:
        return False
    month_numbers = [item.month_number for item in items]
    return (
        len(month_numbers) == len(set(month_numbers))
        and month_numbers == list(range(month_numbers[0], month_numbers[0] + len(month_numbers)))
    )


def _ordered_monthly_supporting_cells(
    run: list[_MonthlyFormula],
    candidate: _MonthlyFormula,
    dominant_signature: str,
) -> list[Cell]:
    by_column = {item.cell.column: item for item in run}
    ordered: list[_MonthlyFormula] = []
    for column in (candidate.cell.column - 1, candidate.cell.column + 1):
        neighbor = by_column.get(column)
        if neighbor is not None and neighbor.signature == dominant_signature:
            ordered.append(neighbor)
    ordered.extend(
        item
        for item in run
        if item.cell.column != candidate.cell.column
        and item.signature == dominant_signature
        and item not in ordered
    )
    return [item.cell for item in ordered]


def _monthly_candidate_evidence(
    *,
    region: str,
    dominant_signature: str,
    current_signature: str,
    supporting_cells: list[Cell],
) -> FormulaPatternEvidence:
    comparison_locations = [cell.coordinate for cell in supporting_cells]
    return FormulaPatternEvidence(
        pattern_type="DOMINANT_NORMALIZED_PATTERN_OUTLIER",
        formula_region=region,
        dominant_pattern_id=_pattern_id(dominant_signature),
        current_pattern_id=_pattern_id(current_signature),
        neighbor_count=len(supporting_cells),
        evidence_locations=comparison_locations,
        detection_basis=(
            "M01~M12 header-to-sheet correspondence was compared across a same-row "
            "binary subtraction run; one interior cell differed while both adjacent "
            "month cells matched the dominant sheet-reference pattern."
        ),
        current_limitations=[
            "Formula results, business rules, and intentional exceptions were not evaluated.",
            "Only explicit M01~M12 month headers with existing sheets and simple "
            "subtraction are supported.",
            "No repair formula or automatic workbook change is generated.",
        ],
        pattern_subtype="REFERENCE_SHEET_DRIFT",
        evidence_summary=(
            "The target formula uses a different monthly sheet-reference pattern from "
            "the surrounding M01~M12 subtraction formulas."
        ),
        dominant_pattern_summary=(
            "Nearby month formulas repeat the same subtraction shape against their "
            "matching M01~M12 sheets."
        ),
        current_pattern_summary=(
            "The target formula does not follow the resolved monthly subtraction "
            "pattern used by the adjacent month cells."
        ),
        comparison_locations=comparison_locations,
        normal_case_possibility=(
            "This may be intentional when the workbook owner has a documented "
            "exception for this one month cell."
        ),
    )


def _pattern_profile(signature: str) -> _PatternProfile:
    """Extract only structural comparison facts from a private signature.

    The profile intentionally drops formula text, actual cell values, and the
    specific sheet identifiers before a Finding is created.
    """
    references = list(_REFERENCE_SIGNATURE.finditer(signature))
    functions = tuple(
        match.group("name")
        for match in _FUNCTION_SIGNATURE.finditer(signature)
        # ``REF`` is the normalizer's internal marker, not an Excel function.
        if match.group("name") != "REF"
    )
    coordinates = tuple(match.group("coordinates") for match in references)
    return _PatternProfile(
        functions=functions,
        sheets=tuple(
            _pattern_id(match.group("sheet")) if match.group("sheet") else "LOCAL"
            for match in references
        ),
        absolute_styles=tuple(
            "ABSOLUTE" if "$" in coordinate else "RELATIVE" for coordinate in coordinates
        ),
        absolute_coordinates=tuple(
            "|".join(re.findall(r"[CR]\$\d+", coordinate)) for coordinate in coordinates
        ),
        range_shapes=tuple(coordinate.count(":") + 1 for coordinate in coordinates),
        relative_shapes=coordinates,
    )


def _outlier_explanation(
    dominant_signature: str, current_signature: str
) -> tuple[str, str, str, str, str]:
    """Return a conservative, source-free subtype and display summaries."""
    dominant = _pattern_profile(dominant_signature)
    current = _pattern_profile(current_signature)
    normal_case = "소계·예외 계산 또는 업무상 의도된 다른 수식일 수 있습니다."
    dominant_summary = "주변 수식은 같은 함수와 참조 구조가 반복되는 패턴입니다."

    if dominant.functions != current.functions:
        return (
            "FUNCTION_PATTERN_DRIFT",
            "대상 수식의 함수 구성이 주변 반복 패턴과 다릅니다.",
            dominant_summary,
            "대상 수식은 주변과 다른 함수 구성을 사용합니다.",
            normal_case,
        )
    same_reference_arity = len(dominant.relative_shapes) == len(current.relative_shapes)
    changed_reference_pairs = (
        [
            (dominant_reference, current_reference)
            for dominant_reference, current_reference in zip(
                dominant.relative_shapes, current.relative_shapes, strict=True
            )
            if dominant_reference != current_reference
        ]
        if same_reference_arity
        else []
    )

    if same_reference_arity and dominant.sheets != current.sheets:
        return (
            "REFERENCE_SHEET_DRIFT",
            "대상 수식의 참조 시트 사용 방식이 주변 반복 패턴과 다릅니다.",
            dominant_summary,
            "대상 수식은 주변과 다른 참조 시트 사용 방식을 보입니다.",
            normal_case,
        )
    if same_reference_arity and dominant.absolute_styles != current.absolute_styles:
        return (
            "ABSOLUTE_REFERENCE_DRIFT",
            "대상 수식의 절대·상대 참조 방식이 주변 반복 패턴과 다릅니다.",
            dominant_summary,
            "대상 수식은 주변과 다른 절대·상대 참조 방식을 사용합니다.",
            normal_case,
        )
    if same_reference_arity and (
        dominant.absolute_coordinates != current.absolute_coordinates
        and any(dominant.absolute_coordinates)
    ):
        return (
            "REFERENCE_CELL_DRIFT",
            "대상 수식의 고정 참조 위치가 주변 반복 패턴과 다릅니다.",
            dominant_summary,
            "대상 수식은 주변과 다른 고정 참조 위치를 사용합니다.",
            normal_case,
        )
    if changed_reference_pairs and any(
        ":" in dominant_reference or ":" in current_reference
        for dominant_reference, current_reference in changed_reference_pairs
    ):
        return (
            "RANGE_BOUNDARY_DRIFT",
            "대상 수식의 범위 끝 위치가 주변 반복 패턴과 다릅니다.",
            dominant_summary,
            "대상 수식은 주변과 다른 범위 끝 위치를 사용합니다.",
            normal_case,
        )
    if dominant.relative_shapes != current.relative_shapes:
        return (
            "RELATIVE_REFERENCE_DRIFT",
            "대상 수식의 상대 참조 위치가 주변 반복 패턴과 다릅니다.",
            dominant_summary,
            "대상 수식은 주변과 다른 상대 참조 위치를 사용합니다.",
            normal_case,
        )
    return (
        "GENERIC_PATTERN_DRIFT",
        "대상 수식의 구조가 주변 반복 패턴과 다릅니다.",
        dominant_summary,
        "대상 수식은 주변과 다른 수식 구조를 사용합니다.",
        normal_case,
    )


def _candidate_evidence(
    *,
    pattern_type: str,
    region: str,
    dominant_signature: str,
    current_signature: str | None,
    supporting_cells: list[Cell],
    detection_basis: str,
    candidate_cell: Cell | None = None,
) -> FormulaPatternEvidence:
    comparison_locations = [cell.coordinate for cell in supporting_cells[:4]]
    if pattern_type == "FORMULA_GAP_OR_CONSTANT":
        is_blank = candidate_cell is not None and candidate_cell.value is None
        subtype = "BLANK_GAP_CANDIDATE" if is_blank else "CONSTANT_OVERRIDE_CANDIDATE"
        evidence_summary = (
            "반복 수식 사이의 대상 셀이 비어 있습니다."
            if is_blank
            else "반복 수식 사이의 대상 셀에 수식이 아닌 값이 있습니다."
        )
        current_summary = (
            "대상 셀은 빈 셀입니다."
            if is_blank
            else "대상 셀에는 수식이 아닌 값이 있습니다."
        )
        normal_case = "의도된 수동 입력·구분 행 또는 예외 값일 수 있습니다."
    elif current_signature is not None:
        (
            subtype,
            evidence_summary,
            dominant_summary,
            current_summary,
            normal_case,
        ) = _outlier_explanation(dominant_signature, current_signature)
    else:
        subtype = "GENERIC_PATTERN_DRIFT"
        evidence_summary = "주변 반복 수식과 다른 구조 후보를 확인했습니다."
        dominant_summary = "주변 수식은 같은 구조가 반복되는 패턴입니다."
        current_summary = None
        normal_case = "업무상 의도된 예외 구조일 수 있습니다."

    if pattern_type == "FORMULA_GAP_OR_CONSTANT":
        dominant_summary = "주변 수식은 같은 구조가 반복되는 패턴입니다."

    return FormulaPatternEvidence(
        pattern_type=pattern_type,  # type: ignore[arg-type]
        formula_region=region,
        dominant_pattern_id=_pattern_id(dominant_signature),
        current_pattern_id=_pattern_id(current_signature) if current_signature else None,
        neighbor_count=len(supporting_cells),
        evidence_locations=comparison_locations,
        detection_basis=detection_basis,
        current_limitations=[
            "수식 계산 결과, 업무 규칙, 의도된 예외 계산은 확인하지 않았습니다.",
            "지원하는 A1 참조 수식의 제한된 주변 패턴만 비교했습니다.",
            "수정 수식 생성·적용이나 자동 수정 가능 여부는 제공하지 않습니다.",
        ],
        pattern_subtype=subtype,  # type: ignore[arg-type]
        evidence_summary=evidence_summary,
        dominant_pattern_summary=dominant_summary,
        current_pattern_summary=current_summary,
        comparison_locations=comparison_locations,
        normal_case_possibility=normal_case,
    )


def inspect_worksheet_formula_patterns(worksheet: Worksheet) -> FormulaPatternWorksheetAudit:
    """Return conservative formula candidates for a single worksheet.

    The legacy comparison is column-local; the bounded horizontal pass is
    limited to same-row unanchored arithmetic formulas. Both require three or
    more surrounding formulas with one dominant normalized pattern. This
    deliberately misses many possible inconsistencies in exchange for lower
    false-positive risk.
    """
    table_bounds = _table_bounds(worksheet)
    summary_rows = _summary_rows(worksheet)
    normalized_by_column: dict[int, list[_NormalizedFormula]] = defaultdict(list)
    horizontal_by_row: dict[int, list[_NormalizedFormula]] = defaultdict(list)
    monthly_by_row: dict[int, list[_MonthlyFormula]] = defaultdict(list)
    supported_formula_count = 0

    for row in worksheet.iter_rows():
        for cell in row:
            formula = _formula_text(cell)
            if formula is None or _is_excluded_location(
                worksheet, cell, table_bounds, summary_rows
            ):
                continue
            signature = normalize_formula(formula, cell)
            if signature is not None:
                supported_formula_count += 1
                normalized_by_column[cell.column].append(
                    _NormalizedFormula(
                        cell=cell,
                        signature=signature,
                        pattern_id=_pattern_id(signature),
                    )
                )
            horizontal_signature = _normalize_horizontal_arithmetic_formula(formula, cell)
            if horizontal_signature is not None:
                horizontal_by_row[cell.row].append(
                    _NormalizedFormula(
                        cell=cell,
                        signature=horizontal_signature,
                        pattern_id=_pattern_id(horizontal_signature),
                    )
                )
            monthly_signature = _normalize_monthly_subtraction_formula(formula, cell, worksheet)
            if monthly_signature is not None:
                monthly_by_row[cell.row].append(monthly_signature)
    candidates: list[FormulaPatternCandidate] = []
    emitted: set[tuple[str, str]] = set()
    for column, normalized_items in normalized_by_column.items():
        normalized_items.sort(key=lambda item: item.cell.row)
        for run in _contiguous_formula_runs(normalized_items):
            if len(run) < 4:
                continue
            counts = Counter(item.signature for item in run)
            dominant_signature, dominant_count = counts.most_common(1)[0]
            if dominant_count < 3 or sum(count == dominant_count for count in counts.values()) != 1:
                continue

            by_row = {item.cell.row: item for item in run}
            for item in run[1:-1]:
                if item.signature == dominant_signature:
                    continue
                above = by_row.get(item.cell.row - 1)
                below = by_row.get(item.cell.row + 1)
                if above is None or below is None:
                    continue
                if above.signature != dominant_signature or below.signature != dominant_signature:
                    continue
                supporting = _supporting_items(
                    run,
                    dominant_signature,
                    item.cell.row,
                    worksheet,
                    table_bounds,
                    summary_rows,
                )
                if len(supporting) < 3:
                    continue
                key = ("FORMULA_PATTERN_OUTLIER", item.cell.coordinate)
                if key in emitted:
                    continue
                emitted.add(key)
                candidates.append(
                    FormulaPatternCandidate(
                        rule_code="FORMULA_PATTERN_OUTLIER",
                        title="주변 수식과 다른 패턴 후보가 있습니다.",
                        description=(
                            "같은 열의 연속 수식 영역에서 주변 수식과 다른 정규화 패턴을 "
                            "확인했습니다. 업무적으로 정상일 수 있으므로 사용자 확인과 "
                            "정밀검증이 필요합니다."
                        ),
                        sheet=worksheet.title,
                        cell=item.cell.coordinate,
                        evidence=_candidate_evidence(
                            pattern_type="DOMINANT_NORMALIZED_PATTERN_OUTLIER",
                            region=_region(column, [entry.cell.row for entry in run]),
                            dominant_signature=dominant_signature,
                            current_signature=item.signature,
                            supporting_cells=[entry.cell for entry in supporting],
                            detection_basis=(
                                "같은 열의 연속 수식 영역에서 최소 3개의 주변 수식이 같은 "
                                "정규화 패턴을 보이고, 대상 셀만 다른 패턴입니다."
                            ),
                            candidate_cell=item.cell,
                        ),
                    )
                )

        by_row = {item.cell.row: item for item in normalized_items}
        for row in range(2, worksheet.max_row):
            candidate_cell = worksheet.cell(row=row, column=column)
            if _formula_text(candidate_cell) is not None or _is_excluded_location(
                worksheet, candidate_cell, table_bounds, summary_rows
            ):
                continue
            if candidate_cell.value is None and not _blank_row_has_context(
                worksheet, row=row, target_column=column
            ):
                continue
            above = by_row.get(row - 1)
            below = by_row.get(row + 1)
            if above is None or below is None or above.signature != below.signature:
                continue
            supporting = [above, below]
            for nearby_row in (row - 2, row + 2):
                nearby = by_row.get(nearby_row)
                if nearby is not None and nearby.signature == above.signature:
                    supporting.append(nearby)
            supporting = [
                item
                for item in supporting
                if not _is_excluded_location(
                    worksheet, item.cell, table_bounds, summary_rows
                )
            ]
            if len(supporting) < 3:
                continue
            key = ("FORMULA_PATTERN_GAP", candidate_cell.coordinate)
            if key in emitted:
                continue
            emitted.add(key)
            candidates.append(
                FormulaPatternCandidate(
                    rule_code="FORMULA_PATTERN_GAP",
                    title="반복 수식 영역의 수식 누락 또는 상수 덮어쓰기 후보가 있습니다.",
                    description=(
                        "같은 열의 주변 수식 패턴 사이에서 대상 셀에 수식이 없습니다. "
                        "의도된 상수나 구분 행일 수 있으므로 사용자 확인과 정밀검증이 필요합니다."
                    ),
                    sheet=worksheet.title,
                    cell=candidate_cell.coordinate,
                    evidence=_candidate_evidence(
                        pattern_type="FORMULA_GAP_OR_CONSTANT",
                        region=_region(column, [item.cell.row for item in supporting] + [row]),
                        dominant_signature=above.signature,
                        current_signature=None,
                        supporting_cells=[
                            item.cell
                            for item in sorted(
                                supporting, key=lambda entry: entry.cell.row
                            )
                        ],
                        detection_basis=(
                            "대상 셀의 바로 위·아래 수식과 최소 3개의 주변 수식이 같은 "
                            "정규화 패턴을 보이지만 대상 셀에는 수식이 없습니다."
                        ),
                        candidate_cell=candidate_cell,
                    ),
                )
            )

    for row, normalized_items in horizontal_by_row.items():
        normalized_items.sort(key=lambda item: item.cell.column)
        for run in _contiguous_horizontal_formula_runs(normalized_items):
            if len(run) < 4:
                continue
            run_columns = [entry.cell.column for entry in run]
            if _has_hidden_column_between(worksheet, run_columns):
                continue
            counts = Counter(item.signature for item in run)
            dominant_signature, dominant_count = counts.most_common(1)[0]
            if dominant_count < 3 or sum(count == dominant_count for count in counts.values()) != 1:
                continue

            by_column = {item.cell.column: item for item in run}
            for item in run[1:-1]:
                if item.signature == dominant_signature:
                    continue
                left = by_column.get(item.cell.column - 1)
                right = by_column.get(item.cell.column + 1)
                if left is None or right is None:
                    continue
                if left.signature != dominant_signature or right.signature != dominant_signature:
                    continue
                supporting = [
                    entry
                    for entry in run
                    if entry.cell.column != item.cell.column
                    and entry.signature == dominant_signature
                    and not _is_excluded_location(
                        worksheet,
                        entry.cell,
                        table_bounds,
                        summary_rows,
                    )
                ]
                if len(supporting) < 3:
                    continue
                key = ("FORMULA_PATTERN_OUTLIER", item.cell.coordinate)
                if key in emitted:
                    continue
                emitted.add(key)
                candidates.append(
                    FormulaPatternCandidate(
                        rule_code="FORMULA_PATTERN_OUTLIER",
                        title="주변 수식과 다른 패턴 후보가 있습니다.",
                        description=(
                            "같은 행의 연속 수식 영역에서 주변 수식과 다른 정규화된 패턴을 "
                            "확인했습니다. 의도된 계산일 수 있으므로 사용자의 확인과 "
                            "정밀 검증이 필요합니다."
                        ),
                        sheet=worksheet.title,
                        cell=item.cell.coordinate,
                        evidence=_candidate_evidence(
                            pattern_type="DOMINANT_NORMALIZED_PATTERN_OUTLIER",
                            region=_horizontal_region(
                                row,
                                [entry.cell.column for entry in run],
                            ),
                            dominant_signature=dominant_signature,
                            current_signature=item.signature,
                            supporting_cells=[
                                entry.cell
                                for entry in sorted(
                                    supporting,
                                    key=lambda entry: entry.cell.column,
                                )
                            ],
                            detection_basis=(
                                "같은 행의 연속 수식 영역에서 최소 3개의 주변 수식이 같은 "
                                "정규화 패턴을 보이고, 대상 셀의 좌우 인접 수식도 그 패턴을 "
                                "따르지만 대상 셀만 다른 패턴입니다."
                            ),
                            candidate_cell=item.cell,
                        ),
                    )
                )


    for row, monthly_items in monthly_by_row.items():
        monthly_items.sort(key=lambda item: item.cell.column)
        for run in _contiguous_horizontal_formula_runs(monthly_items):
            if len(run) < 4:
                continue
            run_columns = [entry.cell.column for entry in run]
            if _has_hidden_column_between(worksheet, run_columns):
                continue
            if not _monthly_headers_are_sequential(run):
                continue
            counts = Counter(item.signature for item in run)
            dominant_signature, dominant_count = counts.most_common(1)[0]
            if not dominant_signature.startswith("MONTHLY_SUB("):
                continue
            if dominant_count < 3 or sum(count == dominant_count for count in counts.values()) != 1:
                continue
            deviations = [item for item in run if item.signature != dominant_signature]
            if len(deviations) != 1:
                continue
            item = deviations[0]
            if item.signature.startswith("MONTHLY_SUB("):
                continue
            if item is run[0] or item is run[-1]:
                continue
            by_column = {entry.cell.column: entry for entry in run}
            left = by_column.get(item.cell.column - 1)
            right = by_column.get(item.cell.column + 1)
            if left is None or right is None:
                continue
            if left.signature != dominant_signature or right.signature != dominant_signature:
                continue
            supporting_cells = _ordered_monthly_supporting_cells(run, item, dominant_signature)
            if len(supporting_cells) < 3:
                continue
            key = ("FORMULA_PATTERN_OUTLIER", item.cell.coordinate)
            if key in emitted:
                continue
            emitted.add(key)
            candidates.append(
                FormulaPatternCandidate(
                    rule_code="FORMULA_PATTERN_OUTLIER",
                    title="월별 시트 참조 패턴과 다른 수식 후보가 있습니다.",
                    description=(
                        "M01~M12 머리글과 같은 이름의 월별 시트를 참조하는 반복 뺄셈 수식 중 "
                        "한 셀만 주변 월 패턴과 다릅니다. 의도된 예외인지 확인이 필요합니다."
                    ),
                    sheet=worksheet.title,
                    cell=item.cell.coordinate,
                    evidence=_monthly_candidate_evidence(
                        region=_horizontal_region(row, run_columns),
                        dominant_signature=dominant_signature,
                        current_signature=item.signature,
                        supporting_cells=supporting_cells,
                    ),
                )
            )

    return FormulaPatternWorksheetAudit(
        candidates=candidates,
        supported_formula_count=supported_formula_count,
    )


def audit_worksheet_formula_patterns(worksheet: Worksheet) -> list[FormulaPatternCandidate]:
    """Return candidates while preserving the original M4-A helper contract."""

    return inspect_worksheet_formula_patterns(worksheet).candidates
