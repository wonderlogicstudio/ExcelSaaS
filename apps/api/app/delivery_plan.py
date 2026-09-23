"""Immutable exact repair proposal; no payment or execution occurs here."""

from __future__ import annotations

import copy
import hashlib
import json
import time
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .config import Settings
from .delivery_calculation import (
    CALCULATION_MODE_LEGACY,
    CALCULATION_MODE_MONTHLY_SHEETS,
    ENGINE_VERSION,
    REGISTRY_VERSION,
    calculate,
    engine_fingerprint,
    translate_formula,
)
from .delivery_inputs import (
    POLICY_VERSION,
    PROFILE_1,
    PROFILE_3,
    PROFILE_COMBINED,
    digest,
    inspect_input,
    policy_items,
    preflight,
    reject,
)
from .repair_rules.formula_restore import formula_restore_replacement
from .repair_rules.monthly_formula import monthly_formula_replacement
from .repair_rules.numeric_text import numeric_text_replacement

PLAN_VERSION = "repair-plan-v1"
TEMPLATE_VERSION = "plain-xlsx-three-artifacts-v1"
REQUIRED_ARTIFACTS = ["REPAIRED_XLSX", "CHANGES_XLSX", "VERIFICATION_HTML"]


def repair_rule_fingerprint() -> str:
    root = Path(__file__).parent
    return digest(
        {
            name: hashlib.sha256((root / name).read_bytes()).hexdigest()
            for name in [
                "formula_patterns.py",
                "repair_rules/monthly_formula.py",
                "repair_rules/formula_restore.py",
                "repair_rules/numeric_text.py",
            ]
        }
    )


def reference_status() -> dict:
    path = Path(__file__).with_name("delivery_reference.json")
    if not path.is_file():
        return {"status": "NOT_RUN"}
    result = json.loads(path.read_text(encoding="utf-8"))
    if result.get("engine_fingerprint") != engine_fingerprint():
        return {"status": "STALE"}
    return {
        "status": result["status"],
        "case_count": result["case_count"],
        "excel_version": result["excel_version"],
        "reference_digest": digest(result),
    }


def typed_source(cell: dict) -> dict:
    return {k: cell[k] for k in ["type", "value", "style"] if k in cell}


def plan_digest(plan: dict) -> str:
    return digest({k: v for k, v in plan.items() if k != "digest"})


def validate_plan(job: dict) -> dict:
    plan = job["state"].get("plan")
    if not plan or plan["digest"] != plan_digest(plan):
        reject("PLAN_INTEGRITY_FAILED", "변경계획을 다시 확인해야 합니다.", 409)
    if plan["expires_at"] <= time.time():
        reject("PLAN_EXPIRED", "변경계획이 만료되었습니다. 다시 계산하세요.", 409)
    for key, value in [
        ("owner", job["owner"]),
        ("job_id", job["id"]),
        ("source_hash", job["snapshot"]["source_hash"]),
        ("inventory_hash", job["snapshot"]["inventory_hash"]),
        ("engine_fingerprint", engine_fingerprint()),
        ("repair_rule_fingerprint", repair_rule_fingerprint()),
    ]:
        if plan.get(key) != value:
            reject("STALE_PLAN", "원본·소유자·계산 기준이 달라져 새 계획이 필요합니다.", 409)
    if plan["policy_digest"] != digest(job["state"]["policy"]):
        reject("STALE_PLAN", "업무 기준이 변경되었습니다.", 409)
    if not plan["patches"]:
        reject("NO_ELIGIBLE_CHANGES", "변경 대상이 없는 계획은 실행하지 않습니다.", 409)
    return plan


def build_plan(job: dict, policy: dict) -> dict:
    items = policy_items(policy)
    monthly_profile = len(items) == 1 and items[0].get("profile") == PROFILE_3
    snapshot = (
        inspect_input(
            "workbook.xlsx",
            job["source"],
            Settings(),
            profile_context=PROFILE_3,
        )
        if monthly_profile
        else job["snapshot"]
    )
    gate = preflight(snapshot, policy)
    if gate["status"] != "PRELIMINARY_ONLY":
        reject(
            "PREVIEW_VALIDATION_FAILED", "선택한 전체 범위가 사전 검사 조건을 충족해야 합니다.", 422
        )
    calculation_mode = (
        CALCULATION_MODE_MONTHLY_SHEETS if monthly_profile else CALCULATION_MODE_LEGACY
    )
    before = calculate(snapshot["cells"], calculation_mode=calculation_mode)
    allowed_before_error = (policy["sheet"], policy["targets"][0]) if monthly_profile else None
    existing_errors = [
        (error_sheet, error_cell, value)
        for error_sheet, rows in before["values"].items()
        for error_cell, value in rows.items()
        if value["type"] == "error"
    ]
    if any((sheet, cell) != allowed_before_error for sheet, cell, _value in existing_errors):
        reject(
            "EXISTING_CALCULATION_ERROR",
            "기존 계산 오류가 있어 이 수정 범위로 진행할 수 없습니다.",
            422,
        )
    if monthly_profile and (
        len(existing_errors) != 1
        or existing_errors[0][2].get("value") != "#VALUE!"
        or existing_errors[0][2].get("provenance") != "ENGINE_CALCULATED"
    ):
        reject(
            "PREVIEW_VALIDATION_FAILED",
            "월별 수식 교체 대상의 기존 계산 오류가 일치하지 않습니다.",
            422,
        )
    after_cells = copy.deepcopy(snapshot["cells"])
    patches = []
    item_gates = [
        gate if len(items) == 1 and item is policy else preflight(snapshot, item)
        for item in items
    ]
    seen: set[tuple[str, str]] = set()
    for index, item in enumerate(items):
        item_gate = item_gates[index]
        if item_gate["status"] != "PRELIMINARY_ONLY":
            reject(
                "PREVIEW_VALIDATION_FAILED",
                "선택한 전체 범위가 사전 검사 조건을 충족해야 합니다.",
                422,
            )
        for target in item_gate["targets"]:
            sheet = target["sheet"]
            address = target["cell"]
            key = (sheet, address)
            if key in seen:
                reject("OVERLAPPING_TARGETS", "같은 셀이 두 개 이상의 수정 기준과 중복됩니다.", 422)
            seen.add(key)
            current = after_cells[sheet].get(
                address, {"type": "blank", "value": None, "style": "0"}
            )
            if item["profile"] == PROFILE_1:
                replacement = numeric_text_replacement(current)
            elif item["profile"] == PROFILE_3:
                workbook = load_workbook(BytesIO(job["source"]), data_only=False, read_only=False)
                try:
                    monthly = monthly_formula_replacement(
                        workbook,
                        snapshot,
                        before["values"].get(sheet, {}).get(address, {}),
                        sheet=sheet,
                        cell=address,
                    )
                finally:
                    workbook.close()
                if monthly is None:
                    reject(
                        "PREVIEW_VALIDATION_FAILED",
                        "월별 시트 수식 교체 조건을 확인할 수 없습니다.",
                        422,
                    )
                replacement = {
                    **current,
                    "type": "formula",
                    "value": monthly.after_formula,
                    "cached": monthly.after_value["value"],
                }
            else:
                replacement = formula_restore_replacement(
                    current,
                    translate_formula(item["anchor_formula"], item["anchor"], address),
                )
            replacement.pop("cached", None)
            after_cells[sheet][address] = replacement
            previous = typed_source(current)
            updated = typed_source(replacement)
            patches.append(
                {
                    "candidate_id": digest(
                        [snapshot["source_hash"], sheet, address, previous, updated]
                    ),
                    "sheet": sheet,
                    "cell": address,
                    "before": previous,
                    "before_hash": digest(previous),
                    "after": updated,
                    "change_kind": target["change_kind"],
                    "policy_digest": item_gate["policy_digest"],
                    "policy_index": index,
                    "profile_version": item["profile"],
                    "anchor": item.get("anchor") if item["profile"] != PROFILE_1 else None,
                }
            )
    after = calculate(after_cells, calculation_mode=calculation_mode)
    impact = []
    auxiliary = []
    for sheet, rows in after_cells.items():
        for address, cell in rows.items():
            prior = (
                before["values"]
                .get(sheet, {})
                .get(address, {"type": "blank", "value": None, "provenance": "SOURCE_VALUE"})
            )
            future = after["values"][sheet][address]
            if future["type"] == "error":
                reject(
                    "PREVIEW_VALIDATION_FAILED",
                    "변경 후 새 계산 오류가 발생해 계획을 차단했습니다.",
                    422,
                )
            if cell["type"] == "formula":
                auxiliary.append(
                    {
                        "kind": "FORMULA_CACHE",
                        "sheet": sheet,
                        "cell": address,
                        "part": snapshot["sheet_parts"][sheet],
                        "after": future,
                    }
                )
            if prior["type"] != future["type"] or prior["value"] != future["value"]:
                impact.append({"sheet": sheet, "cell": address, "before": prior, "after": future})
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

    for sheet, rows in after_cells.items():
        if not rows:
            continue
        coordinates = [coordinate_to_tuple(a) for a in rows]
        bottom, right = max(r for r, c in coordinates), max(c for r, c in coordinates)
        if bottom * right > 10_000:
            reject("UNSUPPORTED_SPARSE_RANGE", "변경 후 사용 범위가 수정 한도를 초과합니다.", 422)
        if sheet in snapshot.get("dimensions", {}):
            top, left = min(r for r, c in coordinates), min(c for r, c in coordinates)
            dimension = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
            if snapshot["dimensions"][sheet] != dimension:
                auxiliary.append(
                    {
                        "kind": "SHEET_DIMENSION",
                        "part": snapshot["sheet_parts"][sheet],
                        "before": snapshot["dimensions"][sheet],
                        "after": dimension,
                    }
                )
    auxiliary.append(
        {
            "kind": "RECALCULATION_FLAGS",
            "part": "xl/workbook.xml",
            "attributes": {"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"},
        }
    )
    created = time.time()
    plan = {
        "schema_version": PLAN_VERSION,
        "owner": job["owner"],
        "job_id": job["id"],
        "product_id": job["product"],
        "sku": policy["profile"] if len(items) == 1 else PROFILE_COMBINED,
        "source_hash": snapshot["source_hash"],
        "inventory_hash": snapshot["inventory_hash"],
        "profile_version": policy["profile"] if len(items) == 1 else PROFILE_COMBINED,
        "policy_version": POLICY_VERSION,
        "policy_digest": gate["policy_digest"],
        "policy_items": copy.deepcopy(items),
        "engine_version": ENGINE_VERSION,
        "calculation_mode": calculation_mode,
        "engine_fingerprint": engine_fingerprint(),
        "repair_rule_fingerprint": repair_rule_fingerprint(),
        "registry_version": REGISTRY_VERSION,
        "template_version": TEMPLATE_VERSION,
        "patches": patches,
        "exact_targets": [[p["sheet"], p["cell"]] for p in patches],
        "impact": impact,
        "coverage": after["coverage"],
        "expected_calculated_values": after["values"],
        "value_provenance": "ENGINE_CALCULATED",
        "source_cache_used": False,
        "reference": reference_status(),
        "technical_changes": auxiliary,
        "created_at": created,
        "expires_at": min(created + 600, job["expires"]),
        "quote_id": None,
        "order_id": None,
        "required_artifacts": REQUIRED_ARTIFACTS,
        "exclusions": [],
    }
    plan["digest"] = plan_digest(plan)
    return plan


def plan_summary(plan: dict) -> dict:
    return {
        "digest": plan["digest"],
        "patch_count": len(plan["patches"]),
        "impact_count": len(plan["impact"]),
        "formula_impact_count": sum(
            any(
                a.get("kind") == "FORMULA_CACHE"
                and a.get("sheet") == row["sheet"]
                and a.get("cell") == row["cell"]
                for a in plan["technical_changes"]
            )
            for row in plan["impact"]
        ),
        "coverage": plan["coverage"],
        "reference": plan["reference"],
        "expires_at": plan["expires_at"],
        "exact_detail_available": False,
        "purchase_enabled": False,
        "source_unchanged": True,
        "status": "PREVIEW_VALIDATED"
        if plan["reference"]["status"] == "PASS"
        else "REFERENCE_NOT_VERIFIED",
    }


def customer_plan(job: dict, *, entitled: bool) -> dict:
    plan = validate_plan(job)
    if not entitled:
        reject(
            "ENTITLEMENT_REQUIRED", "정확한 변경계획은 유효한 수정 패키지 권리가 필요합니다.", 403
        )
    return {
        key: copy.deepcopy(value)
        for key, value in plan.items()
        if key not in {"owner", "expected_calculated_values"}
    }
