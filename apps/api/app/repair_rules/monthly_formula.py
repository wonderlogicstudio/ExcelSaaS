"""Stage-1 monthly formula replacement eligibility.

This module only derives a deterministic candidate. Delivery gates, approval,
patching, and execution remain separate stages.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.workbook.workbook import Workbook

from ..formula_patterns import (
    _formula_text,
    _month_header_for_cell,
    _parse_binary_subtraction_references,
    strict_monthly_reference_sheet_drift_candidate,
)

PROFILE_MONTHLY = "RP03_MONTHLY_SHEET_FORMULA_REPLACEMENT_V1"


@dataclass(frozen=True, slots=True)
class MonthlyFormulaReplacement:
    profile: str
    sheet: str
    cell: str
    before_formula: str
    before_value: dict
    after_formula: str
    after_value: dict
    evidence_cells: tuple[str, ...]
    approval_required: bool = True
    automatic_approval: bool = False


def _quote_sheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def _finite_number(record: dict | None) -> Decimal | None:
    if record is None or record.get("type") != "number":
        return None
    try:
        value = Decimal(str(record.get("value")))
    except (InvalidOperation, ValueError):
        return None
    if not math.isfinite(float(value)):
        return None
    return value


def _display_number(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _finite_display_number(value: Decimal) -> int | float | None:
    try:
        as_float = float(value)
    except (OverflowError, ValueError):
        return None
    if not math.isfinite(as_float):
        return None
    return _display_number(value)


def _snapshot_cell(snapshot: dict, sheet: str, cell: str) -> dict | None:
    return snapshot.get("cells", {}).get(sheet, {}).get(cell)


def _candidate_is_detected(workbook: Workbook, sheet: str, cell: str) -> bool:
    if sheet not in workbook.sheetnames:
        return False
    return strict_monthly_reference_sheet_drift_candidate(workbook[sheet], cell) is not None


def _same_direct_month_operands(
    workbook: Workbook, sheet: str, cell: str
) -> tuple[str, str, str] | None:
    worksheet = workbook[sheet]
    row, column = coordinate_to_tuple(cell)
    target_header = _month_header_for_cell(worksheet, worksheet[cell])
    if target_header is None:
        return None
    target_sheet = f"M{target_header[1]:02d}"
    if target_sheet not in workbook.sheetnames:
        return None

    operands: list[tuple[str, str]] = []
    for neighbor_column in (column - 1, column + 1):
        if neighbor_column < 1:
            return None
        neighbor = worksheet.cell(row=row, column=neighbor_column)
        neighbor_header = _month_header_for_cell(worksheet, neighbor)
        formula = _formula_text(neighbor)
        if neighbor_header is None or formula is None:
            return None
        expected_sheet = f"M{neighbor_header[1]:02d}"
        references = _parse_binary_subtraction_references(formula, neighbor)
        if references is None:
            return None
        first, second = references
        if (
            first.sheet != expected_sheet.casefold()
            or second.sheet != expected_sheet.casefold()
            or first.sheet != second.sheet
        ):
            return None
        operands.append((f"{first.column}{first.row}", f"{second.column}{second.row}"))
    if operands[0] != operands[1]:
        return None
    return target_sheet, operands[0][0], operands[0][1]


def monthly_formula_replacement(
    workbook: Workbook,
    snapshot: dict,
    before_result: dict,
    *,
    sheet: str,
    cell: str,
) -> MonthlyFormulaReplacement | None:
    """Return a monthly formula candidate, or None when any contract gate fails."""
    if not _candidate_is_detected(workbook, sheet, cell):
        return None
    if sheet not in workbook.sheetnames:
        return None
    worksheet = workbook[sheet]
    target = worksheet[cell]
    before_formula = _formula_text(target)
    if before_formula is None:
        return None
    snapshot_record = _snapshot_cell(snapshot, sheet, cell)
    if (
        snapshot_record is None
        or snapshot_record.get("type") != "formula"
        or snapshot_record.get("value") != before_formula
    ):
        return None
    if (
        before_result.get("type") != "error"
        or before_result.get("value") != "#VALUE!"
        or before_result.get("provenance") != "ENGINE_CALCULATED"
    ):
        return None
    references = _parse_binary_subtraction_references(before_formula, target)
    if references is None:
        return None
    first, second = references
    target_column = get_column_letter(target.column)
    if (
        first.sheet is not None
        or second.sheet is not None
        or first.column != target_column
        or second.column != target_column
    ):
        return None
    expected_before = f"={target_column}{first.row}-{target_column}{second.row}"
    if before_formula != expected_before:
        return None

    direct_operands = _same_direct_month_operands(workbook, sheet, cell)
    if direct_operands is None:
        return None
    target_month_sheet, first_operand, second_operand = direct_operands
    first_value = _finite_number(_snapshot_cell(snapshot, target_month_sheet, first_operand))
    second_value = _finite_number(_snapshot_cell(snapshot, target_month_sheet, second_operand))
    if first_value is None or second_value is None:
        return None
    after_number = first_value - second_value
    displayed_after_number = _finite_display_number(after_number)
    if displayed_after_number is None:
        return None
    quoted_target_sheet = _quote_sheet(target_month_sheet)
    after_formula = f"={quoted_target_sheet}!{first_operand}-{quoted_target_sheet}!{second_operand}"
    return MonthlyFormulaReplacement(
        profile=PROFILE_MONTHLY,
        sheet=sheet,
        cell=cell,
        before_formula=before_formula,
        before_value={"type": "error", "value": "#VALUE!"},
        after_formula=after_formula,
        after_value={"type": "number", "value": displayed_after_number},
        evidence_cells=(
            worksheet.cell(row=target.row, column=target.column - 1).coordinate,
            worksheet.cell(row=target.row, column=target.column + 1).coordinate,
        ),
    )
