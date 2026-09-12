"""Exact, typed two-source observations. Never infers a correct source or a repair."""

from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from zipfile import ZipFile

from defusedxml import ElementTree as ET
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

from .config import Settings
from .delivery_inputs import MAX_BYTES, NS, REL, digest, reject, valid_cell
from .security import validate_ooxml

PROFILE = "TWO_LEDGER_EXACT_KEY_KRW_V1"
VERSION = "exact-key-krw-v1"
MAX_ROWS = 20_000
MAX_COLUMNS = 50
MAX_SELECTED_CELLS = 250_000
MAX_FIELD = 4096
STATUSES = ["MATCHED", "AMOUNT_DIFF", "ONLY_A", "ONLY_B", "AMBIGUOUS", "INPUT_ERROR"]


def engine_fingerprint():
    from pathlib import Path

    root = Path(__file__).parent
    return digest(
        {
            name: hashlib.sha256((root / name).read_bytes()).hexdigest()
            for name in ["comparison_engine.py", "comparison_process.py", "comparison_worker.py"]
        }
    )


def source_value(kind, value):
    if isinstance(value, str) and len(value) > MAX_FIELD:
        reject("COMPARISON_FIELD_LIMIT", "단일 필드가 4,096자를 넘습니다.", 413)
    return {"type": kind, "value": value}


def read_source(filename: str, payload: bytes, settings: Settings) -> dict:
    if not payload or len(payload) > MAX_BYTES:
        reject("COMPARISON_INPUT_LIMIT", "현재 베타는 파일당 2MiB 이하를 지원합니다.", 413)
    result = {"source_hash": hashlib.sha256(payload).hexdigest(), "sheets": {}}
    if filename.lower().endswith(".csv"):
        try:
            text = payload.decode("utf-8-sig", errors="strict")
            rows = csv.reader(StringIO(text, newline=""), strict=True)
            cells = {}
            width = 0
            count = 0
            for index, row in enumerate(rows, 1):
                if index > MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                    reject("COMPARISON_TABLE_LIMIT", "CSV의 행·열 한도를 초과했습니다.", 413)
                width = max(width, len(row))
                count = index
                if width * count > MAX_SELECTED_CELLS:
                    reject("COMPARISON_TABLE_LIMIT", "CSV의 셀 한도를 초과했습니다.", 413)
                for col, value in enumerate(row, 1):
                    cells[f"{get_column_letter(col)}{index}"] = source_value("text", value)
            result.update(
                format="CSV", sheets={"CSV": {"cells": cells, "max_row": count, "max_col": width}}
            )
            return result
        except (UnicodeError, csv.Error):
            reject("COMPARISON_CSV_INVALID", "UTF-8 CSV의 문자 인코딩·따옴표·필드를 확인하세요.")
    if not filename.lower().endswith(".xlsx"):
        reject(
            "COMPARISON_FORMAT_UNSUPPORTED", "매크로 없는 XLSX 또는 UTF-8 CSV만 지원합니다.", 415
        )
    envelope = validate_ooxml(
        filename,
        payload,
        settings.model_copy(
            update={"max_uncompressed_mb": 20, "max_zip_entries": 256, "max_compression_ratio": 100}
        ),
    )
    if envelope.has_macros or envelope.external_link_part_count:
        reject(
            "COMPARISON_EXTERNAL_CONTENT",
            "매크로·외부 연결이 있는 자료는 값으로 내보낸 후 비교하세요.",
        )
    try:
        with ZipFile(BytesIO(payload)) as z:
            if len(z.namelist()) != len(set(z.namelist())):
                reject("COMPARISON_PACKAGE_INVALID", "중복 내부 파일이 있습니다.")
            for name in z.namelist():
                if name.endswith(".rels") and any(
                    n.get("TargetMode") == "External" for n in ET.fromstring(z.read(name))
                ):
                    reject(
                        "COMPARISON_EXTERNAL_CONTENT",
                        "외부 연결을 제거한 값 자료를 사용하세요.",
                    )
            workbook = ET.fromstring(z.read("xl/workbook.xml"))
            links = {
                n.get("Id"): n.get("Target", "")
                for n in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            }
            strings = []
            if "xl/sharedStrings.xml" in envelope.names:
                strings = [
                    "".join(n.text or "" for n in si.iter(NS + "t"))
                    for si in ET.fromstring(z.read("xl/sharedStrings.xml"))
                ]
                if any(len(s) > MAX_FIELD for s in strings):
                    reject("COMPARISON_FIELD_LIMIT", "문자 필드 한도를 초과했습니다.", 413)
            date_styles = set()
            if "xl/styles.xml" in envelope.names:
                styles = ET.fromstring(z.read("xl/styles.xml"))
                formats = {
                    **BUILTIN_FORMATS,
                    **{
                        int(n.get("numFmtId")): n.get("formatCode", "")
                        for n in styles.findall(NS + "numFmts/" + NS + "numFmt")
                    },
                }
                date_styles = {
                    i
                    for i, n in enumerate(styles.findall(NS + "cellXfs/" + NS + "xf"))
                    if is_date_format(formats.get(int(n.get("numFmtId", "0")), ""))
                }
            count = 0
            import posixpath

            for sheet in workbook.findall(NS + "sheets/" + NS + "sheet"):
                target = links[sheet.get(REL + "id")]
                path = posixpath.normpath(
                    target.lstrip("/") if target.startswith("/") else "xl/" + target
                )
                if path not in envelope.names or not path.startswith("xl/worksheets/"):
                    reject("COMPARISON_PACKAGE_INVALID", "시트 연결을 확인할 수 없습니다.")
                name = sheet.get("name", "")
                cells = {}
                max_row = 0
                max_col = 0
                tree = ET.fromstring(z.read(path))
                seen_rows = set()
                for parent in tree.findall(NS + "sheetData/" + NS + "row"):
                    row_id = int(parent.get("r"))
                    if row_id in seen_rows or row_id < 1 or row_id > 1048576:
                        reject("COMPARISON_PACKAGE_INVALID", "행 위치가 일관되지 않습니다.")
                    seen_rows.add(row_id)
                    for child in parent.findall(NS + "c"):
                        if coordinate_to_tuple(valid_cell(child.get("r")))[0] != row_id:
                            reject("COMPARISON_PACKAGE_INVALID", "셀과 행 위치가 다릅니다.")
                for cell in tree.findall(NS + "sheetData/" + NS + "row/" + NS + "c"):
                    address = valid_cell(cell.get("r"))
                    row, col = coordinate_to_tuple(address)
                    if address in cells:
                        reject("COMPARISON_PACKAGE_INVALID", "중복 셀 위치가 있습니다.")
                    max_row = max(max_row, row)
                    max_col = max(max_col, col)
                    count += 1
                    if count > MAX_SELECTED_CELLS:
                        reject(
                            "COMPARISON_TABLE_LIMIT",
                            "현재 베타의 파일 전체 셀 한도를 초과했습니다.",
                            413,
                        )
                    raw = cell.find(NS + "v")
                    value = raw.text if raw is not None else None
                    kind = cell.get("t", "n")
                    formula = cell.find(NS + "f")
                    if formula is not None:
                        record = source_value("formula", "=" + (formula.text or ""))
                    elif kind == "s":
                        index = int(value)
                        if index < 0 or index >= len(strings):
                            reject(
                                "COMPARISON_PACKAGE_INVALID", "문자열 참조 범위가 잘못되었습니다."
                            )
                        record = source_value("text", strings[index])
                    elif kind == "inlineStr":
                        record = source_value(
                            "text", "".join(n.text or "" for n in cell.iter(NS + "t"))
                        )
                    elif kind == "str":
                        record = source_value("text", value or "")
                    elif kind == "b":
                        record = source_value("boolean", value == "1")
                    elif kind in {"d", "e"}:
                        record = source_value("date" if kind == "d" else "error", value)
                    elif kind == "n" and value is not None:
                        record = source_value(
                            "date" if int(cell.get("s", "0")) in date_styles else "number", value
                        )
                    elif value is None:
                        record = source_value("blank", None)
                    else:
                        record = source_value("unsupported", value)
                    cells[address] = record
                if not name or name in result["sheets"]:
                    reject("COMPARISON_PACKAGE_INVALID", "시트 이름이 일관되지 않습니다.")
                result["sheets"][name] = {
                    "cells": cells,
                    "max_row": max_row,
                    "max_col": max_col,
                    "hidden_rows_included": True,
                }
            if not result["sheets"] or len(result["sheets"]) > 20:
                reject("COMPARISON_TABLE_LIMIT", "시트 개수 한도를 확인하세요.")
            result["format"] = "XLSX"
            return result
    except (KeyError, IndexError, ValueError, TypeError, ET.ParseError):
        reject("COMPARISON_PACKAGE_INVALID", "파일 내부 구조를 안전하게 읽을 수 없습니다.")


def preview(source):
    return {
        "source_hash": source["source_hash"],
        "format": source["format"],
        "sheets": [
            {
                "name": name,
                "max_row": s["max_row"],
                "max_col": s["max_col"],
                "preview": [
                    {
                        "row": r,
                        "cells": [
                            {
                                "column": get_column_letter(c),
                                **s["cells"].get(
                                    f"{get_column_letter(c)}{r}", {"type": "blank", "value": None}
                                ),
                            }
                            for c in range(1, min(s["max_col"], 5) + 1)
                        ],
                    }
                    for r in range(1, min(s["max_row"], 3) + 1)
                ],
            }
            for name, s in source["sheets"].items()
        ],
    }


def integer_amount(cell, *, csv_format=False):
    kind, value = cell["type"], cell["value"]
    if kind not in {"number", "text"} or not isinstance(value, str):
        return None
    if len(value) > MAX_FIELD:
        return None
    if kind == "text":
        pattern = r"[+-]?[0-9]+" if not csv_format else r"[+-]?(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)"
        if not re.fullmatch(pattern, value):
            return None
        value = value.replace(",", "")
    elif not re.fullmatch(r"[+-]?[0-9]+(?:\.0+)?", value):
        return None
    try:
        number = Decimal(value)
        return int(number) if number.is_finite() and number == number.to_integral_value() else None
    except (InvalidOperation, ValueError):
        return None


def select_rows(source: dict, side: str, selection: dict) -> dict:
    sheet = source["sheets"].get(selection.get("sheet"))
    if sheet is None:
        reject("COMPARISON_SELECTION_INVALID", "원천 시트를 선택하세요.")
    try:
        left, top, right, bottom = range_boundaries(selection.get("range", ""))
        keys = selection["key_columns"]
        amount = selection["amount_column"]
        header = selection["header_row"]
        if (
            not all(type(v) is int for v in [left, top, right, bottom, header])
            or not 1 <= top <= header < bottom <= 1048576
            or not 1 <= left <= right <= 16384
        ):
            raise ValueError()
        if (
            bottom - header > MAX_ROWS
            or right - left + 1 > MAX_COLUMNS
            or (right - left + 1) * (bottom - top + 1) > MAX_SELECTED_CELLS
        ):
            raise ValueError()
        columns = [get_column_letter(c) for c in range(left, right + 1)]
        if (
            not isinstance(keys, list)
            or not 1 <= len(keys) <= 3
            or len(set(keys)) != len(keys)
            or amount in keys
            or any(c not in columns for c in [*keys, amount])
        ):
            raise ValueError()
    except (KeyError, ValueError, TypeError):
        reject(
            "COMPARISON_SELECTION_INVALID",
            "헤더·범위·키 1~3열·금액 열과 20,000행/50열/250,000셀 한도를 확인하세요.",
        )
    cells = sheet["cells"]
    blank = {"type": "blank", "value": None}
    headers = [cells.get(c + str(header), blank) for c in [*keys, amount]]
    if any(v["type"] != "text" or not v["value"].strip() for v in headers) or len(
        {v["value"] for v in headers}
    ) != len(headers):
        reject(
            "COMPARISON_HEADER_INVALID", "선택한 키·금액 열의 실제 헤더가 비어 있거나 중복됩니다."
        )
    exclusions = selection.get("exclusions", [])
    if not isinstance(exclusions, list) or len(exclusions) > MAX_ROWS:
        reject("COMPARISON_EXCLUSION_INVALID", "제외 행과 사유를 확인하세요.")
    excluded = {}
    for item in exclusions:
        if (
            not isinstance(item, dict)
            or type(item.get("row")) is not int
            or not header < item["row"] <= bottom
            or item["row"] in excluded
            or not isinstance(item.get("reason"), str)
            or not item["reason"].strip()
            or len(item["reason"]) > 200
        ):
            reject(
                "COMPARISON_EXCLUSION_INVALID", "각 제외 행을 한 번만 지정하고 사유를 기록하세요."
            )
        excluded[item["row"]] = item["reason"]
    rows = []
    omitted = []
    for physical in range(header + 1, bottom + 1):
        rawkeys = [cells.get(c + str(physical), blank) for c in keys]
        rawamount = cells.get(amount + str(physical), blank)
        if any(v["type"] == "formula" for v in [*rawkeys, rawamount]) and physical not in excluded:
            reject(
                "COMPARISON_SELECTED_FORMULA",
                "선택한 키·금액의 수식을 값으로 내보내세요. 캐시는 사용하지 않습니다.",
            )
        key = (
            [v["value"] for v in rawkeys]
            if all(v["type"] == "text" and v["value"].strip() for v in rawkeys)
            else None
        )
        number = integer_amount(rawamount, csv_format=source["format"] == "CSV")
        row = {
            "source_id": side,
            "source_version": source["source_hash"],
            "source_row_id": f"{side}:{physical}",
            "data_row": physical - header,
            "sheet": selection["sheet"],
            "physical_row": physical,
            "raw_key_parts": rawkeys,
            "canonical_key": key,
            "raw_amount": rawamount,
            "amount": str(number) if number is not None else None,
            "errors": ([] if key is not None else ["INVALID_KEY"])
            + ([] if number is not None else ["INVALID_AMOUNT"]),
        }
        if physical in excluded:
            omitted.append({**row, "exclusion_reason": excluded[physical]})
        else:
            rows.append(row)
    return {
        "rows": rows,
        "excluded": omitted,
        "selected_count": bottom - header,
        "selection": selection,
        "source_hash": source["source_hash"],
    }


def compare_rows(a: dict, b: dict, *, tolerance: str = "0") -> dict:
    if not isinstance(tolerance, str) or not re.fullmatch(r"[0-9]{1,15}", tolerance):
        reject("COMPARISON_POLICY_INVALID", "오차는 0 이상의 정수 원 단위로 입력하세요.")
    selected = {"A": a, "B": b}
    grouped = defaultdict(lambda: {"A": [], "B": []})
    records = []
    for side, data in selected.items():
        for row in data["rows"]:
            if row["canonical_key"] is None:
                records.append(
                    {
                        "key": None,
                        "key_parts": None,
                        "status": "INPUT_ERROR",
                        "reason": "INVALID_KEY",
                        "A": [row] if side == "A" else [],
                        "B": [row] if side == "B" else [],
                        "amount_A": None,
                        "amount_B": None,
                        "delta": None,
                    }
                )
            else:
                grouped[tuple(row["canonical_key"])][side].append(row)
    for key, sides in sorted(grouped.items()):
        x, y = sides["A"], sides["B"]
        av = x[0]["amount"] if len(x) == 1 else None
        bv = y[0]["amount"] if len(y) == 1 else None
        delta = None
        reason = None
        if len(x) > 1 or len(y) > 1:
            status = "AMBIGUOUS"
            reason = "DUPLICATE_KEY"
        elif any(row["amount"] is None for row in [*x, *y]):
            status = "INPUT_ERROR"
            reason = "INVALID_AMOUNT"
        elif x and y:
            delta = str(int(av) - int(bv))
            status = "MATCHED" if abs(int(delta)) <= int(tolerance) else "AMOUNT_DIFF"
        else:
            status = "ONLY_A" if x else "ONLY_B"
        records.append(
            {
                "key": key[0] if len(key) == 1 else list(key),
                "key_parts": list(key),
                "status": status,
                "reason": reason,
                **sides,
                "amount_A": av,
                "amount_B": bv,
                "delta": delta,
            }
        )
    counts = {status: sum(r["status"] == status for r in records) for status in STATUSES}
    known = {
        side: sum(int(row["amount"]) for row in data["rows"] if row["amount"] is not None)
        for side, data in selected.items()
    }
    uncompared = {
        side: sum(
            int(row["amount"])
            for r in records
            if r["status"] in {"INPUT_ERROR", "AMBIGUOUS"}
            for row in r[side]
            if row["amount"] is not None
        )
        for side in ["A", "B"]
    }
    expected = (
        sum(int(r["delta"]) for r in records if r["delta"] is not None)
        + sum(int(r["amount_A"]) for r in records if r["status"] == "ONLY_A")
        - sum(int(r["amount_B"]) for r in records if r["status"] == "ONLY_B")
        + uncompared["A"]
        - uncompared["B"]
    )
    if known["A"] - known["B"] != expected:
        reject(
            "COMPARISON_CONSERVATION_FAILED",
            "알려진 금액 보존식이 일치하지 않아 결과를 차단했습니다.",
            422,
        )
    for side, data in selected.items():
        assigned = [row["source_row_id"] for r in records for row in r[side]]
        source = [row["source_row_id"] for row in data["rows"]]
        if (
            Counter(assigned) != Counter(source)
            or len(set(source)) != len(source)
            or len(source) + len(data["excluded"]) != data["selected_count"]
        ):
            reject(
                "COMPARISON_ROW_PRESERVATION_FAILED",
                "원천 행이 중복되거나 누락되어 보고서를 차단했습니다.",
                422,
            )
    for r in records:
        r["group_id"] = digest(
            [
                r["key_parts"],
                r["status"],
                [row["source_row_id"] for side in ["A", "B"] for row in r[side]],
            ]
        )
    summary = {
        "counts": counts,
        "group_count": len(records),
        "input_rows": {side: d["selected_count"] for side, d in selected.items()},
        "assigned_rows": {side: len(d["rows"]) for side, d in selected.items()},
        "excluded_rows": {side: len(d["excluded"]) for side, d in selected.items()},
        "known_amount_totals": {k: str(v) for k, v in known.items()},
        "uncompared_known_amounts": {k: str(v) for k, v in uncompared.items()},
        "excluded_known_amount_totals": {
            side: str(sum(int(row["amount"]) for row in d["excluded"] if row["amount"] is not None))
            for side, d in selected.items()
        },
        "unknown_amount_row_counts": {
            side: sum(row["amount"] is None for row in d["rows"]) for side, d in selected.items()
        },
        "rows_preserved": True,
        "known_amount_conservation": True,
        "whole_ledger_correctness_verified": False,
    }
    valid = counts["MATCHED"] + counts["AMOUNT_DIFF"]
    eligibility = (
        "NOT_ELIGIBLE"
        if not valid
        else "ELIGIBLE_WITH_LIMITATIONS"
        if counts["AMBIGUOUS"] + counts["INPUT_ERROR"]
        else "ELIGIBLE"
    )
    return {
        "profile": PROFILE,
        "engine_version": VERSION,
        "summary": summary,
        "records": records,
        "sources": selected,
        "eligibility": eligibility,
        "tolerance_krw": tolerance,
        "reference_b_is_truth": False,
        "includes_repaired_workbook": False,
    }


def compare_sources(sources: dict, spec: dict) -> dict:
    policy = spec.get("policy", {})
    if (
        policy.get("confirmed") is not True
        or policy.get("currency") != "KRW"
        or policy.get("unit") != "KRW_WON"
        or not isinstance(policy.get("period"), str)
        or not policy["period"].strip()
        or len(policy["period"]) > 80
        or not isinstance(policy.get("amount_meaning"), str)
        or not policy["amount_meaning"].strip()
        or len(policy["amount_meaning"]) > 80
        or policy.get("normalization") != "NONE"
    ):
        reject(
            "COMPARISON_POLICY_UNCONFIRMED",
            "양쪽 자료의 같은 기간·금액 의미·KRW 원 단위와 문자 키 기준을 직접 확인하세요.",
        )
    for side in ["A", "B"]:
        criteria = spec.get(side, {}).get("criteria")
        expected = {key: policy[key] for key in ["period", "amount_meaning", "currency", "unit"]}
        if criteria != expected:
            reject(
                "COMPARISON_CRITERIA_MISMATCH", "양쪽 자료의 기간·금액 의미·단위가 같아야 합니다."
            )
    a = select_rows(sources["A"], "A", spec.get("A", {}))
    b = select_rows(sources["B"], "B", spec.get("B", {}))
    result = compare_rows(a, b, tolerance=policy.get("tolerance_krw", "0"))
    result["spec_hash"] = digest(
        {
            "source_A": sources["A"]["source_hash"],
            "source_B": sources["B"]["source_hash"],
            "spec": spec,
            "profile": PROFILE,
            "engine": VERSION,
            "engine_fingerprint": engine_fingerprint(),
        }
    )
    result["policy"] = policy
    result["engine_fingerprint"] = engine_fingerprint()
    return result
