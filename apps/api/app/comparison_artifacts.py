"""New values-only comparison report files; neither observation is rewritten."""

from __future__ import annotations

import hashlib
import html
import json
from io import BytesIO

from lxml import html as html_parser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .comparison_engine import STATUSES
from .delivery_artifacts import text_row
from .delivery_inputs import reject

LABELS = {
    "MATCHED": "일치",
    "AMOUNT_DIFF": "금액 차이",
    "ONLY_A": "A에만",
    "ONLY_B": "B에만",
    "AMBIGUOUS": "중복·모호",
    "INPUT_ERROR": "자료오류",
}
SHEETS = {
    "MATCHED": "일치",
    "AMOUNT_DIFF": "금액차이",
    "ONLY_A": "한쪽자료",
    "ONLY_B": "한쪽자료",
    "AMBIGUOUS": "중복모호",
    "INPUT_ERROR": "자료오류",
}
REQUIRED = {"COMPARISON_REPORT_XLSX", "COMPARISON_VERIFICATION_HTML"}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def raw_text(cell):
    return "빈 셀" if cell["type"] == "blank" else str(cell["value"])


def row_values(record, row):
    return [
        record["group_id"],
        record["status"],
        row["source_id"],
        row["source_row_id"],
        row["sheet"],
        str(row["physical_row"]),
        json.dumps(row["raw_key_parts"], ensure_ascii=False),
        raw_text(row["raw_amount"]),
        row["raw_amount"]["type"],
        row["amount"] if row["amount"] is not None else "확인 불가",
        record["delta"] if record["delta"] is not None else "—",
        ", ".join(row["errors"]) or record["reason"] or "",
    ]


def report_xlsx(model, job):
    book = Workbook()
    guide = book.active
    guide.title = "안내"
    summary = book.create_sheet("요약")
    for name in ["금액차이", "한쪽자료", "중복모호", "자료오류", "일치", "범위및제외"]:
        book.create_sheet(name)
    text_row(guide, ["두 파일 비교 검토 보고서", "설명"])
    text_row(
        guide,
        [
            "원본 미변경",
            "A/B 모두 관측 자료입니다. B를 정답으로 사용하거나 원본을 바꾸지 않습니다.",
        ],
    )
    text_row(guide, ["제공 범위", "보고서 전용 · 수정본·정답표·손실확정 보고서가 아닙니다."])
    text_row(
        guide,
        ["합성 검증", "현재 내부 합성 검증 경로에서 생성했습니다. 실제 결제 검증이 아닙니다."],
    )
    for key, value in [
        ("job_id", job["id"]),
        ("spec_hash", model["spec_hash"]),
        ("engine_version", model["engine_version"]),
        ("engine_fingerprint", model["engine_fingerprint"]),
        ("profile", model["profile"]),
        ("source_A_hash", model["sources"]["A"]["source_hash"]),
        ("source_B_hash", model["sources"]["B"]["source_hash"]),
    ]:
        text_row(guide, [key, value])
    text_row(guide, ["기간·시점", model["policy"]["period"]])
    text_row(guide, ["금액 의미", model["policy"]["amount_meaning"]])
    text_row(guide, ["통화·단위", "KRW · 원"])
    text_row(guide, ["허용오차", model["tolerance_krw"] + "원"])
    text_row(
        guide,
        [
            "정밀도·표시",
            "값·차액은 정확한 정수 문자열입니다. 0은 실제 값입니다. "
            "—는 비교 불가, 확인 불가는 금액 오류입니다. 수식·링크를 실행하지 않습니다.",
        ],
    )
    text_row(summary, ["분류", "결과 그룹 수", "A 원천행 수", "B 원천행 수"])
    for status in STATUSES:
        selected = [r for r in model["records"] if r["status"] == status]
        text_row(
            summary,
            [
                LABELS[status],
                len(selected),
                sum(len(r["A"]) for r in selected),
                sum(len(r["B"]) for r in selected),
            ],
        )
    for key, label in [
        ("input_rows", "선택 데이터행"),
        ("assigned_rows", "결과에 배정한 행"),
        ("excluded_rows", "명시 제외행"),
        ("unknown_amount_row_counts", "미상 금액 행"),
        ("known_amount_totals", "알려진 금액 합"),
        ("uncompared_known_amounts", "미비교 알려진 금액"),
        ("excluded_known_amount_totals", "제외한 알려진 금액"),
    ]:
        text_row(summary, [label, "", model["summary"][key]["A"], model["summary"][key]["B"]])
    text_row(summary, ["원천 행 보존", "PASS"])
    text_row(summary, ["알려진 금액 보존식", "PASS"])
    text_row(summary, ["업무 정답 검증", "제공하지 않음"])
    for name in ["금액차이", "한쪽자료", "중복모호", "자료오류", "일치"]:
        text_row(
            book[name],
            [
                "그룹 ID",
                "분류",
                "원천",
                "원천 행 ID",
                "시트",
                "행",
                "키 원값·타입",
                "금액 원값",
                "원 타입",
                "정수 금액",
                "A−B 차액",
                "확인 이유",
            ],
        )
    for record in model["records"]:
        for side in ["A", "B"]:
            for row in record[side]:
                text_row(book[SHEETS[record["status"]]], row_values(record, row))
    for name in ["금액차이", "한쪽자료", "중복모호", "자료오류", "일치"]:
        if book[name].max_row == 1:
            text_row(book[name], ["0건 · 이 분류에 배정한 원천 행 없음"])
    scope = book["범위및제외"]
    text_row(
        scope, ["원천", "선택 시트·범위", "헤더행", "키 열", "금액 열", "제외 행 ID", "제외 사유"]
    )
    for side, source in model["sources"].items():
        selection = source["selection"]
        text_row(
            scope,
            [
                side,
                selection["sheet"] + "!" + selection["range"],
                selection["header_row"],
                ", ".join(selection["key_columns"]),
                selection["amount_column"],
                "",
                "",
            ],
        )
        for row in source["excluded"]:
            text_row(
                scope,
                [
                    side,
                    row["sheet"],
                    row["physical_row"],
                    "",
                    "",
                    row["source_row_id"],
                    row["exclusion_reason"],
                ],
            )
    text_row(
        scope,
        [
            "범위 안내",
            "범위 밖은 비교하지 않았습니다. 범위 안의 숨김·필터행은 포함합니다.",
        ],
    )
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="153D35")
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 24 if sheet.title == "요약" else 48
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=11, color="173B36")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F6F3")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[__import__("openpyxl").utils.get_column_letter(col)].width = 20
        if sheet == guide:
            sheet.column_dimensions["A"].width = 28
            sheet.column_dimensions["B"].width = 92
        elif sheet.title in SHEETS.values():
            sheet.column_dimensions["A"].hidden = True
            sheet.column_dimensions["G"].width = 42
            sheet.column_dimensions["H"].width = 24
            sheet.column_dimensions["L"].width = 25
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:1"
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def verification_html(model):
    def e(value):
        return html.escape(str(value), quote=True)

    summary = model["summary"]
    groups = "".join(
        f'<li data-status="{status}">{e(LABELS[status])}: '
        f"<strong>{summary['counts'][status]}그룹</strong> · "
        f"A {sum(len(r['A']) for r in model['records'] if r['status'] == status)}행 · "
        f"B {sum(len(r['B']) for r in model['records'] if r['status'] == status)}행</li>"
        for status in STATUSES
    )
    rows = "".join(
        f'<tr data-source-row="{e(row["source_row_id"])}">'
        f"<td>{e(LABELS[r['status']])}</td>"
        f"<td>{e(row['source_id'])} · {e(row['sheet'])} {row['physical_row']}행</td>"
        f"<td>{e(json.dumps(row['raw_key_parts'], ensure_ascii=False))}</td>"
        f"<td>{e(raw_text(row['raw_amount']))}</td>"
        f"<td>{e(row['amount'] if row['amount'] is not None else '확인 불가')}</td>"
        f"<td>{e(r['delta'] if r['delta'] is not None else '—')}</td></tr>"
        for r in model["records"]
        for side in ["A", "B"]
        for row in r[side]
    )
    info = "".join(
        f'<dt>{e(key)}</dt><dd data-info="{e(key)}">{e(value)}</dd>'
        for key, value in [
            ("비교 기간", model["policy"]["period"]),
            ("금액 의미", model["policy"]["amount_meaning"]),
            ("통화·단위", "KRW · 원"),
            ("허용오차", model["tolerance_krw"] + "원"),
            ("spec_hash", model["spec_hash"]),
            ("source_A_hash", model["sources"]["A"]["source_hash"]),
            ("source_B_hash", model["sources"]["B"]["source_hash"]),
            ("engine_version", model["engine_version"]),
            ("engine_fingerprint", model["engine_fingerprint"]),
        ]
    )
    template = (
        __import__("pathlib")
        .Path(__file__)
        .with_name("comparison_verification_template.html")
        .read_text(encoding="utf-8")
    )
    values = {
        "groups": groups,
        "rows": rows,
        "info": info,
        "total": str(summary["group_count"]),
        "a_rows": str(summary["input_rows"]["A"]),
        "b_rows": str(summary["input_rows"]["B"]),
        "a_total": summary["known_amount_totals"]["A"],
        "b_total": summary["known_amount_totals"]["B"],
        "a_unknown": str(summary["unknown_amount_row_counts"]["A"]),
        "b_unknown": str(summary["unknown_amount_row_counts"]["B"]),
    }
    for key, value in values.items():
        template = template.replace("@@" + key + "@@", value)
    return template.encode()


def make_comparison_artifacts(model, job):
    artifacts = {
        "COMPARISON_REPORT_XLSX": {
            "name": f"comparison_report_{job['id']}.xlsx",
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "data": report_xlsx(model, job),
        },
        "COMPARISON_VERIFICATION_HTML": {
            "name": f"comparison_verification_{job['id']}.html",
            "mime": "text/html;charset=utf-8",
            "data": verification_html(model),
        },
    }
    if sum(len(v["data"]) for v in artifacts.values()) > 20 * 1024**2:
        reject("COMPARISON_REPORT_LIMIT", "전체 보고서가 출력 한도를 초과했습니다.", 422)
    manifest = {
        "product_id": "TWO_FILE_COMPARISON",
        "job_id": job["id"],
        "spec_hash": model["spec_hash"],
        "source_A_hash": model["sources"]["A"]["source_hash"],
        "source_B_hash": model["sources"]["B"]["source_hash"],
        "engine_version": model["engine_version"],
        "engine_fingerprint": model["engine_fingerprint"],
        "profile": model["profile"],
        "expires_at": job["expires"],
        "source_unchanged": True,
        "includes_repaired_workbook": False,
        "files": {
            k: {
                "name": v["name"],
                "mime": v["mime"],
                "sha256": sha(v["data"]),
                "bytes": len(v["data"]),
            }
            for k, v in artifacts.items()
        },
    }
    package = {"artifacts": artifacts, "manifest": manifest}
    validate_comparison_artifacts(package, model)
    return package


def validate_comparison_artifacts(package, model):
    artifacts = package["artifacts"]
    manifest = package["manifest"]
    if set(artifacts) != REQUIRED or set(manifest["files"]) != REQUIRED:
        reject(
            "COMPARISON_REPORT_INCOMPLETE", "필수 보고서 두 파일을 모두 준비하지 못했습니다.", 422
        )
    for kind, artifact in artifacts.items():
        if sha(artifact["data"]) != manifest["files"][kind]["sha256"]:
            reject("COMPARISON_REPORT_INVALID", "보고서 무결성 확인에 실패했습니다.", 422)
    book = load_workbook(
        BytesIO(artifacts["COMPARISON_REPORT_XLSX"]["data"]), read_only=False, data_only=False
    )
    try:
        if book.sheetnames != [
            "안내",
            "요약",
            "금액차이",
            "한쪽자료",
            "중복모호",
            "자료오류",
            "일치",
            "범위및제외",
        ]:
            reject("COMPARISON_REPORT_INVALID", "보고서 시트가 누락되었습니다.", 422)
        metadata = dict(book["안내"].values)
        for key in ["spec_hash", "source_A_hash", "source_B_hash"]:
            value = (
                model["spec_hash"]
                if key == "spec_hash"
                else model["sources"][key[7]]["source_hash"]
            )
            if metadata.get(key) != value or manifest.get(key) != value:
                reject("COMPARISON_REPORT_INVALID", "원본·기준 참조가 다릅니다.", 422)
        for key, value in [
            ("기간·시점", model["policy"]["period"]),
            ("금액 의미", model["policy"]["amount_meaning"]),
            ("통화·단위", "KRW · 원"),
            ("허용오차", model["tolerance_krw"] + "원"),
        ]:
            if metadata.get(key) != value:
                reject("COMPARISON_REPORT_INVALID", "업무 확인 기준이 다릅니다.", 422)
        for index, status in enumerate(STATUSES, 2):
            records = [r for r in model["records"] if r["status"] == status]
            expected_summary = [
                LABELS[status],
                str(len(records)),
                str(sum(len(r["A"]) for r in records)),
                str(sum(len(r["B"]) for r in records)),
            ]
            if [c.value for c in book["요약"][index]] != expected_summary:
                reject("COMPARISON_REPORT_INVALID", "요약 분류·원천 행 수가 다릅니다.", 422)
        for index, key in enumerate(
            [
                "input_rows",
                "assigned_rows",
                "excluded_rows",
                "unknown_amount_row_counts",
                "known_amount_totals",
                "uncompared_known_amounts",
                "excluded_known_amount_totals",
            ],
            8,
        ):
            for col, side in [(3, "A"), (4, "B")]:
                if book["요약"].cell(index, col).value != str(model["summary"][key][side]):
                    reject("COMPARISON_REPORT_INVALID", "요약 금액·제외 수가 다릅니다.", 422)
        actual = {}
        for name in set(SHEETS.values()):
            for row in list(book[name].values)[1:]:
                if row[0] and str(row[0]).startswith("0건"):
                    continue
                if row[3] in actual:
                    reject(
                        "COMPARISON_ROW_PRESERVATION_FAILED",
                        "보고서에 중복 원천 행이 있습니다.",
                        422,
                    )
                actual[row[3]] = [v if v is not None else "" for v in row]
        expected = {
            row["source_row_id"]: [str(v) for v in row_values(r, row)]
            for r in model["records"]
            for side in ["A", "B"]
            for row in r[side]
        }
        if actual != expected:
            reject(
                "COMPARISON_ROW_PRESERVATION_FAILED",
                "보고서 행·타입·원천·정확한 금액이 결과와 다릅니다.",
                422,
            )
        if any(c.data_type == "f" or c.hyperlink for sheet in book for row in sheet for c in row):
            reject("COMPARISON_UNSAFE_REPORT", "보고서에 실행 수식 또는 링크가 있습니다.", 422)
        document = artifacts["COMPARISON_VERIFICATION_HTML"]["data"].decode()
        if document.count("data-source-row=") != len(expected):
            reject("COMPARISON_REPORT_INVALID", "HTML의 원천 행 수가 다릅니다.", 422)
        parsed = html_parser.fromstring(document)
        observed_rows = {}
        for element in parsed.xpath("//tr[@data-source-row]"):
            row_id = element.get("data-source-row")
            if row_id in observed_rows:
                reject("COMPARISON_REPORT_INVALID", "HTML 원천 행이 중복되었습니다.", 422)
            observed_rows[row_id] = [cell.text_content() for cell in element.findall("td")]
        html_expected = {
            row["source_row_id"]: [
                LABELS[record["status"]],
                f"{row['source_id']} · {row['sheet']} {row['physical_row']}행",
                json.dumps(row["raw_key_parts"], ensure_ascii=False),
                raw_text(row["raw_amount"]),
                row["amount"] if row["amount"] is not None else "확인 불가",
                record["delta"] if record["delta"] is not None else "—",
            ]
            for record in model["records"]
            for side in ["A", "B"]
            for row in record[side]
        }
        if observed_rows != html_expected:
            reject("COMPARISON_REPORT_INVALID", "HTML의 원천 행·위치·타입·금액이 다릅니다.", 422)
        for side in ["A", "B"]:
            total = parsed.xpath(f'//*[@data-total="{side}"]/text()')
            if total != [model["summary"]["known_amount_totals"][side]]:
                reject("COMPARISON_REPORT_INVALID", "HTML 알려진 금액 합이 다릅니다.", 422)
        for key in ["spec_hash", "source_A_hash", "source_B_hash"]:
            if f'data-info="{key}">{manifest[key]}</dd>' not in document:
                reject("COMPARISON_REPORT_INVALID", "HTML 검증 참조가 다릅니다.", 422)
    finally:
        book.close()
