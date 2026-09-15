"""Private, real customer artifacts, generated only after output verification."""

from __future__ import annotations

import hashlib
import html
import time
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .delivery_inputs import reject
from .delivery_plan import REQUIRED_ARTIFACTS

REPORT_VERSION = "repair-reports-v1"
LABELS = {
    "SOURCE_IMMUTABLE": "원본 SHA-256 보존",
    "EXACT_APPROVED_PATCH": "승인한 셀과 실제 변경 일치",
    "UNTOUCHED_MEMBERS": "비대상 내부 파일 보존",
    "NON_TARGET_XML_AND_STYLES": "비대상 셀·서식·구조 보존",
    "TYPED_FORMULA_CACHES": "계획된 계산 캐시와 타입 일치",
    "WHOLE_WORKBOOK_POST_CALCULATION": "전체 지원 수식 실제 후계산",
    "NO_NEW_CALCULATION_ERRORS": "계산 오류 없음",
    "NO_NEW_STATIC_FINDINGS": "신규 정적 위험 신호 없음",
    "EXCEL_FIXTURE_COMPATIBILITY": "지원 합성 파일의 Excel 재개봉 대조",
}
LABELS.update(
    {
        "TARGET_STATIC_FINDINGS_RESOLVED": "승인 대상 정적 발견 해소",
        "FORMULA_AUDIT_COMPLETED": "수식 후보 재검증 완료",
        "NO_NEW_FORMULA_CANDIDATES": "새 수식 후보 없음",
        "TARGET_FORMULA_CANDIDATES_RESOLVED": "승인 대상 수식 후보 해소",
    }
)

REQUIRED_VALIDATION_CODES = {
    "SOURCE_IMMUTABLE",
    "EXACT_APPROVED_PATCH",
    "UNTOUCHED_MEMBERS",
    "NON_TARGET_XML_AND_STYLES",
    "TYPED_FORMULA_CACHES",
    "WHOLE_WORKBOOK_POST_CALCULATION",
    "NO_NEW_CALCULATION_ERRORS",
    "NO_NEW_STATIC_FINDINGS",
    "EXCEL_FIXTURE_COMPATIBILITY",
}



def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def cell_text(cell: dict) -> str:
    if cell["type"] == "blank":
        return "빈 셀"
    if cell["type"] == "number" and isinstance(cell["value"], float) and cell["value"].is_integer():
        return str(int(cell["value"]))
    return str(cell["value"])


def text_row(sheet, values):
    sheet.append([str(v) if v is not None else "" for v in values])
    # Customer formula/string content is literal report text, never an executable formula.
    for c in sheet[sheet.max_row]:
        c.data_type = "s"
        c.alignment = Alignment(vertical="top", wrap_text=True)


def _detector_detail(item: dict) -> str:
    parts = [item.get("rule_code", "")]
    if item.get("pattern_subtype"):
        parts.append(str(item["pattern_subtype"]))
    return " / ".join(part for part in parts if part)


def _audit_status_label(status: str) -> str:
    return {
        "COMPLETED": "완료",
        "ABSTAINED_INSUFFICIENT_EVIDENCE": "증거 부족으로 판단 보류",
        "FAILED": "실패",
    }.get(status, "완료하지 못함")


def detector_rows(verification: dict):
    labels = {"static": "정적 구조 검사", "formula": "수식 후보 검사"}
    statuses = {
        "target_before": "원본 승인 대상 발견",
        "resolved": "수정본에서 해소됨",
        "target_remaining": "수정본에 승인 대상 남음",
        "new": "수정본 신규 발견",
        "remaining": "수정본에 남은 비대상 발견",
    }
    detectors = verification.get("detectors") or {}
    for detector, label in labels.items():
        summary = detectors.get(detector) or {}
        for status in ["target_before", "resolved", "target_remaining", "new", "remaining"]:
            for item in summary.get(status, []):
                yield [
                    label,
                    statuses[status],
                    item.get("sheet", ""),
                    item.get("cell", ""),
                    "같은 검사 기준으로 다시 확인했습니다.",
                    _detector_detail(item),
                ]
    if detectors.get("formula_comparison_status") == "NOT_ESTABLISHED":
        before_status = _audit_status_label(detectors.get("formula_audit_before_status", ""))
        after_status = _audit_status_label(detectors.get("formula_audit_after_status", ""))
        limitations = list(detectors.get("formula_audit_before_limitations") or []) + list(
            detectors.get("formula_audit_after_limitations") or []
        )
        detail = "; ".join(str(item) for item in limitations)
        yield [
            "수식 후보 검사",
            "비교 미확정",
            "",
            "",
            "수식 후보 비교에 필요한 증거가 부족해 완료 여부를 판단하지 않았습니다.",
            f"원본 {before_status}, 수정본 {after_status}"
            + (f"; 제한: {detail}" if detail else ""),
        ]


def detector_html(verification: dict) -> str:
    rows = []
    for detector, status, sheet, cell, description, detail in detector_rows(verification):
        shown_detail = detail if status == "비교 미확정" else ""
        rows.append(
            "<tr>"
            f"<td>{html.escape(detector, quote=True)}</td>"
            f"<td>{html.escape(status, quote=True)}</td>"
            f"<td>{html.escape(sheet, quote=True)}</td>"
            f"<td>{html.escape(cell, quote=True)}</td>"
            f"<td>{html.escape(description, quote=True)}</td>"
            f"<td>{html.escape(shown_detail, quote=True)}</td>"
            "</tr>"
        )
    if not rows:
        rows.append('<tr><td colspan="6">재검증에서 표시할 발견 또는 후보가 없습니다.</td></tr>')
    return (
        '<div class="table"><table><thead><tr>'
        "<th>검사</th><th>상태</th><th>시트</th><th>셀</th><th>설명</th><th>상세</th>"
        "</tr></thead><tbody>"
        + "".join(rows)
        + "</tbody></table></div>"
    )

def detector_summary_text(verification: dict) -> str:
    formula_summary = (
        "수식 후보: 비교 미확정(증거 부족)"
        if (verification.get("detectors") or {}).get("formula_comparison_status")
        == "NOT_ESTABLISHED"
        else f"수식 후보: 원본 {verification['formula_candidates_before']}건 → "
        f"수정본 {verification['formula_candidates_remaining']}건"
    )
    return (
        f"정적 구조 발견: 원본 {verification['static_before']}건 → "
        f"수정본 {verification['static_remaining']}건 · "
        + formula_summary
    )

def report_book(plan: dict, verification: dict) -> bytes:
    book = Workbook()
    overview = book.active
    overview.title = "전체 요약"
    text_row(overview, ["승인한 변경의 납품 내역", "내용"])
    text_row(overview, ["변경 수", len(plan["patches"])])
    text_row(overview, ["수정 종류", plan["profile_version"]])
    text_row(overview, ["원본", "변경하지 않음 · 수정본은 별도 파일"])
    text_row(
        overview, ["범위", "고객이 지정한 업무 기준의 승인 셀만 변경. 업무 정답을 보장하지 않음."]
    )
    text_row(overview, ["잔여 정적 발견", verification["static_remaining"]])
    text_row(
        overview,
        [
            "제외 항목",
            "선택 범위 밖은 변경하지 않음. 미지원 부분의 수정 또는 업무 정답 검증은 제공하지 않음.",
        ],
    )
    text_row(
        overview,
        ["계산 범위", f"전체 수식 {plan['coverage']['formula_count']}개 · 저장 캐시 미사용"],
    )
    changes = book.create_sheet("변경 셀")
    text_row(
        changes,
        [
            "시트",
            "셀",
            "이전 타입",
            "이전 값·수식",
            "이후 타입",
            "이후 값·수식",
            "변경 이유",
            "후계산 결과",
        ],
    )
    for p in plan["patches"]:
        calculated = plan["expected_calculated_values"][p["sheet"]][p["cell"]]
        text_row(
            changes,
            [
                p["sheet"],
                p["cell"],
                p["before"]["type"],
                cell_text(p["before"]),
                p["after"]["type"],
                cell_text(p["after"]),
                "승인한 타입 변환"
                if p["change_kind"] == "TYPE_NORMALIZATION"
                else "승인한 수식 복원",
                cell_text(calculated),
            ],
        )
    impacts = book.create_sheet("계산 영향")
    text_row(impacts, ["시트", "셀", "이전 타입", "이전 계산", "이후 타입", "이후 계산"])
    for p in plan["impact"]:
        text_row(
            impacts,
            [
                p["sheet"],
                p["cell"],
                p["before"]["type"],
                cell_text(p["before"]),
                p["after"]["type"],
                cell_text(p["after"]),
            ],
        )
    metadata = book.create_sheet("검증 정보")
    text_row(metadata, ["항목", "값"])
    for key in [
        "source_hash",
        "output_hash",
        "plan_digest",
        "engine_version",
        "static_scanner_version",
        "formula_audit_rule_set_version",
    ]:
        text_row(metadata, [key, verification[key]])
    text_row(metadata, ["report_version", REPORT_VERSION])
    for check in verification["checks"]:
        text_row(metadata, [LABELS.get(check["code"], check["code"]), check["status"]])
    detectors = book.create_sheet("탐지 재검증")
    text_row(detectors, ["검사", "상태", "시트", "셀", "설명", "상세"])
    rows = list(detector_rows(verification))
    if rows:
        for row in rows:
            text_row(detectors, row)
    else:
        text_row(detectors, ["정적 구조 검사", "표시할 항목 없음", "", "", "", ""])
        text_row(detectors, ["수식 후보 검사", "표시할 항목 없음", "", "", "", ""])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for c in sheet[1]:
            c.fill = PatternFill("solid", fgColor="153D35")
            c.font = Font(color="FFFFFF", bold=True, size=11)
        sheet.row_dimensions[1].height = 30
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 48
            for c in row:
                c.font = Font(name="맑은 고딕", size=11, color="173B36")
                if c.row % 2 == 0:
                    c.fill = PatternFill("solid", fgColor="F0F6F3")
        if sheet in [overview, metadata]:
            sheet.column_dimensions["A"].width = 34
            sheet.column_dimensions["B"].width = 92
        else:
            widths = {"A": 12, "B": 10, "C": 12, "D": 30, "E": 12, "F": 34, "G": 20, "H": 16}
            for col, width in widths.items():
                sheet.column_dimensions[col].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def verification_html(plan: dict, verification: dict) -> bytes:
    types = {
        "number": "숫자",
        "text": "문자",
        "formula": "수식",
        "blank": "빈 셀",
        "boolean": "논리값",
    }

    def shown(cell):
        value = f"{float(cell['value']):,.15g}" if cell["type"] == "number" else cell_text(cell)
        return html.escape(value, quote=True)

    def e(value):
        return html.escape(str(value), quote=True)

    def typed_shown(cell):
        if cell["type"] == "blank":
            return shown(cell)
        return e(types.get(cell["type"], cell["type"])) + " · " + shown(cell)

    checks = "".join(
        f'<li data-check="{e(c["code"])}">'
        f"<strong>{e(c['status'])}</strong> · {e(LABELS.get(c['code'], c['code']))}</li>"
        for c in verification["checks"]
    )
    changes = "".join(
        f'<tr data-change-cell="{e(p["cell"])}">'
        f"<td>{e(p['sheet'])} · {e(p['cell'])}</td>"
        f"<td>{typed_shown(p['before'])}</td>"
        f"<td>{typed_shown(p['after'])}"
        + (
            "<br>계산 결과: " + shown(plan["expected_calculated_values"][p["sheet"]][p["cell"]])
            if p["after"]["type"] == "formula"
            else ""
        )
        + "</td>"
        "</tr>"
        for p in plan["patches"]
    )
    impacts = "".join(
        f'<li data-result-cell="{e(p["cell"])}">'
        f"{e(p['sheet'])} · {e(p['cell'])}: {shown(p['before'])} → "
        f"<strong>{shown(p['after'])}</strong>"
        f"</li>"
        for p in plan["impact"]
        if not any(
            p["sheet"] == target["sheet"] and p["cell"] == target["cell"]
            for target in plan["patches"]
        )
    )
    provenance = "".join(
        f'<dt>{e(k)}</dt><dd data-provenance="{e(k)}">{e(verification[k])}</dd>'
        for k in [
            "source_hash",
            "output_hash",
            "plan_digest",
            "engine_version",
            "static_scanner_version",
            "formula_audit_rule_set_version",
        ]
    )
    from pathlib import Path

    rendered = (
        Path(__file__).with_name("delivery_verification_template.html").read_text(encoding="utf-8")
    )
    detector_section = (
        "    <section>\n"
        "<h2>탐지 재검증</h2>\n"
        "<p>원본에서 발견된 정적 구조 발견과 수식 후보를 같은 검사로 다시 확인했습니다. "
        "승인한 대상에서 원래 발견된 항목은 수정본에서 사라져야 하며, "
        "승인 범위 밖의 남은 항목은 그대로 표시합니다. "
        "이 표는 파일 전체가 업무적으로 맞다는 보장이 아닙니다.</p>\n"
        + detector_html(verification)
        + "\n</section>\n"
    )
    rendered = rendered.replace(
        "    <section>\n<details><summary>",
        detector_section + "    <section>\n<details><summary>",
    )
    values = [
        len(plan["patches"]),
        plan["coverage"]["formula_count"],
        verification["static_remaining"],
        verification["static_resolved"],
        checks,
        changes,
        impacts,
        provenance,
        detector_summary_text(verification),
    ]
    for index, value in enumerate(values):
        rendered = rendered.replace("@@" + str(index) + "@@", str(value))
    return rendered.encode()


def make_artifacts(
    job: dict,
    plan: dict,
    repaired: bytes,
    verification: dict,
    *,
    compatibility_bootstrap: bool = False,
) -> dict:
    if any(c["status"] != "PASS" for c in verification["checks"]):
        reject("REQUIRED_VALIDATION_INCOMPLETE", "필수 검증을 완료하지 못했습니다.", 422)
    artifacts = {
        "REPAIRED_XLSX": {
            "name": f"workbookcare_repaired_{job['id']}.xlsx",
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "data": repaired,
        },
        "CHANGES_XLSX": {
            "name": f"changes_{job['id']}.xlsx",
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "data": report_book(plan, verification),
        },
        "VERIFICATION_HTML": {
            "name": f"verification_{job['id']}.html",
            "mime": "text/html;charset=utf-8",
            "data": verification_html(plan, verification),
        },
    }
    manifest = {
        "product_id": job["product"],
        "job_id": job["id"],
        "sku": plan["sku"],
        "source_hash": plan["source_hash"],
        "output_hash": sha(repaired),
        "plan_digest": plan["digest"],
        "engine_version": plan["engine_version"],
        "engine_fingerprint": plan["engine_fingerprint"],
        "profile_version": plan["profile_version"],
        "report_version": REPORT_VERSION,
        "created_at": time.time(),
        "expires_at": job["expires"],
        "verification": verification,
        "files": {
            k: {
                "name": v["name"],
                "mime": v["mime"],
                "sha256": sha(v["data"]),
                "bytes": len(v["data"]),
            }
            for k, v in artifacts.items()
        },
    }
    package = {"manifest": manifest, "artifacts": artifacts}
    validate_artifacts(package, plan, require_compatibility=not compatibility_bootstrap)
    return package


def validate_artifacts(package: dict, plan: dict, *, require_compatibility: bool = True):
    artifacts = package["artifacts"]
    manifest = package["manifest"]
    if set(artifacts) != set(REQUIRED_ARTIFACTS) or set(manifest["files"]) != set(
        REQUIRED_ARTIFACTS
    ):
        reject("INCOMPLETE_PACKAGE", "필수 세 파일을 모두 준비하지 못했습니다.", 422)
    if sum(len(v["data"]) for v in artifacts.values()) > 8 * 1024**2:
        reject("PACKAGE_LIMIT_EXCEEDED", "납품 파일 크기 한도를 초과했습니다.", 422)
    for kind, artifact in artifacts.items():
        if (
            not artifact["data"]
            or sha(artifact["data"]) != manifest["files"][kind]["sha256"]
            or len(artifact["data"]) != manifest["files"][kind]["bytes"]
        ):
            reject("ARTIFACT_INTEGRITY_FAILED", "파일 무결성을 확인하지 못했습니다.", 422)
    expected = {
        "source_hash": plan["source_hash"],
        "output_hash": sha(artifacts["REPAIRED_XLSX"]["data"]),
        "plan_digest": plan["digest"],
    }
    required = set(REQUIRED_VALIDATION_CODES) - (
        set() if require_compatibility else {"EXCEL_FIXTURE_COMPATIBILITY"}
    )
    passed = {c["code"] for c in manifest["verification"]["checks"] if c["status"] == "PASS"}
    if not required.issubset(passed):
        reject("REQUIRED_VALIDATION_INCOMPLETE", "필수 검증이 누락되었거나 실패했습니다.", 422)
    book = load_workbook(
        BytesIO(artifacts["CHANGES_XLSX"]["data"]), read_only=True, data_only=False
    )
    try:
        metadata = dict(book["검증 정보"].values)
        rows = list(book["변경 셀"].values)[1:]
        if len(rows) != len(plan["patches"]):
            reject("REPORT_MISMATCH", "변경내역 행 수가 다릅니다.", 422)
        for row, p in zip(rows, plan["patches"], strict=True):
            expected_row = [
                p["sheet"],
                p["cell"],
                p["before"]["type"],
                cell_text(p["before"]),
                p["after"]["type"],
                cell_text(p["after"]),
            ]
            if list(row[:6]) != expected_row:
                reject("REPORT_MISMATCH", "변경내역과 계획이 다릅니다.", 422)
        rendered = artifacts["VERIFICATION_HTML"]["data"].decode()
        for key, value in expected.items():
            if (
                metadata.get(key) != value
                or manifest.get(key) != value
                or manifest["verification"].get(key) != value
                or f'data-provenance="{key}">{value}</dd>' not in rendered
            ):
                reject("REPORT_MISMATCH", "원본·계획·수정본 참조가 다릅니다.", 422)
        if any(c.data_type == "f" for sheet in book for row in sheet for c in row):
            reject("REPORT_FORMULA_FORBIDDEN", "변경내역에 실행 수식이 있습니다.", 422)
    finally:
        book.close()
